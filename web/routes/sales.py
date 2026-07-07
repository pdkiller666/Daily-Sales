import asyncio
import io
import logging
from datetime import date
from typing import Annotated
from fastapi import APIRouter, Request, Form
from fastapi.responses import RedirectResponse, StreamingResponse, JSONResponse
from timezone_utils import DEFAULT_TZ
from web.response_utils import content_disposition as _cd

router = APIRouter()
PAGE_SIZE = 50


def _summary_empty():
    return (0, 0, 0, 0)


def _get_internal_uid(db, telegram_id: int):
    """Return internal users.id for this telegram_id, or None."""
    try:
        conn = db.get_connection()
        try:
            cur = conn.cursor()
            cur.execute("SELECT id FROM users WHERE telegram_id = ?", (telegram_id,))
            row = cur.fetchone()
        finally:
            conn.close()
        return row[0] if row else None
    except Exception:
        return None


def _get_user_shop_from_db(db, telegram_id: int):
    """Fallback: read user's shop_name directly from the org DB users table."""
    try:
        conn = db.get_connection()
        try:
            cur = conn.cursor()
            cur.execute("SELECT shop_name FROM users WHERE telegram_id = ?", (telegram_id,))
            row = cur.fetchone()
        finally:
            conn.close()
        return row[0] if row and row[0] else None
    except Exception:
        return None


def _get_user_allowed_shops(telegram_id: int, db) -> list:
    """Return list of shop names the user is allowed to access.
    Applies scope restrictions; owners/unrestricted users get all shops.
    Falls back to the user's shop_name from the org DB users table on any failure."""
    from db_utils import get_user_org_scope
    all_shops = db.get_all_shops() or []
    try:
        scope_type, scope_values = get_user_org_scope(telegram_id)
        if not scope_type or scope_type == "all":
            return all_shops
        if scope_type == "shop":
            filtered = [s for s in all_shops if s in scope_values]
            if filtered:
                return filtered
        elif scope_type in ("city", "network", "trade_network"):
            # Whitelist: col is only ever "city" or "trade_network" — no user input reaches here
            col = "city" if scope_type == "city" else "trade_network"
            assert col in ("city", "trade_network"), f"Unexpected col: {col}"
            placeholders = ",".join("?" * len(scope_values))
            conn = db.get_connection()
            try:
                cur = conn.cursor()
                cur.execute(
                    f"SELECT DISTINCT shop_name FROM users WHERE {col} IN ({placeholders}) AND shop_name IS NOT NULL",
                    scope_values,
                )
                allowed = {row[0] for row in cur.fetchall()}
            finally:
                conn.close()
            filtered = [s for s in all_shops if s in allowed]
            if filtered:
                return filtered
    except Exception:
        pass
    # Fallback: read user's shop directly from org DB
    user_shop = _get_user_shop_from_db(db, telegram_id)
    if user_shop and user_shop in all_shops:
        return [user_shop]
    return all_shops


@router.get("/api/product-motivation")
def api_product_motivation(request: Request, product_id: int = 0):
    """Return motivation info for a given product (for sale preview)."""
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        if not product_id:
            return JSONResponse({"none": True})
        db = get_web_db(telegram_id, org_db)
        info = db.get_product_motivation(product_id)
        if not info:
            return JSONResponse({"none": True})
        return JSONResponse({
            "none": False,
            "motivation_type": info["motivation_type"],
            "motivation_value": float(info["motivation_value"] or 0),
        })
    except Exception as exc:
        return JSONResponse({"none": True, "error": "Внутренняя ошибка сервера"})


@router.get("/api/products-for-shop")
def api_products_for_shop(request: Request, shop: str = ""):
    """Return JSON list of products with stock > 0 in the given shop."""
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return JSONResponse({"error": "Unauthorized", "products": []}, status_code=401)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        allowed_shops = _get_user_allowed_shops(telegram_id, db)
        if shop and shop not in allowed_shops:
            # Admins can query any org shop (cross-shop)
            if user.get("role") not in ("owner", "admin", "super_admin"):
                return JSONResponse({"error": "Forbidden", "products": []}, status_code=403)
        raw = db.get_all_inventory(shop_name=shop if shop else None) or []
        # id[0] product_id[1] shop_name[2] quantity[3] name[6] category[7] price[8]
        products = sorted(
            [
                {"id": r[1], "name": r[6] or "", "category": r[7] or "",
                 "price": float(r[8] or 0), "stock": int(r[3] or 0)}
                for r in raw if int(r[3] or 0) > 0
            ],
            key=lambda p: p["name"].lower(),
        )
        return JSONResponse({"products": products})
    except Exception as e:
        return JSONResponse({"error": "Внутренняя ошибка сервера", "products": []})


@router.get("/sales")
def sales_page(
    request: Request,
    date_from: str = "",
    date_to: str = "",
    shop: str = "",
    seller_id: int = 0,
    category: str = "",
    product_id: int = 0,
    page: int = 1,
    sort_order: str = "desc",
    sort_col: str = "date",
    only_backdated: int = 0,
):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    sort_order = sort_order if sort_order in ("asc", "desc") else "desc"
    sort_col = sort_col if sort_col in ("date", "amount", "qty", "product", "shop") else "date"
    ctx: dict = {
        "request": request, "user": user,
        "is_admin": user.get("role") in ("owner", "admin", "super_admin"),
        "sales": [], "shops": [], "sellers": [], "categories": [],
        "date_from": date_from, "date_to": date_to, "selected_shop": shop,
        "selected_seller_id": seller_id, "selected_category": category,
        "selected_product_id": product_id, "product_name": "",
        "page": 1, "total_pages": 1, "total_count": 0,
        "summary": _summary_empty(), "error": None,
        "csrf_token": get_csrf_token(request),
        "flash_ok": request.query_params.get("ok") == "1",
        "flash_err": request.query_params.get("error", ""),
        "current_user_id": None,
        "user_tz": DEFAULT_TZ,
        "sort_order": sort_order,
        "sort_col": sort_col,
        "only_backdated": bool(only_backdated),
    }

    try:
        db = get_web_db(telegram_id, org_db)

        from timezone_utils import get_current_user_time
        tz = db.get_user_timezone(telegram_id)
        ctx["user_tz"] = tz or DEFAULT_TZ
        today = get_current_user_time(tz).date()
        month_start = today.replace(day=1)

        if not date_from:
            date_from = month_start.isoformat()
        if not date_to:
            date_to = today.isoformat()
        ctx["date_from"] = date_from
        ctx["date_to"] = date_to

        ctx["shops"] = _get_user_allowed_shops(telegram_id, db)
        ctx["current_user_id"] = _get_internal_uid(db, telegram_id)
        try:
            ctx["all_shops"] = db.get_all_shops() or []
        except Exception:
            ctx["all_shops"] = ctx["shops"]

        # Build sellers list for dropdown
        try:
            all_users = db.get_all_users() or []
            # users: id[0] telegram_id[1] first_name[2] last_name[3] ... shop_name[8]
            sellers = [
                {"id": u[0], "name": f"{u[2] or ''} {u[3] or ''}".strip()}
                for u in all_users if u[8] not in ("Системный", "System", None)
            ]
            ctx["sellers"] = sorted(sellers, key=lambda x: x["name"].lower())
        except Exception:
            pass

        # Build categories list for dropdown
        try:
            cats = db.get_all_categories() or []
            ctx["categories"] = [c for c in cats if c]
        except Exception:
            pass

        # Determine if this user has a restricted shop scope
        # Applies to non-admin AND scoped admin (admin with limited shop list)
        _all_shops_set = set(ctx.get("all_shops", ctx["shops"]))
        _user_scoped = (
            bool(ctx["shops"])
            and set(ctx["shops"]) != _all_shops_set
        )

        kwargs: dict = {"start_date": date_from, "end_date": date_to}
        sum_kwargs: dict = {"start_date": date_from, "end_date": date_to}
        if shop:
            # Validate requested shop against allowed scope for all users
            if _user_scoped and shop not in ctx["shops"]:
                shop = ctx["shops"][0] if ctx["shops"] else shop
                ctx["selected_shop"] = shop
            kwargs["shop_name"] = shop
            sum_kwargs["shop_name"] = shop
        elif _user_scoped:
            # Auto-apply scope: user only sees sales from their allowed shops
            kwargs["shop_names"] = ctx["shops"]
            sum_kwargs["shop_names"] = ctx["shops"]

        if category:
            kwargs["category"] = category
            sum_kwargs["category"] = category

        if seller_id:
            kwargs["user_id"] = seller_id
        if product_id:
            kwargs["product_id"] = product_id

        all_sales = db.get_sales_report(**kwargs) or []

        # Backdated filter: keep only sales whose date != today in user's timezone
        if only_backdated:
            import zoneinfo as _zi
            from datetime import datetime as _dt
            _tz_obj = _zi.ZoneInfo(tz or DEFAULT_TZ)
            _today = _dt.now(_tz_obj).date()

            def _sale_is_backdated(sd):
                if not sd:
                    return False
                try:
                    parsed = _dt.strptime(str(sd)[:19], "%Y-%m-%d %H:%M:%S")
                    local = parsed.replace(tzinfo=_zi.ZoneInfo("UTC")).astimezone(_tz_obj)
                    return local.date() != _today
                except Exception:
                    return False

            all_sales = [s for s in all_sales if _sale_is_backdated(s[6])]

        # Имя продукта для заголовка (product_id уже передан как SQL-фильтр)
        if product_id:
            try:
                _p = db.get_product(product_id)
                if _p:
                    ctx["product_name"] = _p[1]
            except Exception:
                pass

        # Apply sorting: DB returns DESC by date by default
        if sort_col == "date":
            if sort_order == "asc":
                all_sales = list(reversed(all_sales))
        elif sort_col == "amount":
            all_sales = sorted(
                all_sales,
                key=lambda s: float(s[3] or 0) * float(s[4] or 0),
                reverse=(sort_order == "desc"),
            )
        elif sort_col == "qty":
            all_sales = sorted(
                all_sales,
                key=lambda s: float(s[3] or 0),
                reverse=(sort_order == "desc"),
            )
        elif sort_col == "product":
            all_sales = sorted(
                all_sales,
                key=lambda s: (s[7] or "").lower(),
                reverse=(sort_order == "desc"),
            )
        elif sort_col == "shop":
            all_sales = sorted(
                all_sales,
                key=lambda s: (s[2] or "").lower(),
                reverse=(sort_order == "desc"),
            )

        total = len(all_sales)
        total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
        page = max(1, min(page, total_pages))
        start = (page - 1) * PAGE_SIZE

        ctx["sales"] = all_sales[start : start + PAGE_SIZE]
        ctx["total_count"] = total
        ctx["total_pages"] = total_pages
        ctx["page"] = page
        if product_id:
            # get_sales_summary has no product filter — recompute from filtered rows
            _cnt = len(all_sales)
            _qty = sum(int(s[3] or 0) for s in all_sales)
            _rev = sum(float(s[3] or 0) * float(s[4] or 0) for s in all_sales)
            ctx["summary"] = (_cnt, _qty, _rev, (_rev / _cnt if _cnt else 0))
        else:
            ctx["summary"] = db.get_sales_summary(**sum_kwargs) or _summary_empty()

    except Exception as exc:
        ctx["error"] = "Произошла внутренняя ошибка. Попробуйте позже."

    return request.app.state.templates.TemplateResponse(
        request, "sales/index.html", ctx
    )


@router.get("/sales/export.xlsx")
def sales_export_xlsx(
    request: Request,
    date_from: str = "",
    date_to: str = "",
    shop: str = "",
):
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        db = get_web_db(telegram_id, org_db)

        today = date.today()
        if not date_from:
            date_from = today.replace(day=1).isoformat()
        if not date_to:
            date_to = today.isoformat()

        allowed_shops = _get_user_allowed_shops(telegram_id, db)
        all_shops = db.get_all_shops() or []
        # Apply scope to all users including scoped admins
        _scoped = bool(allowed_shops) and set(allowed_shops) != set(all_shops)

        kwargs: dict = {"start_date": date_from, "end_date": date_to}
        if shop:
            # Validate shop against allowed scope
            if _scoped and shop not in allowed_shops:
                shop = allowed_shops[0] if allowed_shops else shop
            kwargs["shop_name"] = shop
        elif _scoped:
            kwargs["shop_names"] = allowed_shops

        # id[0] pid[1] shop[2] qty[3] price[4] uid[5] date[6]
        # product_name[7] category[8] first_name[9] last_name[10]
        sales = db.get_sales_report(**kwargs) or []

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Продажи"

        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        hdr_fill = PatternFill("solid", fgColor="1E40AF")
        hdr_font = Font(bold=True, color="FFFFFF", size=11)
        even_fill = PatternFill("solid", fgColor="F0F4FF")
        tot_fill = PatternFill("solid", fgColor="DBEAFE")
        meta_font = Font(italic=True, color="4B5563", size=10)

        # ── Мета-заголовок (строки 1-3) ──────────────────────────────────────
        ws["A1"] = f"Период: {date_from} — {date_to}"
        ws["A1"].font = meta_font
        ws["A2"] = f"Магазин: {shop or 'Все'}"
        ws["A2"].font = meta_font
        ws["A3"] = f"Экспортировано: {today.strftime('%d.%m.%Y')}"
        ws["A3"].font = meta_font

        # ── Заголовки таблицы (строка 5) ─────────────────────────────────────
        # 9 колонок: добавлена «Время» и переименована «Цена»
        headers = ["Дата", "Время", "Товар", "Категория", "Магазин",
                   "Кол-во", "Цена (₽)", "Сумма (₽)", "Продавец"]
        col_widths = [12, 10, 30, 18, 20, 8, 14, 14, 22]
        HDR_ROW = 5
        for i, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=HDR_ROW, column=i, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[HDR_ROW].height = 28
        ws.freeze_panes = f"A{HDR_ROW + 1}"

        total_amount = 0.0
        total_qty = 0
        tx_count = 0
        for row_idx, s in enumerate(sales, HDR_ROW + 1):
            qty = int(s[3] or 0)
            price = float(s[4] or 0)
            amount = qty * price
            total_amount += amount
            total_qty += qty
            tx_count += 1
            seller = f"{s[9] or ''} {s[10] or ''}".strip()
            row_fill = even_fill if row_idx % 2 == 0 else None

            raw_dt = str(s[6] or "")
            try:
                from timezone_utils import get_user_time as _gut
                _tz = db.get_user_timezone(telegram_id) or DEFAULT_TZ
                _local = _gut(raw_dt, _tz)
                if _local:
                    date_part = _local.strftime("%Y-%m-%d")
                    time_part = _local.strftime("%H:%M")
                else:
                    date_part = raw_dt[:10]
                    time_part = raw_dt[11:16] if len(raw_dt) > 10 else ""
            except Exception:
                date_part = raw_dt[:10]
                time_part = raw_dt[11:16] if len(raw_dt) > 10 else ""

            for col_idx, val in enumerate(
                [date_part, time_part, s[7] or "", s[8] or "", s[2] or "",
                 qty, price, amount, seller], 1
            ):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = border
                if row_fill:
                    cell.fill = row_fill
                if col_idx in (7, 8):
                    cell.number_format = '#,##0.00 ₽'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx == 6:
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in (1, 2):
                    cell.alignment = Alignment(horizontal="center")

        # ── Итого ────────────────────────────────────────────────────────────
        if sales:
            tr = len(sales) + HDR_ROW + 1
            for col in range(1, 10):
                ws.cell(row=tr, column=col).border = border
                ws.cell(row=tr, column=col).fill = tot_fill
            ws.cell(row=tr, column=1, value="ИТОГО").font = Font(bold=True)
            ws.cell(row=tr, column=6, value=total_qty).font = Font(bold=True)
            total_cell = ws.cell(row=tr, column=8, value=total_amount)
            total_cell.font = Font(bold=True)
            total_cell.number_format = '#,##0.00 ₽'
            total_cell.alignment = Alignment(horizontal="right")

            avg_check = total_amount / tx_count if tx_count else 0
            tr2 = tr + 1
            ws.cell(row=tr2, column=1, value="Ср. чек").font = Font(italic=True, color="6B7280")
            avg_cell = ws.cell(row=tr2, column=8, value=avg_check)
            avg_cell.number_format = '#,##0.00 ₽'
            avg_cell.font = Font(italic=True, color="6B7280")
            avg_cell.alignment = Alignment(horizontal="right")

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f"sales_{date_from}__{date_to}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": _cd(filename)},
        )

    except Exception as exc:
        logging.getLogger(__name__).error(f"sales_export error: {exc}", exc_info=True)
        return RedirectResponse(url=f"/sales?error=Ошибка+при+экспорте.+Попробуйте+позже.", status_code=302)


@router.post("/sales/create")
def sales_create(
    request: Request,
    product_id: Annotated[int, Form()],
    shop_name: Annotated[str, Form()],
    quantity: Annotated[int, Form()],
    sale_price: Annotated[float, Form()],
    csrf_token: str = Form(default=""),
    cross_shop: str = Form(default=""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/sales?error=CSRF+error", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        allowed_shops = _get_user_allowed_shops(telegram_id, db)
        if shop_name not in allowed_shops:
            # Cross-shop: admins can sell from any org shop
            if cross_shop == "1" and user.get("role") in ("owner", "admin", "super_admin"):
                all_org_shops = db.get_all_shops() or []
                if shop_name not in all_org_shops:
                    return RedirectResponse(url="/sales?error=Магазин+не+найден", status_code=302)
            else:
                return RedirectResponse(url="/sales?error=Магазин+недоступен", status_code=302)
        internal_uid = _get_internal_uid(db, telegram_id)
        if not internal_uid:
            return RedirectResponse(url="/sales?error=Пользователь+не+найден", status_code=302)
        if quantity < 1:
            return RedirectResponse(url="/sales?error=Неверное+количество", status_code=302)

        # Subscription limit: enforce monthly sale cap in web layer (mirrors bot check_sales_limit)
        try:
            from subscription_utils import check_sales_limit
            _ok, _msg = check_sales_limit(telegram_id)
            if not _ok:
                from urllib.parse import quote as _q
                return RedirectResponse(url=f"/sales?error={_q(_msg or 'Достигнут лимит продаж по тарифу')}", status_code=302)
        except Exception as _lim_err:
            import logging as _lg
            _lg.error(f"check_sales_limit failed (fail-open): {_lim_err}")

        result = db.add_sale(
            product_id=product_id,
            shop_name=shop_name,
            quantity_sold=quantity,
            user_id=internal_uid,
            sale_price=sale_price,
        )
        if result is None:
            return RedirectResponse(url="/sales?error=Недостаточно+товара+на+складе", status_code=302)

        # Fire post-sale side effects (GSheets, shift alerts, plan milestones)
        # non-blocking: response returns immediately, effects run in background
        try:
            from web.app import _main_loop
            from web.sale_events import post_sale_effects
            if _main_loop is not None:
                asyncio.run_coroutine_threadsafe(
                    post_sale_effects(org_db, result, shop_name, telegram_id),
                    _main_loop,
                )
        except Exception as _pse:
            logging.warning(f"sales_create: post_sale_effects schedule error: {_pse}")

    except Exception as e:
        logging.error(f"sales_create error: {e}")
        return RedirectResponse(url="/sales?error=Ошибка+записи", status_code=302)

    return RedirectResponse(url="/sales?ok=1", status_code=302)


@router.get("/api/sales/{sale_id}")
def api_get_sale(request: Request, sale_id: int):
    """Return JSON with sale data for pre-filling the edit modal."""
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)

    is_admin_role = user.get("role") in ("owner", "admin", "super_admin")
    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        sale = db.get_sale_by_id(sale_id)
        if not sale:
            return JSONResponse({"error": "Not found"}, status_code=404)

        internal_uid = _get_internal_uid(db, telegram_id)
        if is_admin_role:
            allowed_shops = _get_user_allowed_shops(telegram_id, db)
            if sale[2] not in allowed_shops:
                return JSONResponse({"error": "Forbidden"}, status_code=403)
        else:
            # Regular user: can only view/edit their own sales
            if not internal_uid or sale[5] != internal_uid:
                return JSONResponse({"error": "Нет доступа к этой продаже"}, status_code=403)

        return JSONResponse({
            "id": sale[0],
            "product_id": sale[1],
            "shop_name": sale[2],
            "quantity_sold": sale[3],
            "sale_price": float(sale[4] or 0),
            "sale_date": (sale[6] or "")[:10],
            "product_name": sale[7] or "",
        })
    except Exception as e:
        return JSONResponse({"error": "Внутренняя ошибка сервера"}, status_code=500)


def _get_inventory_qty(db, shop_name: str, product_id: int) -> int:
    """Return current inventory qty for a product in a shop, or 0 on error."""
    try:
        conn = db.get_connection()
        try:
            cur = conn.cursor()
            cur.execute(
                "SELECT quantity FROM inventory WHERE shop_name = ? AND product_id = ?",
                (shop_name, product_id),
            )
            row = cur.fetchone()
        finally:
            conn.close()
        return int(row[0]) if row else 0
    except Exception:
        return 0


@router.post("/sales/{sale_id}/edit")
def sales_edit(
    request: Request,
    sale_id: int,
    shop_name: Annotated[str, Form()],
    quantity: Annotated[int, Form()],
    sale_price: Annotated[float, Form()],
    sale_date: str = Form(default=""),
    csrf_token: str = Form(default=""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return JSONResponse({"ok": False, "error": "Unauthorized"}, status_code=401)
    if not verify_csrf_token(request, csrf_token):
        return JSONResponse({"ok": False, "error": "CSRF error"}, status_code=403)

    is_admin_role = user.get("role") in ("owner", "admin", "super_admin")
    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)

        if quantity < 1:
            return JSONResponse({"ok": False, "error": "Неверное количество"}, status_code=400)

        # Load existing sale first
        sale = db.get_sale_by_id(sale_id)
        if not sale:
            return JSONResponse({"ok": False, "error": "Продажа не найдена"}, status_code=404)

        old_shop = sale[2]
        old_qty = sale[3]
        product_id = sale[1]
        internal_uid = _get_internal_uid(db, telegram_id)

        if is_admin_role:
            allowed_shops = _get_user_allowed_shops(telegram_id, db)
            if shop_name not in allowed_shops:
                return JSONResponse({"ok": False, "error": "Магазин недоступен"}, status_code=403)
            if old_shop not in allowed_shops:
                return JSONResponse({"ok": False, "error": "Нет доступа к этой продаже"}, status_code=403)
        else:
            # Regular user: can only edit their own sales, shop stays fixed
            if not internal_uid or sale[5] != internal_uid:
                return JSONResponse({"ok": False, "error": "Нет доступа к этой продаже"}, status_code=403)
            shop_name = old_shop  # users cannot change shop

        # Stock sufficiency check
        shop_changed = (shop_name != old_shop)
        if shop_changed:
            # Moving to new shop — need quantity units available there
            available = _get_inventory_qty(db, shop_name, product_id)
            if available < quantity:
                return JSONResponse(
                    {"ok": False, "error": f"Недостаточно товара в {shop_name}: {available} шт."},
                    status_code=400,
                )
        else:
            # Same shop: only a net increase requires a stock check
            extra = quantity - old_qty
            if extra > 0:
                available = _get_inventory_qty(db, shop_name, product_id)
                if available < extra:
                    return JSONResponse(
                        {"ok": False, "error": f"Недостаточно товара: {available} шт. в наличии"},
                        status_code=400,
                    )

        # Validate optional sale_date (YYYY-MM-DD)
        validated_date = None
        if sale_date and sale_date.strip():
            import re as _re
            if _re.match(r'^\d{4}-\d{2}-\d{2}$', sale_date.strip()):
                validated_date = sale_date.strip()

        internal_uid = _get_internal_uid(db, telegram_id)
        ok = db.update_sale_full(
            sale_id=sale_id,
            quantity_sold=quantity,
            sale_price=sale_price,
            shop_name=shop_name,
            changed_by=internal_uid,
            sale_date=validated_date,
        )
        if not ok:
            return JSONResponse({"ok": False, "error": "Ошибка обновления"}, status_code=500)
        return JSONResponse({"ok": True})
    except Exception as e:
        logging.error(f"sales_edit error: {e}")
        return JSONResponse({"ok": False, "error": "Ошибка сохранения"}, status_code=500)


@router.post("/sales/{sale_id}/delete")
def sales_delete(
    request: Request,
    sale_id: int,
    csrf_token: str = Form(default=""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/sales?error=CSRF+error", status_code=302)

    is_admin_role = user.get("role") in ("owner", "admin", "super_admin")
    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        sale = db.get_sale_by_id(sale_id)
        if not sale:
            return RedirectResponse(url="/sales?error=Продажа+не+найдена", status_code=302)

        internal_uid = _get_internal_uid(db, telegram_id)
        if is_admin_role:
            allowed_shops = _get_user_allowed_shops(telegram_id, db)
            if sale[2] not in allowed_shops:
                return RedirectResponse(url="/sales?error=Нет+доступа", status_code=302)
        else:
            if not internal_uid or sale[5] != internal_uid:
                return RedirectResponse(url="/sales?error=Нет+доступа", status_code=302)

        db.delete_sale(sale_id)
    except Exception as e:
        logging.error(f"sales_delete error: {e}")

    return RedirectResponse(url="/sales", status_code=302)


@router.post("/api/sales/{sale_id}/delete")
def api_sales_delete(
    request: Request,
    sale_id: int,
    csrf_token: str = Form(default=""),
):
    """JSON endpoint for deleting a sale from the edit modal."""
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return JSONResponse({"ok": False, "error": "Unauthorized"}, status_code=401)
    if not verify_csrf_token(request, csrf_token):
        return JSONResponse({"ok": False, "error": "CSRF error"}, status_code=403)

    is_admin_role = user.get("role") in ("owner", "admin", "super_admin")
    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        sale = db.get_sale_by_id(sale_id)
        if not sale:
            return JSONResponse({"ok": False, "error": "Продажа не найдена"}, status_code=404)

        internal_uid = _get_internal_uid(db, telegram_id)
        if is_admin_role:
            allowed_shops = _get_user_allowed_shops(telegram_id, db)
            if sale[2] not in allowed_shops:
                return JSONResponse({"ok": False, "error": "Нет доступа"}, status_code=403)
        else:
            if not internal_uid or sale[5] != internal_uid:
                return JSONResponse({"ok": False, "error": "Нет доступа"}, status_code=403)

        db.delete_sale(sale_id)
        return JSONResponse({"ok": True})
    except Exception as e:
        logging.error(f"api_sales_delete error: {e}")
        return JSONResponse({"ok": False, "error": "Ошибка удаления"}, status_code=500)


@router.get("/api/sales/{sale_id}/audit")
def api_sales_audit(request: Request, sale_id: int):
    """Return JSON audit log for a sale (history of quantity/price changes)."""
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)

    is_admin_role = user.get("role") in ("owner", "admin", "super_admin")
    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        sale = db.get_sale_by_id(sale_id)
        if not sale:
            return JSONResponse({"error": "Not found"}, status_code=404)

        internal_uid = _get_internal_uid(db, telegram_id)
        if is_admin_role:
            allowed_shops = _get_user_allowed_shops(telegram_id, db)
            if sale[2] not in allowed_shops:
                return JSONResponse({"error": "Forbidden"}, status_code=403)
        else:
            if not internal_uid or sale[5] != internal_uid:
                return JSONResponse({"error": "Forbidden"}, status_code=403)

        rows = db.get_sale_audit_log(sale_id) or []
        records = []
        for r in rows:
            records.append({
                "editor": r[2] or "—",
                "old_qty": r[3],
                "new_qty": r[4],
                "old_price": float(r[5]) if r[5] is not None else None,
                "new_price": float(r[6]) if r[6] is not None else None,
                "changed_at": (r[7] or "")[:16].replace("T", " "),
            })
        return JSONResponse({"records": records})
    except Exception as e:
        logging.error(f"api_sales_audit error: {e}")
        return JSONResponse({"error": "Внутренняя ошибка"}, status_code=500)


@router.get("/api/recent-products")
def api_recent_products(request: Request, shop: str = ""):
    """Return recent products for the current user (used in sale modal)."""
    from web.auth import get_session_user
    from web.deps import get_web_db
    user = get_session_user(request)
    if not user:
        return JSONResponse({"error": "Unauthorized", "products": []}, status_code=401)
    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        uid = _get_internal_uid(db, telegram_id)
        if not uid:
            return JSONResponse({"products": []})
        recent_rows = db.get_user_recent_products(uid, limit=10) or []
        # get current stock for each
        # Build inventory index once (not N times)
        inv = db.get_all_inventory(shop_name=shop if shop else None) or []
        stock_map: dict = {}
        for r in inv:
            pid_r, qty_r = r[1], int(r[3] or 0)
            if qty_r > 0:
                stock_map[pid_r] = stock_map.get(pid_r, 0) + qty_r

        result = []
        for row in recent_rows:
            pid, name, price, cat = row[0], row[1], float(row[2] or 0), (row[3] or "")
            stock = stock_map.get(pid, 0)
            if stock > 0:
                result.append({"id": pid, "name": name, "price": price, "category": cat, "stock": stock})
        return JSONResponse({"products": result})
    except Exception as e:
        return JSONResponse({"error": "Внутренняя ошибка сервера", "products": []})


@router.get("/api/favorite-products")
def api_favorite_products(request: Request, shop: str = ""):
    """Return favourite products for the current user (used in sale modal)."""
    from web.auth import get_session_user
    from web.deps import get_web_db
    user = get_session_user(request)
    if not user:
        return JSONResponse({"error": "Unauthorized", "products": []}, status_code=401)
    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        uid = _get_internal_uid(db, telegram_id)
        if not uid:
            return JSONResponse({"products": []})
        fav_ids = set(db.get_favorite_products(uid) or [])
        if not fav_ids:
            return JSONResponse({"products": []})
        inv = db.get_all_inventory(shop_name=shop if shop else None) or []
        products = []
        seen = set()
        for r in inv:
            pid = r[1]
            if pid in fav_ids and int(r[3] or 0) > 0 and pid not in seen:
                seen.add(pid)
                products.append({"id": pid, "name": r[6] or "", "price": float(r[8] or 0),
                                 "category": r[7] or "", "stock": int(r[3] or 0)})
        products.sort(key=lambda p: p["name"].lower())
        return JSONResponse({"products": products})
    except Exception as e:
        return JSONResponse({"error": "Внутренняя ошибка сервера", "products": []})
