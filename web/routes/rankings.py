import io
import logging
import traceback
from fastapi import APIRouter, Request
from fastapi.responses import RedirectResponse, StreamingResponse
from web.response_utils import content_disposition as _cd

logger = logging.getLogger(__name__)

router = APIRouter()

MEDALS = ["🥇", "🥈", "🥉"]


def _get_own_db_uid(db, telegram_id: int):
    """Return internal users.id for this telegram_id, or None."""
    try:
        conn = db.get_connection()
        try:
            cur = conn.cursor()
            cur.execute("SELECT id FROM users WHERE telegram_id = ?", (telegram_id,))
            row = cur.fetchone()
        finally:
            conn.close()
        return row[0] if row else None
    except Exception:
        return None


def _period_dates(period: str, today):
    from datetime import timedelta
    if period == "week":
        return (today - timedelta(days=6)).isoformat(), today.isoformat()
    if period == "prev_month":
        ms = today.replace(day=1)
        end = ms - timedelta(days=1)
        return end.replace(day=1).isoformat(), end.isoformat()
    if period == "all":
        return None, None
    # month
    return today.replace(day=1).isoformat(), today.isoformat()


@router.get("/rankings")
def rankings_page(
    request: Request,
    tab: str = "sellers",
    period: str = "month",
    date_from: str = "",
    date_to: str = "",
    city: str = "",
):
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    from billing_utils import has_module
    if not has_module(telegram_id, "analytics"):
        return RedirectResponse(url="/subscription?msg=analytics_locked", status_code=302)
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    if date_from and date_to:
        period = "custom"

    ctx: dict = {
        "request": request, "user": user,
        "is_admin": user.get("role") in ("owner", "admin", "super_admin"),
        "tab": tab, "period": period,
        "ranking": [], "medals": MEDALS, "error": None,
        "selected_city": city,
        "date_from": date_from, "date_to": date_to,
        "own_rank": None,
        "prev_revenue": None, "growth_pct": None, "cur_revenue": None,
    }

    try:
        db = get_web_db(telegram_id, org_db)

        from timezone_utils import get_current_user_time
        tz = db.get_user_timezone(telegram_id)
        today = get_current_user_time(tz).date()

        if date_from and date_to:
            df, dt = date_from, date_to
        else:
            df, dt = _period_dates(period, today)
        ctx["date_from"] = df or ""
        ctx["date_to"] = dt or ""

        kwargs: dict = {}
        if df:
            kwargs["start_date"] = df
        if dt:
            kwargs["end_date"] = dt

        if tab == "shops":
            raw = db.get_shop_ranking(**kwargs) or []
            # shop_name[0] total_sold[1] total_revenue[2] active_sellers[3] total_sales[4]
            if city:
                cmap = db.get_shop_city_map() or {}
                raw = [r for r in raw if cmap.get(r[0]) == city]
            max_rev = float(raw[0][2]) if raw else 1.0
            ranking = []
            for i, r in enumerate(raw):
                rev = float(r[2] or 0)
                ranking.append({
                    "pos": i + 1, "medal": MEDALS[i] if i < 3 else "",
                    "label": r[0] or "—", "sub": f"{r[3]} продавцов",
                    "qty": int(r[1] or 0), "revenue": rev,
                    "count": int(r[4] or 0), "earnings": float(r[5] or 0),
                    "pct": round(rev / max_rev * 100) if max_rev else 0,
                })
            ctx["ranking"] = ranking

        elif tab == "cities":
            raw = db.get_city_ranking(**kwargs) or []
            # city[0] total_sold[1] total_revenue[2] active_sellers[3] total_sales[4] total_earnings[5]
            max_rev = float(raw[0][2]) if raw else 1.0
            ranking = []
            for i, r in enumerate(raw):
                if not r[0]:
                    continue
                rev = float(r[2] or 0)
                ranking.append({
                    "pos": i + 1, "medal": MEDALS[i] if i < 3 else "",
                    "label": r[0], "sub": f"{r[3]} продавцов",
                    "qty": int(r[1] or 0), "revenue": rev,
                    "count": int(r[4] or 0), "earnings": float(r[5] or 0),
                    "pct": round(rev / max_rev * 100) if max_rev else 0,
                })
            ctx["ranking"] = ranking

        else:  # sellers
            raw = db.get_sales_ranking(**kwargs) or []
            # first_name[0] last_name[1] shop_name[2] total_sold[3] total_revenue[4]
            # total_sales[5] total_earnings[6] user_db_id[7] username[8]
            max_rev = float(raw[0][4]) if raw else 1.0
            ranking = []
            own_uid = _get_own_db_uid(db, telegram_id)
            for i, r in enumerate(raw):
                fname = (r[0] or "").strip()
                lname = (r[1] or "").strip()
                name = f"{fname} {lname}".strip() or (f"@{r[8]}" if r[8] else "—")
                rev = float(r[4] or 0)
                is_me = (own_uid is not None and r[7] == own_uid)
                entry = {
                    "pos": i + 1, "medal": MEDALS[i] if i < 3 else "",
                    "label": name, "sub": r[2] or "—",
                    "username": r[8] or "",
                    "user_db_id": r[7],
                    "qty": int(r[3] or 0), "revenue": rev,
                    "count": int(r[5] or 0), "earnings": float(r[6] or 0),
                    "pct": round(rev / max_rev * 100) if max_rev else 0,
                    "is_me": is_me,
                }
                ranking.append(entry)
                if is_me:
                    ctx["own_rank"] = entry
            ctx["ranking"] = ranking

        # Period comparison vs previous period
        try:
            if df and dt:
                import datetime as _dt_r
                from datetime import timedelta as _tdelta
                cur_start = _dt_r.date.fromisoformat(df)
                cur_end = _dt_r.date.fromisoformat(dt)
                span = (cur_end - cur_start).days
                prev_end = cur_start - _tdelta(days=1)
                prev_start = prev_end - _tdelta(days=span)
                prev_s = db.get_sales_summary(
                    start_date=prev_start.isoformat(), end_date=prev_end.isoformat()
                ) or (0, 0, 0, 0)
                cur_s = db.get_sales_summary(start_date=df, end_date=dt) or (0, 0, 0, 0)
                prev_rev = float(prev_s[2] or 0)
                cur_rev = float(cur_s[2] or 0)
                if prev_rev > 0:
                    gpct = round((cur_rev - prev_rev) / prev_rev * 100, 1)
                elif cur_rev > 0:
                    gpct = 100.0
                else:
                    gpct = 0.0
                ctx["prev_revenue"] = int(prev_rev)
                ctx["cur_revenue"] = int(cur_rev)
                ctx["growth_pct"] = gpct
        except Exception:
            pass

    except Exception as exc:
        ctx["error"] = "Произошла внутренняя ошибка. Попробуйте позже."

    return request.app.state.templates.TemplateResponse(
        request, "rankings/index.html", ctx
    )


@router.get("/rankings/export.xlsx")
def rankings_export_xlsx(request: Request, tab: str = "sellers", period: str = "month", date_from: str = "", date_to: str = ""):
    """Export current ranking tab to Excel."""
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    from billing_utils import has_module
    if not has_module(int(user["sub"]), "analytics"):
        return RedirectResponse(url="/subscription?msg=analytics_locked", status_code=302)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from timezone_utils import get_current_user_time

        telegram_id = int(user["sub"])
        org_db = user.get("org_db")
        if not org_db:
            return RedirectResponse("/dashboard", status_code=302)
        db = get_web_db(telegram_id, org_db)

        tz = db.get_user_timezone(telegram_id)
        today = get_current_user_time(tz).date()
        if date_from and date_to:
            df, dt = date_from, date_to
        else:
            df, dt = _period_dates(period, today)
        kwargs: dict = {}
        if df:
            kwargs["start_date"] = df
        if dt:
            kwargs["end_date"] = dt

        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        hdr_fill = PatternFill("solid", fgColor="1E3A5F")
        hdr_font = Font(bold=True, color="FFFFFF", size=11)
        even_fill = PatternFill("solid", fgColor="F0F4FF")

        wb = openpyxl.Workbook()
        ws = wb.active

        period_labels = {"week": "7 дней", "month": "Месяц", "prev_month": "Прошлый месяц", "all": "Всё время"}
        period_label = period_labels.get(period, period)

        # ── Строка с периодом над таблицей ───────────────────────────────────
        ws.cell(row=1, column=1, value=f"Период: {period_label}").font = Font(bold=True, size=11)
        HDR_ROW = 3

        if tab == "shops":
            raw = db.get_shop_ranking(**kwargs) or []
            ws.title = "Рейтинг магазинов"
            # добавлены Ср.чек и Доля%
            headers = ["#", "Магазин", "Продавцов", "Продано (шт.)", "Транзакций", "Выручка (₽)", "Ср. чек (₽)", "Доля (%)"]
            col_widths = [5, 28, 12, 14, 14, 18, 14, 10]
            total_rev = sum(float(r[2] or 0) for r in raw)
            rows = []
            for i, r in enumerate(raw):
                rev = float(r[2] or 0)
                tx = int(r[4] or 0)
                avg = round(rev / tx, 2) if tx else 0.0
                pct = round(rev / total_rev * 100, 1) if total_rev else 0.0
                rows.append((i + 1, r[0] or "—", int(r[3] or 0), int(r[1] or 0), tx, rev, avg, pct))
        elif tab == "cities":
            raw = db.get_city_ranking(**kwargs) or []
            ws.title = "Рейтинг городов"
            headers = ["#", "Город", "Продавцов", "Продано (шт.)", "Транзакций", "Выручка (₽)", "Ср. чек (₽)", "Доля (%)"]
            col_widths = [5, 24, 12, 14, 14, 18, 14, 10]
            total_rev = sum(float(r[2] or 0) for r in raw if r[0])
            rows = []
            for i, r in enumerate(raw):
                if not r[0]:
                    continue
                rev = float(r[2] or 0)
                tx = int(r[4] or 0)
                avg = round(rev / tx, 2) if tx else 0.0
                pct = round(rev / total_rev * 100, 1) if total_rev else 0.0
                rows.append((i + 1, r[0] or "—", int(r[3] or 0), int(r[1] or 0), tx, rev, avg, pct))
        else:
            raw = db.get_sales_ranking(**kwargs) or []
            ws.title = "Рейтинг продавцов"
            # «З/П» → «Мотивация», добавлены Ср.чек и Доля%
            headers = ["#", "Продавец", "Магазин", "Продано (шт.)", "Транзакций",
                       "Выручка (₽)", "Мотивация (₽)", "Ср. чек (₽)", "Доля (%)"]
            col_widths = [5, 28, 22, 14, 14, 18, 16, 14, 10]
            total_rev = sum(float(r[4] or 0) for r in raw)
            rows = []
            for i, r in enumerate(raw):
                fname = (r[0] or "").strip()
                lname = (r[1] or "").strip()
                name = f"{fname} {lname}".strip() or (f"@{r[8]}" if len(r) > 8 and r[8] else "—")
                rev = float(r[4] or 0)
                tx = int(r[5] or 0)
                avg = round(rev / tx, 2) if tx else 0.0
                pct = round(rev / total_rev * 100, 1) if total_rev else 0.0
                rows.append((i + 1, name, r[2] or "—", int(r[3] or 0), tx, rev,
                              float(r[6] or 0), avg, pct))

        for i, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=HDR_ROW, column=i, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[HDR_ROW].height = 26
        ws.freeze_panes = f"A{HDR_ROW + 1}"

        medals = {1: "🥇", 2: "🥈", 3: "🥉"}
        money_cols_sellers = {6, 7, 8}
        money_cols_other = {6, 7}
        pct_col = len(headers)
        for row_idx, row in enumerate(rows, HDR_ROW + 1):
            row_fill = even_fill if row_idx % 2 == 0 else None
            for col_idx, val in enumerate(row, 1):
                display_val = val
                if col_idx == 1:
                    display_val = f"{medals.get(val, '')} {val}".strip()
                cell = ws.cell(row=row_idx, column=col_idx, value=display_val)
                cell.border = border
                if row_fill:
                    cell.fill = row_fill
                money_cols = money_cols_sellers if tab == "sellers" else money_cols_other
                if col_idx in money_cols:
                    cell.number_format = '#,##0.00 ₽'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx == pct_col:
                    cell.number_format = '0.0"%"'
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in (3, 4, 5):
                    cell.alignment = Alignment(horizontal="center")

        # ── ИТОГО ─────────────────────────────────────────────────────────────
        tot_fill_r = PatternFill("solid", fgColor="DBEAFE")
        tr = len(rows) + HDR_ROW + 1
        for col in range(1, len(headers) + 1):
            ws.cell(row=tr, column=col).border = border
            ws.cell(row=tr, column=col).fill = tot_fill_r
        ws.cell(row=tr, column=1, value="ИТОГО").font = Font(bold=True)
        ws.cell(row=tr, column=4, value=sum(r[3] for r in rows)).font = Font(bold=True)
        ws.cell(row=tr, column=5, value=sum(r[4] for r in rows)).font = Font(bold=True)
        rev_idx = 6
        total_rev_sum = sum(r[rev_idx - 1] for r in rows)
        rev_tot = ws.cell(row=tr, column=rev_idx, value=round(total_rev_sum, 2))
        rev_tot.font = Font(bold=True)
        rev_tot.number_format = '#,##0.00 ₽'
        rev_tot.alignment = Alignment(horizontal="right")

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        tab_names = {"sellers": "sellers", "shops": "shops", "cities": "cities"}
        filename = f"ranking_{tab_names.get(tab, tab)}_{period}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": _cd(filename)},
        )
    except Exception as exc:
        logger.error(f"rankings_export error: {exc}\n{traceback.format_exc()}")
        return RedirectResponse(url=f"/rankings?tab={tab}&period={period}&error=Ошибка+при+экспорте.+Попробуйте+позже.", status_code=302)
