"""
Web-маршруты модуля CRM (клиенты).
GET  /clients                   — список клиентов
GET  /clients/new               — форма добавления
POST /clients/new               — создать клиента
GET  /clients/{id}              — карточка клиента
POST /clients/{id}/edit         — редактировать
POST /clients/{id}/delete       — удалить
GET  /clients/export.xlsx       — экспорт Excel
"""
import io
import json
import logging
import re
from datetime import datetime

from fastapi import APIRouter, Form, Request
from fastapi.responses import RedirectResponse, Response, StreamingResponse

router = APIRouter()
logger = logging.getLogger(__name__)

_PAGE_SIZE = 50


def _get_ctx(request: Request):
    return {
        "user": request.session.get("user", {}),
        "csrf_token": request.session.get("csrf_token", ""),
    }


def _parse_tags(tags_json: str) -> list:
    try:
        v = json.loads(tags_json or "[]")
        return v if isinstance(v, list) else []
    except Exception:
        return []


def _format_phone(phone: str) -> str:
    digits = re.sub(r"\D", "", phone or "")
    if len(digits) == 11 and digits[0] in ("7", "8"):
        return f"+7 ({digits[1:4]}) {digits[4:7]}-{digits[7:9]}-{digits[9:11]}"
    return phone


@router.get("/clients")
def clients_list(request: Request, q: str = "", page: int = 1, tag: str = ""):
    from web.deps import get_web_db
    from web.auth import get_session_user
    from billing_utils import has_module
    user = get_session_user(request)
    if not user:
        return RedirectResponse("/login", status_code=302)
    tg_id = user.get("telegram_id", 0)
    if not has_module(tg_id, "crm"):
        return RedirectResponse("/subscription?msg=crm_locked", status_code=302)

    org_db = user.get("org_db", "")
    db = get_web_db(tg_id, org_db)
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    conn = db.get_connection()
    try:
        offset = (page - 1) * _PAGE_SIZE
        params = []
        where = "1=1"
        if q:
            where += " AND (lower(first_name||' '||last_name) LIKE ? OR phone LIKE ? OR email LIKE ?)"
            like = f"%{q.lower()}%"
            params += [like, like, like]
        if tag:
            where += " AND tags_json LIKE ?"
            params.append(f'%"{tag}"%')

        total = conn.execute(f"SELECT COUNT(*) FROM clients WHERE {where}", params).fetchone()[0]
        rows = conn.execute(
            f"SELECT id, first_name, last_name, phone, email, birth_date, tags_json, created_at "
            f"FROM clients WHERE {where} ORDER BY created_at DESC LIMIT ? OFFSET ?",
            params + [_PAGE_SIZE, offset],
        ).fetchall()

        clients = []
        all_tags = set()
        for r in rows:
            tags = _parse_tags(r[6])
            all_tags.update(tags)
            clients.append({
                "id": r[0], "first_name": r[1], "last_name": r[2],
                "phone": _format_phone(r[3]), "email": r[4],
                "birth_date": r[5], "tags": tags, "created_at": r[7],
            })

        all_tags_list = sorted(all_tags)
        total_pages = max(1, (total + _PAGE_SIZE - 1) // _PAGE_SIZE)

        ctx = _get_ctx(request)
        ctx.update({
            "clients": clients, "total": total, "q": q, "tag": tag,
            "page": page, "total_pages": total_pages,
            "all_tags": all_tags_list, "is_admin": user.get("role") in ("owner", "admin"),
        })
        return request.app.state.templates.TemplateResponse(request, "clients/index.html", ctx)
    finally:
        conn.close()


@router.get("/clients/new")
def clients_new_form(request: Request):
    from web.auth import get_session_user
    from billing_utils import has_module
    user = get_session_user(request)
    if not user:
        return RedirectResponse("/login", status_code=302)
    if not has_module(user.get("telegram_id", 0), "crm"):
        return RedirectResponse("/clients", status_code=302)
    ctx = _get_ctx(request)
    ctx.update({"client": None, "error": request.query_params.get("error", "")})
    return request.app.state.templates.TemplateResponse(request, "clients/form.html", ctx)


@router.post("/clients/new")
def clients_create(
    request: Request,
    csrf_token: str = Form(""),
    first_name: str = Form(""),
    last_name: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    birth_date: str = Form(""),
    source: str = Form(""),
    notes: str = Form(""),
    tags: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module
    user = get_session_user(request)
    if not user:
        return RedirectResponse("/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse("/clients/new?error=csrf", status_code=302)
    if not has_module(user.get("telegram_id", 0), "crm"):
        return RedirectResponse("/clients", status_code=302)
    if not first_name.strip():
        return RedirectResponse("/clients/new?error=name_required", status_code=302)

    tags_list = [t.strip() for t in tags.split(",") if t.strip()] if tags else []
    tags_json = json.dumps(tags_list, ensure_ascii=False)
    birth_val = birth_date.strip() or None

    tg_id = user.get("telegram_id", 0)
    org_db = user.get("org_db", "")
    db = get_web_db(tg_id, org_db)
    conn = db.get_connection()
    try:
        cur = conn.execute(
            "INSERT INTO clients (first_name, last_name, phone, email, birth_date, source, notes, tags_json, created_by) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (first_name.strip(), last_name.strip(), phone.strip(), email.strip(),
             birth_val, source.strip(), notes.strip(), tags_json,
             user.get("telegram_id")),
        )
        conn.commit()
        client_id = cur.lastrowid
        return RedirectResponse(f"/clients/{client_id}", status_code=303)
    finally:
        conn.close()


@router.get("/clients/{client_id}")
def client_detail(request: Request, client_id: int):
    from web.auth import get_session_user
    from web.deps import get_web_db
    from billing_utils import has_module
    user = get_session_user(request)
    if not user:
        return RedirectResponse("/login", status_code=302)
    if not has_module(user.get("telegram_id", 0), "crm"):
        return RedirectResponse("/dashboard", status_code=302)

    tg_id = user.get("telegram_id", 0)
    org_db = user.get("org_db", "")
    db = get_web_db(tg_id, org_db)
    conn = db.get_connection()
    try:
        r = conn.execute(
            "SELECT id, first_name, last_name, phone, email, birth_date, source, notes, tags_json, telegram_id, created_at "
            "FROM clients WHERE id=?", (client_id,)
        ).fetchone()
        if not r:
            return RedirectResponse("/clients", status_code=302)

        client = {
            "id": r[0], "first_name": r[1], "last_name": r[2],
            "phone": r[3], "phone_fmt": _format_phone(r[3]),
            "email": r[4], "birth_date": r[5], "source": r[6],
            "notes": r[7], "tags": _parse_tags(r[8]), "tags_json": r[8],
            "telegram_id": r[9], "created_at": r[10],
        }

        # История продаж
        sales_history = conn.execute(
            "SELECT s.id, p.name, s.quantity_sold, s.sale_price, s.sale_date, s.shop_name "
            "FROM sales s LEFT JOIN products p ON s.product_id=p.id "
            "WHERE s.client_id=? ORDER BY s.sale_date DESC LIMIT 50",
            (client_id,),
        ).fetchall()
        sales = [{"id": r[0], "product": r[1] or "—", "qty": r[2],
                  "price": r[3], "date": r[4], "shop": r[5]} for r in sales_history]
        sales_total = sum((r[2] or 0) * ((r[3]) or 0) for r in sales_history)

        # История записей
        appts_history = conn.execute(
            "SELECT a.id, sv.name, a.start_time, a.end_time, a.status, a.price "
            "FROM appointments a LEFT JOIN services sv ON a.service_id=sv.id "
            "WHERE a.client_id=? ORDER BY a.start_time DESC LIMIT 50",
            (client_id,),
        ).fetchall()
        appts = [{"id": r[0], "service": r[1] or "—", "start": r[2],
                  "end": r[3], "status": r[4], "price": r[5]} for r in appts_history]
        appts_total = sum((r[5] or 0) for r in appts_history)

        ltv = sales_total + appts_total

        ctx = _get_ctx(request)
        ctx.update({
            "client": client, "sales": sales, "sales_total": sales_total,
            "appts": appts, "appts_total": appts_total, "ltv": ltv,
            "is_admin": user.get("role") in ("owner", "admin"),
            "has_services": has_module(user.get("telegram_id", 0), "services"),
        })
        return request.app.state.templates.TemplateResponse(request, "clients/detail.html", ctx)
    finally:
        conn.close()


@router.post("/clients/{client_id}/edit")
def client_edit(
    request: Request,
    client_id: int,
    csrf_token: str = Form(""),
    first_name: str = Form(""),
    last_name: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    birth_date: str = Form(""),
    source: str = Form(""),
    notes: str = Form(""),
    tags: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module
    user = get_session_user(request)
    if not user:
        return RedirectResponse("/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(f"/clients/{client_id}?error=csrf", status_code=302)
    if not has_module(user.get("telegram_id", 0), "crm"):
        return RedirectResponse("/clients", status_code=302)
    if not first_name.strip():
        return RedirectResponse(f"/clients/{client_id}?error=name_required", status_code=302)

    tags_list = [t.strip() for t in tags.split(",") if t.strip()] if tags else []
    tags_json = json.dumps(tags_list, ensure_ascii=False)
    birth_val = birth_date.strip() or None

    tg_id = user.get("telegram_id", 0)
    org_db = user.get("org_db", "")
    db = get_web_db(tg_id, org_db)
    conn = db.get_connection()
    try:
        conn.execute(
            "UPDATE clients SET first_name=?, last_name=?, phone=?, email=?, birth_date=?, "
            "source=?, notes=?, tags_json=?, updated_at=datetime('now') WHERE id=?",
            (first_name.strip(), last_name.strip(), phone.strip(), email.strip(),
             birth_val, source.strip(), notes.strip(), tags_json, client_id),
        )
        conn.commit()
        return RedirectResponse(f"/clients/{client_id}", status_code=303)
    finally:
        conn.close()


@router.post("/clients/{client_id}/delete")
def client_delete(request: Request, client_id: int, csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module
    user = get_session_user(request)
    if not user:
        return RedirectResponse("/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(f"/clients/{client_id}?error=csrf", status_code=302)
    if user.get("role") not in ("owner", "admin"):
        return RedirectResponse(f"/clients/{client_id}", status_code=302)
    if not has_module(user.get("telegram_id", 0), "crm"):
        return RedirectResponse("/clients", status_code=302)

    tg_id = user.get("telegram_id", 0)
    org_db = user.get("org_db", "")
    db = get_web_db(tg_id, org_db)
    conn = db.get_connection()
    try:
        conn.execute("UPDATE sales SET client_id=NULL WHERE client_id=?", (client_id,))
        conn.execute("DELETE FROM clients WHERE id=?", (client_id,))
        conn.commit()
        return RedirectResponse("/clients", status_code=303)
    finally:
        conn.close()


@router.get("/clients/export.xlsx")
def clients_export(request: Request):
    from web.auth import get_session_user
    from web.deps import get_web_db
    from billing_utils import has_module
    user = get_session_user(request)
    if not user:
        return RedirectResponse("/login", status_code=302)
    if not has_module(user.get("telegram_id", 0), "crm"):
        return RedirectResponse("/clients", status_code=302)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return Response("openpyxl not installed", status_code=500)

    tg_id = user.get("telegram_id", 0)
    org_db = user.get("org_db", "")
    db = get_web_db(tg_id, org_db)
    conn = db.get_connection()
    try:
        rows = conn.execute(
            "SELECT first_name, last_name, phone, email, birth_date, source, tags_json, created_at "
            "FROM clients ORDER BY created_at DESC"
        ).fetchall()
    finally:
        conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Клиенты"
    headers = ["Имя", "Фамилия", "Телефон", "Email", "Дата рождения", "Источник", "Теги", "Добавлен"]
    hdr_fill = PatternFill("solid", fgColor="4F46E5")
    hdr_font = Font(color="FFFFFF", bold=True)
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

    for ri, r in enumerate(rows, 2):
        tags = ", ".join(_parse_tags(r[6]))
        ws.append([r[0], r[1], r[2], r[3], r[4] or "", r[5] or "", tags, r[7]])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=clients.xlsx"},
    )
