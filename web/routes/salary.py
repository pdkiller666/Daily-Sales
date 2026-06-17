import io
import logging
import traceback
from typing import Annotated
from fastapi import APIRouter, Request, Form
from fastapi.responses import RedirectResponse, StreamingResponse
from datetime import date

logger = logging.getLogger(__name__)

router = APIRouter()

MONTH_NAMES = {
    1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
    5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
    9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь",
}


def _adjacent_month(year: int, month: int, delta: int):
    """Return (year, month) shifted by delta months."""
    total = (year - 1) * 12 + (month - 1) + delta
    return (total // 12 + 1, total % 12 + 1)


EARNINGS_PAGE_SIZE = 20

# Метки источника мотивации (task #49 — таргетинг по оргструктуре)
MOTIVATION_SOURCE_LABELS = {
    "global": "Общая",
    "trade_network": "Сеть",
    "city": "Город",
    "shop": "Магазин",
    "user": "Сотрудник",
    "schedule": "Месячная",
}


def _source_label(src):
    src = (src or "global")
    return MOTIVATION_SOURCE_LABELS.get(src, src)


# Порядок отображения источников мотивации в сводке (от частного к общему)
_SOURCE_ORDER = ["user", "shop", "city", "trade_network", "schedule", "global"]


def _commission_by_source(earnings):
    """Группирует список начислений по источнику мотивации.

    Принимает список dict-ов с ключами 'source' и 'commission'.
    Возвращает список {'source', 'source_label', 'total'} только для источников
    с ненулевой суммой, упорядоченный от частного (Сотрудник) к общему (Общая).
    """
    totals = {}
    for e in earnings:
        comm = float(e.get("commission") or 0)
        if comm == 0:
            continue
        src = e.get("source") or "global"
        totals[src] = totals.get(src, 0.0) + comm
    ordered = sorted(
        totals.items(),
        key=lambda kv: (_SOURCE_ORDER.index(kv[0]) if kv[0] in _SOURCE_ORDER else len(_SOURCE_ORDER)),
    )
    return [
        {"source": src, "source_label": _source_label(src), "total": round(total, 2)}
        for src, total in ordered
    ]


def _salary_user_earnings(request, user, year: int, month: int, page: int = 1):
    """Personal earnings view for user role."""
    from web.auth import get_csrf_token
    from web.deps import get_web_db
    from datetime import date

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    today = date.today()
    if not year:
        year = today.year
    if not month:
        month = today.month

    prev_y, prev_m = _adjacent_month(year, month, -1)
    next_y, next_m = _adjacent_month(year, month, 1)

    ctx: dict = {
        "request": request, "user": user,
        "is_admin": False,
        "year": year, "month": month,
        "month_name": MONTH_NAMES.get(month, str(month)),
        "prev_y": prev_y, "prev_m": prev_m,
        "next_y": next_y, "next_m": next_m,
        "is_future": (year, month) > (today.year, today.month),
        "csrf_token": get_csrf_token(request),
        "earnings": [],
        "total_commission": 0.0,
        "commission_before_coeff": 0.0,
        "plan_coeff": None,
        "plan_coeff_details": [],
        "total_base": 0.0,
        "total_adj": 0.0,
        "contest_rewards": 0.0,
        "contest_details": [],
        "grand_total": 0.0,
        "worked_days": 0,
        "paid_absence_days": 0,
        "daily_rate": 0.0,
        "adj_rows": [],
        "page": 1,
        "total_pages": 1,
        "total_earnings_count": 0,
        "error": None,
        "user_tz": "Europe/Moscow",
    }

    try:
        db = get_web_db(telegram_id, org_db)
        try:
            ctx["user_tz"] = db.get_user_timezone(telegram_id) or "Europe/Moscow"
        except Exception:
            pass
        import sqlite3 as _sq
        conn_m = _sq.connect("data/main.db")
        uid_row = conn_m.execute(
            "SELECT org_id FROM user_org_mapping WHERE telegram_id=? AND is_active=1",
            (telegram_id,)
        ).fetchone()
        conn_m.close()

        conn_u = db.get_connection()
        try:
            user_row = conn_u.execute(
                "SELECT id FROM users WHERE telegram_id=?", (telegram_id,)
            ).fetchone()
        finally:
            conn_u.close()

        if not user_row:
            ctx["error"] = "Пользователь не найден в базе"
            return request.app.state.templates.TemplateResponse(
                request, "salary/earnings.html", ctx
            )

        user_db_id = user_row[0]
        start_date = f"{year}-{month:02d}-01"
        if month == 12:
            end_date = f"{year}-12-31"
        else:
            import calendar as _cal
            last_day = _cal.monthrange(year, month)[1]
            end_date = f"{year}-{month:02d}-{last_day}"

        # Earnings from commissions
        raw_earnings = db.get_seller_earnings(user_db_id, start_date, end_date) or []
        # cols: commission_amount[0] motivation_type[1] motivation_value[2]
        #       product_name[3] quantity_sold[4] sale_price[5] sale_date[6] shop_name[7]
        earnings = []
        total_commission = 0.0
        for row in raw_earnings:
            comm = float(row[0] or 0)
            total_commission += comm
            mtype = row[1] or "percentage"
            mval = float(row[2] or 0)
            earnings.append({
                "date": str(row[6] or "")[:10],
                "product": row[3] or "—",
                "qty": int(row[4] or 0),
                "price": float(row[5] or 0),
                "mtype": mtype,
                "mval": mval,
                "rate_display": f"{mval:g}%" if mtype == "percentage" else f"{int(mval):,}".replace(",", "\u00a0") + "\u00a0₽/ед.",
                "commission": comm,
                "shop": row[7] or "—",
                "source": (row[8] if len(row) > 8 else "global") or "global",
                "source_label": _source_label(row[8] if len(row) > 8 else "global"),
            })

        # Учитываем joint-бонус (совместный режим мотивации, если настроен)
        commission_raw = total_commission  # сумма из seller_earnings до joint-корр.
        try:
            joint_adj = db.get_joint_bonus_adjustment(user_db_id, start_date, end_date)
            total_commission = round(total_commission + joint_adj, 2)
        except Exception:
            joint_adj = 0.0

        # Коэффициент выполнения недельных планов
        plan_coeff = None
        plan_coeff_details = []
        commission_before_coeff = total_commission
        try:
            ns = db.get_notification_settings(user_db_id)
            if ns.get('plan_coeff_enabled'):
                raw_coeff, plan_coeff_details = db.get_plan_motivation_coefficient(user_db_id, year, month)
                if ns.get('plan_coeff_cap', True):
                    raw_coeff = min(raw_coeff, 1.0)
                # Применяем только если есть планы (details не пустой)
                if plan_coeff_details:
                    plan_coeff = raw_coeff
                    total_commission = round(total_commission * plan_coeff, 2)
        except Exception:
            pass

        # Base salary from schedule × rate (+ paid approved absences)
        worked = db.get_worked_days_count(user_db_id, year, month)
        paid_abs = db.get_paid_absence_days_count(user_db_id, year, month)
        rate = db.get_salary_rate(user_db_id)
        base_salary = (worked + paid_abs) * rate
        adj_sum = db.get_salary_adjustments_sum(user_db_id, year, month)
        adj_rows = db.get_salary_adjustments(user_db_id, year, month) or []

        # Contest prizes for this user in the period
        contest_rewards = 0.0
        contest_details = []
        try:
            contest_rewards, contest_details = db.get_user_contest_rewards_detail(
                telegram_id, start_date, end_date
            )
        except Exception:
            pass

        # Pagination for earnings rows
        total_earnings_count = len(earnings)
        total_pages = max(1, (total_earnings_count + EARNINGS_PAGE_SIZE - 1) // EARNINGS_PAGE_SIZE)
        page = max(1, min(page, total_pages))
        start = (page - 1) * EARNINGS_PAGE_SIZE
        earnings_page = earnings[start:start + EARNINGS_PAGE_SIZE]

        ctx.update({
            "earnings": earnings_page,
            "total_earnings_count": total_earnings_count,
            "commission_by_source": _commission_by_source(earnings),
            "total_commission": round(total_commission, 2),
            "commission_raw": round(commission_raw, 2),
            "joint_adj": round(joint_adj, 2),
            "commission_before_coeff": round(commission_before_coeff, 2),
            "plan_coeff": plan_coeff,
            "plan_coeff_details": plan_coeff_details,
            "total_base": round(base_salary, 2),
            "total_adj": round(adj_sum, 2),
            "contest_rewards": round(contest_rewards, 2),
            "contest_details": contest_details,
            "grand_total": round(base_salary + total_commission + adj_sum + contest_rewards, 2),
            "worked_days": worked,
            "paid_absence_days": paid_abs,
            "daily_rate": rate,
            "adj_rows": adj_rows,
            "page": page,
            "total_pages": total_pages,
        })

    except Exception as exc:
        import logging
        logging.error(f"_salary_user_earnings error: {exc}")
        ctx["error"] = "Произошла внутренняя ошибка. Попробуйте позже."

    return request.app.state.templates.TemplateResponse(
        request, "salary/earnings.html", ctx
    )


@router.get("/salary")
def salary_page(
    request: Request,
    year: int = 0,
    month: int = 0,
    user_id: int = 0,
    page: int = 1,
    detail_page: int = 1,
):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    # Non-admin users see their personal earnings view
    if user.get("role") == "user":
        return _salary_user_earnings(request, user, year, month, page)

    telegram_id = int(user["sub"])
    from billing_utils import has_module
    if not has_module(telegram_id, "team"):
        return RedirectResponse(url="/dashboard?msg=module_team_required", status_code=302)
    org_db = user.get("org_db")

    today = date.today()
    if not year:
        year = today.year
    if not month:
        month = today.month
    year  = max(2015, min(year,  2040))
    month = max(1,    min(month, 12))

    prev_y, prev_m = _adjacent_month(year, month, -1)
    next_y, next_m = _adjacent_month(year, month, 1)
    is_future = (year, month) > (today.year, today.month)

    ctx: dict = {
        "request": request, "user": user,
        "is_admin": user.get("role") in ("owner", "admin", "super_admin"),
        "year": year, "month": month,
        "month_name": MONTH_NAMES.get(month, str(month)),
        "prev_y": prev_y, "prev_m": prev_m,
        "next_y": next_y, "next_m": next_m,
        "is_future": is_future,
        "staff_salary": [], "selected_user_id": user_id,
        "detail_user": None, "work_days_set": set(),
        "adjustments": [], "adj_sum": 0.0,
        "detail_earnings": [], "detail_motivation_total": 0.0,
        "detail_joint_adj": 0.0, "detail_raw_commission": 0.0,
        "detail_plan_coeff": None, "detail_plan_coeff_details": [],
        "detail_contest_rewards": 0.0, "detail_contest_details": [],
        "detail_page": 1, "detail_total_pages": 1, "detail_total_count": 0,
        "total_salary_fund": 0.0, "error": None,
        "csrf_token": get_csrf_token(request),
    }

    try:
        from env_manager import env_manager
        import calendar as _cal
        last_day = _cal.monthrange(year, month)[1]
        start_date = f"{year}-{month:02d}-01"
        end_date = f"{year}-{month:02d}-{last_day}"

        db = get_web_db(telegram_id, org_db)

        # All employees with their rates (super_admin excluded)
        all_rates = db.get_all_salary_rates() or []
        # (user_id[0], first_name[1], last_name[2], daily_rate[3], telegram_id[4])

        staff_salary = []
        total_fund = 0.0

        # Bulk-fetch worked/adj_sum in 2 GROUP BY queries → avoids N+1 for these fields
        bulk = db.get_salary_bulk_stats(year, month, start_date, end_date)

        # Bulk-fetch paid absence days for all users (2 queries instead of N)
        non_admin_uids = [r[0] for r in all_rates if not env_manager.is_super_admin(r[4])]
        paid_abs_bulk: dict = {}
        try:
            paid_abs_bulk = db.get_paid_absence_days_bulk(year, month, non_admin_uids)
        except Exception:
            pass

        # Contest rewards for all users (computed once per contest)
        contest_bulk: dict = {}
        try:
            contest_bulk = db.get_bulk_contest_rewards_by_telegram(start_date, end_date)
        except Exception:
            pass

        for row in all_rates:
            if env_manager.is_super_admin(row[4]):
                continue
            uid = row[0]
            rate = float(row[3] or 0)
            bk = bulk.get(uid, {'worked': 0, 'adj_sum': 0.0, 'motivation': 0.0})
            worked = bk['worked']
            adj_sum = bk['adj_sum']
            paid_abs = paid_abs_bulk.get(uid, 0)
            base = rate * (worked + paid_abs)
            # Correct motivation: includes joint_bonus + plan_coeff
            earn = db.get_seller_total_earnings(uid, start_date=start_date, end_date=end_date) or {}
            motivation = round(float(earn.get('total_earnings', 0.0) or 0), 2)
            # Contest prizes
            contest_rewards = round(contest_bulk.get(row[4], 0.0), 2)
            total = base + adj_sum + motivation + contest_rewards
            total_fund += total

            staff_salary.append({
                "user_id": uid,
                "first_name": row[1] or "",
                "last_name": row[2] or "",
                "telegram_id": row[4],
                "daily_rate": rate,
                "worked_days": worked,
                "paid_absence_days": paid_abs,
                "base_salary": base,
                "adj_sum": adj_sum,
                "motivation": motivation,
                "contest_rewards": contest_rewards,
                "total": total,
                "shop": "",
            })

        # Sort: by total desc
        staff_salary.sort(key=lambda x: x["total"], reverse=True)
        ctx["staff_salary"] = staff_salary
        ctx["total_salary_fund"] = total_fund

        # If a specific user is selected, show their calendar + adjustments + motivation
        if user_id:
            work_days = db.get_work_schedule(user_id, year, month)
            adj_rows = db.get_salary_adjustments(user_id, year, month) or []
            adj_sum_val = db.get_salary_adjustments_sum(user_id, year, month)
            rate_row = next((s for s in staff_salary if s["user_id"] == user_id), None)

            # Build calendar grid: list of weeks, each week = list of (day_num | 0)
            first_weekday, days_in_month = _cal.monthrange(year, month)
            cal_grid: list[list[int]] = []
            week: list[int] = [0] * first_weekday
            for d in range(1, days_in_month + 1):
                week.append(d)
                if len(week) == 7:
                    cal_grid.append(week)
                    week = []
            if week:
                week += [0] * (7 - len(week))
                cal_grid.append(week)

            # Detail motivation earnings (individual raw commissions, for line-by-line breakdown)
            detail_earnings_raw = db.get_seller_earnings(user_id, start_date, end_date) or []
            detail_earnings = []
            detail_raw_commission = 0.0
            for erow in detail_earnings_raw:
                comm = float(erow[0] or 0)
                detail_raw_commission += comm
                mtype = erow[1] or "percentage"
                mval = float(erow[2] or 0)
                detail_earnings.append({
                    "date": str(erow[6] or "")[:10],
                    "product": erow[3] or "—",
                    "qty": int(erow[4] or 0),
                    "price": float(erow[5] or 0),
                    "mtype": mtype,
                    "mval": mval,
                    "rate_display": f"{mval:g}%" if mtype == "percentage" else f"{int(mval):,}".replace(",", "\u00a0") + "\u00a0₽/ед.",
                    "commission": comm,
                    "shop": erow[7] or "—",
                    "source": (erow[8] if len(erow) > 8 else "global") or "global",
                    "source_label": _source_label(erow[8] if len(erow) > 8 else "global"),
                })

            # Adjusted motivation total (with joint_bonus + plan_coeff) for summary line
            detail_motivation_total = rate_row["motivation"] if rate_row else round(detail_raw_commission, 2)
            # Joint bonus adjustment (совместный режим — для отображения в разбивке)
            try:
                detail_joint_adj = round(db.get_joint_bonus_adjustment(user_id, start_date, end_date), 2)
            except Exception:
                detail_joint_adj = 0.0
            # Plan coefficient details for selected user
            detail_plan_coeff = None
            detail_plan_coeff_details = []
            try:
                ns = db.get_notification_settings(user_id)
                if ns.get('plan_coeff_enabled'):
                    raw_c, detail_plan_coeff_details = db.get_plan_motivation_coefficient(user_id, year, month)
                    if ns.get('plan_coeff_cap', True):
                        raw_c = min(raw_c, 1.0)
                    if detail_plan_coeff_details:
                        detail_plan_coeff = raw_c
            except Exception:
                pass

            # Contest rewards for selected user
            detail_contest_rewards = 0.0
            detail_contest_details = []
            try:
                if rate_row:
                    detail_contest_rewards, detail_contest_details = \
                        db.get_user_contest_rewards_detail(rate_row["telegram_id"], start_date, end_date)
            except Exception:
                pass

            # Pagination for detail earnings
            detail_total_count = len(detail_earnings)
            detail_total_pages = max(1, (detail_total_count + EARNINGS_PAGE_SIZE - 1) // EARNINGS_PAGE_SIZE)
            detail_page = max(1, min(detail_page, detail_total_pages))
            d_start = (detail_page - 1) * EARNINGS_PAGE_SIZE
            detail_earnings_page = detail_earnings[d_start:d_start + EARNINGS_PAGE_SIZE]

            ctx["detail_user"] = rate_row
            ctx["work_days_set"] = {int(d[8:10]) for d in work_days}
            ctx["adjustments"] = adj_rows
            ctx["adj_sum"] = adj_sum_val
            ctx["cal_grid"] = cal_grid
            ctx["detail_earnings"] = detail_earnings_page
            ctx["detail_total_count"] = detail_total_count
            ctx["detail_page"] = detail_page
            ctx["detail_total_pages"] = detail_total_pages
            ctx["detail_commission_by_source"] = _commission_by_source(detail_earnings)
            ctx["detail_raw_commission"] = round(detail_raw_commission, 2)
            ctx["detail_joint_adj"] = detail_joint_adj
            ctx["detail_motivation_total"] = round(detail_motivation_total, 2)
            ctx["detail_plan_coeff"] = detail_plan_coeff
            ctx["detail_plan_coeff_details"] = detail_plan_coeff_details
            ctx["detail_contest_rewards"] = round(detail_contest_rewards, 2)
            ctx["detail_contest_details"] = detail_contest_details

    except Exception as exc:
        ctx["error"] = "Произошла внутренняя ошибка. Попробуйте позже."

    return request.app.state.templates.TemplateResponse(
        request, "salary/index.html", ctx
    )


@router.get("/salary/export.xlsx")
def salary_export_xlsx(request: Request, year: int = 0, month: int = 0):
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/salary", status_code=302)

    telegram_id = int(user["sub"])
    from billing_utils import has_extension
    if not has_extension(telegram_id, "salary_export"):
        return RedirectResponse(url="/salary?msg=ext_salary_export_required", status_code=302)
    org_db = user.get("org_db")

    today = date.today()
    if not year:
        year = today.year
    if not month:
        month = today.month
    year  = max(2015, min(year,  2040))
    month = max(1,    min(month, 12))

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from env_manager import env_manager
        import calendar as _cal
        last_day = _cal.monthrange(year, month)[1]
        start_date = f"{year}-{month:02d}-01"
        end_date = f"{year}-{month:02d}-{last_day}"

        db = get_web_db(telegram_id, org_db)

        all_rates = db.get_all_salary_rates() or []
        # (user_id[0], first_name[1], last_name[2], daily_rate[3], telegram_id[4])

        # Contest rewards bulk (computed once per contest)
        xls_contest_bulk: dict = {}
        try:
            xls_contest_bulk = db.get_bulk_contest_rewards_by_telegram(start_date, end_date)
        except Exception:
            pass

        # Bulk paid absence days for Excel export
        xls_non_admin_uids = [r[0] for r in all_rates if not env_manager.is_super_admin(r[4])]
        xls_paid_abs_bulk: dict = {}
        try:
            xls_paid_abs_bulk = db.get_paid_absence_days_bulk(year, month, xls_non_admin_uids)
        except Exception:
            pass

        rows: list = []
        total_fund = 0.0
        motivation_details: dict = {}
        for row in all_rates:
            if env_manager.is_super_admin(row[4]):
                continue
            uid = row[0]
            rate = float(row[3] or 0)
            worked = db.get_worked_days_count(uid, year, month)
            paid_abs = xls_paid_abs_bulk.get(uid, 0)
            adj = db.get_salary_adjustments_sum(uid, year, month)
            earn = db.get_seller_total_earnings(uid, start_date=start_date, end_date=end_date) or {}
            motivation = round(float(earn.get('total_earnings', 0.0) or 0), 2)
            contest_r = round(xls_contest_bulk.get(row[4], 0.0), 2)
            base = rate * (worked + paid_abs)
            total = base + adj + motivation + contest_r
            total_fund += total
            name = f"{row[1] or ''} {row[2] or ''}".strip()
            rows.append((name, rate, worked, paid_abs, base, motivation, adj, contest_r, total, uid))
            try:
                earnings_detail = db.get_seller_earnings(uid, start_date=start_date, end_date=end_date) or []
                if earnings_detail:
                    motivation_details[name] = earnings_detail
            except Exception:
                pass

        rows.sort(key=lambda r: -r[8])

        wb = openpyxl.Workbook()
        ws = wb.active
        mn = MONTH_NAMES.get(month, str(month))
        ws.title = f"Зарплата {mn} {year}"

        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        hdr_fill = PatternFill("solid", fgColor="1E3A5F")
        hdr_font = Font(bold=True, color="FFFFFF", size=11)
        even_fill = PatternFill("solid", fgColor="F0F8FF")
        tot_fill = PatternFill("solid", fgColor="DCFCE7")
        meta_font = Font(italic=True, color="4B5563", size=10)

        # ── Заголовок ведомости ───────────────────────────────────────────────
        ws["A1"] = f"Зарплатная ведомость — {mn} {year}"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"Формула: Оклад = Ставка × (Смен + Оплач. отсутствия)  |  Итого = Оклад + Мотивация + Корректировки + Конкурсы"
        ws["A2"].font = meta_font
        ws.merge_cells("A2:I2")

        # ── 9 колонок: добавлен «Оплач. отсутств.», расшифровано «Корр.» ─────
        headers = ["Сотрудник", "Ставка/день", "Смен (раб.)", "Оплач. отсутств.",
                   "Оклад", "Мотивация", "Корректировки", "Конкурсы", "Итого"]
        col_widths = [28, 14, 12, 16, 16, 14, 16, 14, 16]
        HDR_ROW = 4
        for i, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=HDR_ROW, column=i, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[HDR_ROW].height = 28
        ws.freeze_panes = f"A{HDR_ROW + 1}"

        for row_idx, r in enumerate(rows, HDR_ROW + 1):
            row_fill = even_fill if row_idx % 2 == 0 else None
            for col_idx, val in enumerate(r[:9], 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = border
                if row_fill:
                    cell.fill = row_fill
                if col_idx in (2, 5, 6, 7, 8, 9):
                    cell.number_format = '#,##0.00 ₽'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in (3, 4):
                    cell.alignment = Alignment(horizontal="center")
                if col_idx == 9:
                    cell.font = Font(bold=True, color="166534")

        # ── ИТОГО с разбивкой по компонентам ─────────────────────────────────
        tr = len(rows) + HDR_ROW + 1
        for col in range(1, 10):
            ws.cell(row=tr, column=col).border = border
            ws.cell(row=tr, column=col).fill = tot_fill

        def _bold_money(row, col, val):
            c = ws.cell(row=row, column=col, value=round(val, 2))
            c.font = Font(bold=True)
            c.number_format = '#,##0.00 ₽'
            c.alignment = Alignment(horizontal="right")

        ws.cell(row=tr, column=1, value="ИТОГО").font = Font(bold=True)
        ws.cell(row=tr, column=3, value=sum(r[2] for r in rows)).font = Font(bold=True)
        ws.cell(row=tr, column=4, value=sum(r[3] for r in rows)).font = Font(bold=True)
        _bold_money(tr, 5, sum(r[4] for r in rows))
        _bold_money(tr, 6, sum(r[5] for r in rows))
        _bold_money(tr, 7, sum(r[6] for r in rows))
        _bold_money(tr, 8, sum(r[7] for r in rows))
        _bold_money(tr, 9, total_fund)

        # ── Лист 2: Мотивация (детализация комиссий по продавцам) ────────────
        if motivation_details:
            ws_m = wb.create_sheet("Мотивация")
            ws_m["A1"] = f"Детализация мотивации — {mn} {year}"
            ws_m["A1"].font = Font(bold=True, size=13)
            m_hdr_fill = PatternFill("solid", fgColor="166534")
            m_hdr_font = Font(bold=True, color="FFFFFF", size=10)
            m_headers = ["Сотрудник", "Дата", "Товар", "Магазин", "Кол-во", "Цена (₽)", "Комиссия (₽)", "Источник мотивации"]
            m_widths = [28, 12, 30, 20, 8, 12, 16, 18]
            for i, (h, w) in enumerate(zip(m_headers, m_widths), 1):
                cell = ws_m.cell(row=3, column=i, value=h)
                cell.font = m_hdr_font
                cell.fill = m_hdr_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
                ws_m.column_dimensions[get_column_letter(i)].width = w
            ws_m.freeze_panes = "A4"
            m_row = 4
            m_even = PatternFill("solid", fgColor="F0FFF4")
            source_totals = {}
            for seller_name, detail_rows in sorted(motivation_details.items()):
                for dr in detail_rows:
                    # commission[0] type[1] value[2] product[3] qty[4] price[5] date[6] shop[7] source[8]
                    commission = float(dr[0] or 0)
                    if commission == 0:
                        continue
                    src_key = (dr[8] if len(dr) > 8 else "global") or "global"
                    source_totals[src_key] = source_totals.get(src_key, 0.0) + commission
                    src_label = _source_label(src_key)
                    row_fill = m_even if m_row % 2 == 0 else None
                    for col_idx, val in enumerate(
                        [seller_name, str(dr[6] or "")[:10], dr[3] or "—",
                         dr[7] or "—", int(dr[4] or 0), float(dr[5] or 0), commission,
                         src_label], 1
                    ):
                        cell = ws_m.cell(row=m_row, column=col_idx, value=val)
                        cell.border = border
                        if row_fill:
                            cell.fill = row_fill
                        if col_idx in (6, 7):
                            cell.number_format = '#,##0.00 ₽'
                            cell.alignment = Alignment(horizontal="right")
                        elif col_idx in (2, 5, 8):
                            cell.alignment = Alignment(horizontal="center")
                    m_row += 1
            # Итого по мотивации
            if m_row > 4:
                for col in range(1, 9):
                    ws_m.cell(row=m_row, column=col).border = border
                    ws_m.cell(row=m_row, column=col).fill = PatternFill("solid", fgColor="DCFCE7")
                ws_m.cell(row=m_row, column=1, value="ИТОГО").font = Font(bold=True)
                tot_m = ws_m.cell(row=m_row, column=7,
                                   value=round(sum(r[5] for r in rows), 2))
                tot_m.font = Font(bold=True)
                tot_m.number_format = '#,##0.00 ₽'
                tot_m.alignment = Alignment(horizontal="right")

                # Свод по уровням нацеливания мотивации
                if source_totals:
                    m_row += 2
                    ws_m.cell(row=m_row, column=1, value="Мотивация по уровням нацеливания").font = Font(bold=True, size=11)
                    m_row += 1
                    ordered_src = sorted(
                        source_totals.items(),
                        key=lambda kv: (_SOURCE_ORDER.index(kv[0]) if kv[0] in _SOURCE_ORDER else len(_SOURCE_ORDER)),
                    )
                    for src_key, src_total in ordered_src:
                        lbl_cell = ws_m.cell(row=m_row, column=1, value=_source_label(src_key))
                        lbl_cell.border = border
                        val_cell = ws_m.cell(row=m_row, column=2, value=round(src_total, 2))
                        val_cell.border = border
                        val_cell.number_format = '#,##0.00 ₽'
                        val_cell.alignment = Alignment(horizontal="right")
                        m_row += 1

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f"salary_{year}_{month:02d}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    except Exception as exc:
        logger.error(f"salary_export error: {exc}\n{traceback.format_exc()}")
        return RedirectResponse(url=f"/salary?year={year}&month={month}&error=Ошибка+при+экспорте.+Попробуйте+позже.", status_code=302)


def _get_internal_uid(db, telegram_id: int):
    try:
        conn = db.get_connection()
        cur = conn.cursor()
        cur.execute("SELECT id FROM users WHERE telegram_id = ?", (telegram_id,))
        row = cur.fetchone()
        conn.close()
        return row[0] if row else None
    except Exception:
        return None


@router.post("/salary/adjustment/add")
def salary_adj_add(
    request: Request,
    target_user_id: Annotated[int, Form()],
    amount: Annotated[float, Form()],
    year: Annotated[int, Form()],
    month: Annotated[int, Form()],
    comment: str = Form(default=""),
    csrf_token: str = Form(default=""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/salary", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/salary?error=CSRF", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    year  = max(2015, min(year,  2040))
    month = max(1,    min(month, 12))
    amount = max(-1_000_000.0, min(amount, 1_000_000.0))
    try:
        db = get_web_db(telegram_id, org_db)
        creator_uid = _get_internal_uid(db, telegram_id)

        # Validate target_user_id exists in the org and within the admin's scope
        from web.routes.sales import _get_user_allowed_shops
        from db_utils import get_user_org_scope
        allowed_shops = _get_user_allowed_shops(telegram_id, db)
        all_shops = db.get_all_shops() or []
        if set(allowed_shops) != set(all_shops):
            # Scoped admin — verify target user belongs to allowed shops
            conn = db.get_connection()
            target_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE id = ?", (target_user_id,)
            ).fetchone()
            conn.close()
            if not target_row or target_row[1] not in allowed_shops:
                return RedirectResponse(url="/salary?error=access_denied", status_code=302)

        db.add_salary_adjustment(
            user_id=target_user_id,
            year=year,
            month=month,
            amount=amount,
            comment=(comment or "").strip()[:500] or None,
            created_by=creator_uid,
        )
    except RedirectResponse:
        raise
    except Exception as e:
        logging.error(f"salary_adj_add error: {e}")

    return RedirectResponse(
        url=f"/salary?year={year}&month={month}&user_id={target_user_id}&msg=adj_added",
        status_code=302,
    )


@router.post("/salary/adjustment/{adj_id}/delete")
def salary_adj_delete(
    request: Request,
    adj_id: int,
    year: Annotated[int, Form()],
    month: Annotated[int, Form()],
    target_user_id: Annotated[int, Form()],
    csrf_token: str = Form(default=""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/salary", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/salary?error=CSRF", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    try:
        db = get_web_db(telegram_id, org_db)
        db.delete_salary_adjustment(adj_id, user_id=target_user_id)
    except Exception as e:
        logging.error(f"salary_adj_delete error: {e}")

    return RedirectResponse(
        url=f"/salary?year={year}&month={month}&user_id={target_user_id}&msg=adj_deleted",
        status_code=302,
    )


@router.post("/salary/rate/set")
def salary_rate_set(
    request: Request,
    csrf_token: str = Form(default=""),
    target_user_id: int = Form(...),
    daily_rate: str = Form(default="0"),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/salary", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/salary?error=CSRF", status_code=302)

    try:
        rate = float(daily_rate.replace(",", ".").strip())
        if rate < 0:
            raise ValueError("Ставка не может быть отрицательной")
    except (ValueError, AttributeError) as exc:
        from urllib.parse import quote as _q
        return RedirectResponse(url=f"/salary?error={_q(str(exc))}", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    try:
        db = get_web_db(telegram_id, org_db)
        db.set_salary_rate(target_user_id, rate, updated_by=target_user_id)
        logging.info(f"Salary rate set: user={target_user_id} rate={rate} by={telegram_id}")
    except Exception as exc:
        logging.error(f"salary_rate_set error: {exc}")
        return RedirectResponse(url="/salary?error=Ошибка+сохранения+ставки.+Попробуйте+позже.", status_code=303)

    return RedirectResponse(url=f"/salary?rate_saved=1&user_id={target_user_id}", status_code=303)
