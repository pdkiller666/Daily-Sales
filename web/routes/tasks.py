"""
Web-маршруты модуля задач.
GET  /tasks                              — список задач
GET  /tasks/new                          — форма создания
POST /tasks/new                          — создать задачу
GET  /tasks/{task_id}                    — детальная карточка
POST /tasks/{task_id}/status             — изменить статус
POST /tasks/{task_id}/comment            — добавить комментарий
POST /tasks/{task_id}/checklist/{iid}/toggle — отметить пункт
POST /tasks/{task_id}/edit               — редактировать (admin)
POST /tasks/{task_id}/delete             — удалить (admin)
GET  /tasks/topics                       — управление темами (admin)
POST /tasks/topics/new                   — создать тему
POST /tasks/topics/{tid}/delete          — удалить тему
"""
import html as _html
import json
import logging
import mimetypes
import os
import re
import threading
import urllib.request
import uuid
from datetime import date, datetime
from typing import List

from fastapi import APIRouter, Form, Request, UploadFile, File
from fastapi.responses import FileResponse, RedirectResponse, Response

router = APIRouter()
logger = logging.getLogger(__name__)

MAX_TASK_FILE_SIZE = 20 * 1024 * 1024  # 20 МБ на файл
MAX_TASK_FILES     = 10                 # до 10 файлов на задачу


def _uploads_dir_tasks(org_db: str) -> str:
    base = os.path.splitext(org_db)[0]
    d = base + "_uploads/tasks"
    os.makedirs(d, exist_ok=True)
    return d


def _safe_filename_tasks(original: str) -> str:
    name = os.path.basename(original)
    name = re.sub(r'[^\w.\-]', '_', name)
    return name[:120] or "file"


def _fmt_filesize(size: int) -> str:
    if size < 1024:
        return f"{size} Б"
    if size < 1024 * 1024:
        return f"{size // 1024} КБ"
    return f"{size / 1024 / 1024:.1f} МБ"


PRIORITY_LABELS = {
    'low':    '🟢 Низкий',
    'normal': '🔵 Обычный',
    'high':   '🟡 Высокий',
    'urgent': '🔴 Срочно',
}
PRIORITY_CSS = {
    'low':    'bg-emerald-50 text-emerald-700 border-emerald-200',
    'normal': 'bg-blue-50 text-blue-700 border-blue-200',
    'high':   'bg-amber-50 text-amber-700 border-amber-200',
    'urgent': 'bg-red-50 text-red-700 border-red-200',
}
STATUS_LABELS = {
    'new':         '🆕 Новая',
    'in_progress': '🔄 В работе',
    'review':      '🔍 На проверке',
    'done':        '✅ Выполнена',
    'cancelled':   '🚫 Отменена',
}
STATUS_CSS = {
    'new':         'bg-slate-100 text-slate-600 border-slate-300',
    'in_progress': 'bg-blue-50 text-blue-700 border-blue-200',
    'review':      'bg-amber-50 text-amber-700 border-amber-200',
    'done':        'bg-emerald-50 text-emerald-700 border-emerald-200',
    'cancelled':   'bg-red-50 text-red-500 border-red-200',
}
TOPIC_COLORS = {
    'blue':   'bg-blue-100 text-blue-700 border-blue-200',
    'green':  'bg-emerald-100 text-emerald-700 border-emerald-200',
    'amber':  'bg-amber-100 text-amber-700 border-amber-200',
    'red':    'bg-red-100 text-red-700 border-red-200',
    'purple': 'bg-purple-100 text-purple-700 border-purple-200',
    'slate':  'bg-slate-100 text-slate-600 border-slate-300',
}


def _send_tg_task_notify(telegram_id: int, text: str) -> None:
    """Push-уведомление в Telegram пользователю (fire-and-forget с логированием ошибок)."""
    token = os.environ.get("BOT_TOKEN", "")
    if not token or not telegram_id:
        logger.warning("task notify SKIP: token=%s tg_id=%s", bool(token), telegram_id)
        return
    try:
        url = f"https://api.telegram.org/bot{token}/sendMessage"
        payload = json.dumps({
            "chat_id": telegram_id,
            "text": text,
            "parse_mode": "HTML",
            "reply_markup": {
                "inline_keyboard": [[{"text": "✅ Прочитано", "callback_data": "notif_read"}]]
            },
        }).encode("utf-8")
        req = urllib.request.Request(
            url, data=payload,
            headers={"Content-Type": "application/json"},
        )

        def _do_send():
            try:
                resp = urllib.request.urlopen(req, timeout=10)
                body = resp.read(512)
                logger.info("task notify OK tg_id=%s status=%s body=%s",
                            telegram_id, resp.status, body[:200])
            except Exception as err:
                logger.error("task notify FAIL tg_id=%s: %s", telegram_id, err)

        threading.Thread(target=_do_send, daemon=True).start()
    except Exception as e:
        logger.error("task web notify setup: %s", e)


def _fmt_dt_with_offset(dt, user_tz: str) -> str:
    """Format datetime string with user timezone and UTC offset label (e.g. '26.06.2026 14:30 (UTC+3)')."""
    from timezone_utils import get_user_time as _gut
    try:
        user_dt = _gut(dt, user_tz)
        if user_dt is None:
            return "Неизвестно"
        offset = user_dt.utcoffset()
        total_secs = int(offset.total_seconds())
        hours = total_secs // 3600
        sign = '+' if hours >= 0 else '-'
        return user_dt.strftime('%d.%m.%Y %H:%M') + f" (UTC{sign}{abs(hours)})"
    except Exception:
        return "Неизвестно"


def _get_user_tg_id(db, user_db_id: int) -> int | None:
    """Получить telegram_id сотрудника по его users.id."""
    try:
        conn = db.get_connection()
        try:
            row = conn.execute(
                "SELECT telegram_id FROM users WHERE id = ?", (user_db_id,)
            ).fetchone()
        finally:
            conn.close()
        if row and row[0] is not None:
            return row[0]
        # telegram_id не задан (email-only пользователь или пустой профиль)
        logger.debug("_get_user_tg_id: no telegram_id for uid=%s", user_db_id)
        return None
    except Exception as e:
        logger.error("_get_user_tg_id error uid=%s: %s", user_db_id, e)
        return None


def _notify_recurring_spawn(db, assigned_to, title: str) -> None:
    """Уведомить исполнителя о создании следующей повторяющейся задачи.

    Используется всеми веб-путями закрытия (маршрут статуса, канбан, массовое
    действие), чтобы поведение было одинаковым независимо от способа закрытия.
    """
    if not assigned_to:
        return
    try:
        tg_id = _get_user_tg_id(db, assigned_to)
        if not tg_id:
            return
        try:
            db.add_notification_to_history(
                assigned_to, 'task_assigned',
                f"🔁 Создана следующая задача: {title}")
        except Exception:
            pass
        _send_tg_task_notify(
            tg_id,
            f"🔁 <b>Создана следующая задача</b>\n\n<b>{_html.escape(title)}</b>")
        try:
            from web.push_utils import send_web_push as _swp
            import threading as _th
            _th.Thread(
                target=_swp,
                args=(int(tg_id), "🔁 Создана следующая задача", title, "/tasks"),
                daemon=True,
            ).start()
        except Exception:
            pass
    except Exception as e:
        logger.warning("_notify_recurring_spawn: %s", e)


def _get_staff_list(db) -> list:
    """Список сотрудников org для выбора исполнителя."""
    try:
        rows = db.get_all_users()
        result = []
        for r in rows:
            # SELECT * → id(0) telegram_id(1) first_name(2) last_name(3) middle_name(4)
            #             phone(5) email(6) trade_network(7) shop_name(8) city(9)
            #             timezone(10) created_at(11) username(12)
            uid = r[0]
            name = f"{r[2] or ''} {r[3] or ''}".strip() or r[12] or f"User#{uid}"
            shop = r[8] or ""
            result.append({"id": uid, "name": name, "shop": shop})
        return result
    except Exception:
        return []


def _get_shops_list(db) -> list:
    """Список магазинов из users.shop_name (distinct)."""
    try:
        conn = db.get_connection()
        try:
            rows = conn.execute(
                "SELECT DISTINCT shop_name FROM users "
                "WHERE shop_name IS NOT NULL AND shop_name != '' "
                "ORDER BY shop_name"
            ).fetchall()
        finally:
            conn.close()
        return [{"id": r[0], "name": r[0]} for r in rows]
    except Exception:
        return []


def _build_search_items_json(staff_list: list, shops_list: list) -> str:
    """JSON-массив для динамического поиска получателей задачи."""
    items = []
    for s in staff_list:
        label = s['name'] + (f" ({s['shop']})" if s.get('shop') else '')
        items.append({"type": "user", "id": s['id'], "name": s['name'],
                      "label": label, "key": f"u{s['id']}"})
    for sh in shops_list:
        items.append({"type": "shop", "id": None, "name": sh['name'],
                      "label": f"🏪 {sh['name']}", "key": f"s{sh['name']}"})
    return json.dumps(items, ensure_ascii=False)


def _build_init_selected_json(task: dict | None) -> tuple[str, str]:
    """Возвращает (init_selected_json, init_assign_all) для формы редактирования."""
    if not task:
        return "[]", "false"
    if task.get('assign_all'):
        return "[]", "true"
    if task.get('assigned_to') and task.get('assigned_name'):
        name = task['assigned_name']
        uid = task['assigned_to']
        return json.dumps([{"type": "user", "id": uid, "name": name,
                            "label": name, "key": f"u{uid}"}], ensure_ascii=False), "false"
    if task.get('assigned_shop'):
        shop = task['assigned_shop']
        return json.dumps([{"type": "shop", "id": None, "name": shop,
                            "label": f"🏪 {shop}", "key": f"s{shop}"}], ensure_ascii=False), "false"
    return "[]", "false"


def _get_shop_members_tg_ids(db, shop_name: str) -> list[tuple]:
    """Возвращает [(users.id, telegram_id)] всех активных сотрудников магазина."""
    try:
        conn = db.get_connection()
        try:
            rows = conn.execute(
                "SELECT id, telegram_id FROM users "
                "WHERE shop_name = ? AND telegram_id IS NOT NULL",
                (shop_name,)
            ).fetchall()
        finally:
            conn.close()
        return [(r[0], r[1]) for r in rows]
    except Exception:
        return []


def _get_all_members_tg_ids(db) -> list[tuple]:
    """Возвращает [(users.id, telegram_id)] всех сотрудников орга."""
    try:
        conn = db.get_connection()
        try:
            rows = conn.execute(
                "SELECT id, telegram_id FROM users "
                "WHERE telegram_id IS NOT NULL"
            ).fetchall()
        finally:
            conn.close()
        return [(r[0], r[1]) for r in rows]
    except Exception:
        return []


def _fmt_deadline(d: str | None) -> str:
    if not d:
        return ""
    try:
        s = d[:16].replace("T", " ")
        has_time = len(d) >= 13 and ("T" in d or " " in d[10:])
        dt = datetime.fromisoformat(s) if has_time else date.fromisoformat(d[:10])
        if has_time:
            return dt.strftime("%d.%m.%Y %H:%M")
        return dt.strftime("%d.%m.%Y")
    except Exception:
        return d


def _is_overdue(deadline: str | None, status: str) -> bool:
    if not deadline or status in ('done', 'cancelled'):
        return False
    try:
        has_time = len(deadline) >= 13 and ("T" in deadline or " " in deadline[10:])
        if has_time:
            dl_dt = datetime.fromisoformat(deadline[:16].replace("T", " "))
            return dl_dt < datetime.now()
        dl = date.fromisoformat(deadline[:10])
        return dl < date.today()
    except Exception:
        return False


# ─── LIST ────────────────────────────────────────────────────────────────────

_WEB_STATUS_VALUES = {'', 'active', 'new', 'in_progress', 'review', 'done', 'cancelled'}

_SHARED_PREF_VALUES = {'active', 'done', 'all'}

_WEB_TO_PREF = {
    '': 'all',
    'active': 'active',
    'new': 'active',
    'in_progress': 'active',
    'review': 'active',
    'done': 'done',
    'cancelled': 'done',
}

_PREF_TO_WEB = {
    'active': 'active',
    'done': 'done',
    'all': '',
}


@router.get("/tasks")
def tasks_list(request: Request, status: str = "", topic_id: int = 0,
               assigned_filter: int = 0, shop_filter: str = "", msg: str = "",
               q: str = "", page: int = 1, sort: str = "", group: str = "",
               view_id: int = 0, priority: str = ""):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db
    import math as _math
    import urllib.parse as _urlparse

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    from billing_utils import has_module as _has_module, has_extension as _has_ext
    tasks_pro = _has_module(telegram_id, 'tasks_pro')
    tasks_ai = tasks_pro and _has_ext(telegram_id, 'tasks_ai')

    status_in_url = 'status' in request.query_params

    ctx = {
        "request": request, "user": user, "is_admin": is_admin,
        "tasks": [], "topics": [], "staff_list": [], "shops_list": [],
        "status_filter": status, "topic_filter": topic_id,
        "assigned_filter": assigned_filter, "shop_filter": shop_filter,
        "q": q,
        "sort": sort, "group": group, "active_view_id": view_id,
        "saved_views": [],
        "status_labels": STATUS_LABELS, "status_css": STATUS_CSS,
        "priority_labels": PRIORITY_LABELS, "priority_css": PRIORITY_CSS,
        "topic_colors": TOPIC_COLORS,
        "csrf_token": get_csrf_token(request),
        "msg": msg, "error": None,
        "fmt_deadline": _fmt_deadline, "is_overdue": _is_overdue,
        "tasks_pro": tasks_pro, "tasks_ai": tasks_ai,
        "page": 1, "total_pages": 1, "total_tasks": 0,
        "pagination_base": "/tasks?page=",
    }

    _PAGE_SIZE = 50

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or "") if my_row else ""

        if my_db_id:
            try:
                if not status_in_url:
                    _raw_pref = db.get_user_task_pref(my_db_id, 'status_filter', 'active')
                    _canonical = _raw_pref if _raw_pref in _SHARED_PREF_VALUES else 'active'
                    status = _PREF_TO_WEB.get(_canonical, '')
                    ctx["status_filter"] = status
                else:
                    _pref_val = _WEB_TO_PREF.get(status, 'active')
                    db.set_user_task_pref(my_db_id, 'status_filter', _pref_val)
            except Exception:
                pass

        ctx["topics"] = db.get_task_topics()

        _filter_kwargs = dict(
            status=status or None,
            topic_id=topic_id or None,
            assigned_to=assigned_filter if assigned_filter else None,
            shop_filter=shop_filter or None,
            priority=priority or None,
            is_admin=is_admin,
            my_user_id=my_db_id,
            my_shop=my_shop or None,
            q=q.strip() or None,
        )

        total_tasks = db.count_tasks(**_filter_kwargs)
        total_pages = max(1, _math.ceil(total_tasks / _PAGE_SIZE))
        page = max(1, min(page, total_pages))

        tasks = db.get_tasks(
            **_filter_kwargs,
            sort=sort or None,
            limit=_PAGE_SIZE,
            offset=(page - 1) * _PAGE_SIZE,
        )
        ctx["tasks"] = tasks
        ctx["page"] = page
        ctx["total_pages"] = total_pages
        ctx["total_tasks"] = total_tasks

        _parts = []
        if status: _parts.append(f"status={_urlparse.quote(status)}")
        if priority: _parts.append(f"priority={_urlparse.quote(priority)}")
        if topic_id: _parts.append(f"topic_id={topic_id}")
        if assigned_filter: _parts.append(f"assigned_filter={assigned_filter}")
        if shop_filter: _parts.append(f"shop_filter={_urlparse.quote(shop_filter)}")
        if q: _parts.append(f"q={_urlparse.quote(q)}")
        if sort: _parts.append(f"sort={_urlparse.quote(sort)}")
        if group: _parts.append(f"group={_urlparse.quote(group)}")
        if view_id: _parts.append(f"view_id={view_id}")
        _base = "/tasks?" + ("&".join(_parts) + "&" if _parts else "") + "page="
        ctx["pagination_base"] = _base

        if is_admin:
            ctx["staff_list"] = _get_staff_list(db)
            ctx["shops_list"] = _get_shops_list(db)
        ctx["my_db_id"] = my_db_id
        try:
            ctx["saved_views"] = db.get_task_saved_views(my_db_id) if my_db_id else []
        except Exception:
            ctx["saved_views"] = []
    except Exception as e:
        logger.error("tasks_list: %s", e)
        ctx["error"] = "Ошибка загрузки задач."

    if request.headers.get("HX-Request") == "true":
        return request.app.state.templates.TemplateResponse(
            request, "tasks/_list_fragment.html", ctx
        )

    return request.app.state.templates.TemplateResponse(
        request, "tasks/index.html", ctx
    )


# ─── QUICK-ADD ───────────────────────────────────────────────────────────────

@router.post("/tasks/quick-add")
def tasks_quick_add(
    request: Request,
    csrf_token: str = Form(""),
    title: str = Form(""),
    topic_id: str = Form(""),
    priority: str = Form("normal"),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks?msg=csrf_error", status_code=303)

    title = title.strip()[:200]
    if not title:
        return RedirectResponse(url="/tasks?msg=error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        tid = int(topic_id) if topic_id and topic_id.isdigit() else None
        pri = priority if priority in ("low", "normal", "high", "urgent") else "normal"
        new_tid = db.create_task(title=title, created_by=my_db_id, topic_id=tid, priority=pri)
        try:
            from task_automation import run_rules as _run_rules
            _new_task = db.get_task(new_tid) if new_tid else None
            if _new_task:
                _run_rules(db, 'task_created', dict(_new_task), my_db_id)
        except Exception as _are:
            logger.warning("tasks_quick_add automation: %s", _are)
    except Exception as e:
        logger.error("tasks_quick_add: %s", e)
        return RedirectResponse(url="/tasks?msg=error", status_code=303)
    return RedirectResponse(url="/tasks?msg=created_1", status_code=303)


# ─── SAVED VIEWS ─────────────────────────────────────────────────────────────

@router.post("/tasks/views/save")
def tasks_views_save(
    request: Request,
    csrf_token: str = Form(""),
    view_name: str = Form(""),
    status: str = Form(""),
    topic_id: str = Form(""),
    assigned_filter: str = Form(""),
    shop_filter: str = Form(""),
    q: str = Form(""),
    sort: str = Form(""),
    group: str = Form(""),
):
    import json as _json
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks?msg=csrf_error", status_code=303)

    view_name = view_name.strip()[:80]
    if not view_name:
        return RedirectResponse(url="/tasks?msg=error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        filters = {}
        if status: filters["status"] = status
        if topic_id and topic_id.isdigit(): filters["topic_id"] = int(topic_id)
        if assigned_filter and assigned_filter.isdigit(): filters["assigned_filter"] = int(assigned_filter)
        if shop_filter: filters["shop_filter"] = shop_filter
        if q: filters["q"] = q
        if sort: filters["sort"] = sort
        if group: filters["group"] = group
        db.create_task_saved_view(my_db_id, view_name, _json.dumps(filters, ensure_ascii=False))
    except Exception as e:
        logger.error("tasks_views_save: %s", e)
        return RedirectResponse(url="/tasks?msg=error", status_code=303)
    return RedirectResponse(url="/tasks?msg=view_saved", status_code=303)


@router.post("/tasks/views/{view_id}/delete")
def tasks_views_delete(
    request: Request,
    view_id: int,
    csrf_token: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        db.delete_task_saved_view(view_id, my_db_id)
    except Exception as e:
        logger.error("tasks_views_delete: %s", e)
    return RedirectResponse(url="/tasks?msg=view_deleted", status_code=303)


# ─── INLINE EDIT ─────────────────────────────────────────────────────────────

@router.post("/tasks/{task_id}/inline-edit")
def task_inline_edit(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
    field: str = Form(""),
    value: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return JSONResponse({"error": "not_auth"}, status_code=401)
    if not verify_csrf_token(request, csrf_token):
        return JSONResponse({"error": "csrf"}, status_code=403)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    ALLOWED_FIELDS = {"title", "priority", "deadline"}
    if is_admin:
        ALLOWED_FIELDS.add("assigned_to")
    if field not in ALLOWED_FIELDS:
        return JSONResponse({"error": "invalid_field"}, status_code=400)

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
            task_row = conn.execute(
                "SELECT id, title, priority, deadline, assigned_to, status, created_by FROM tasks WHERE id = ?",
                (task_id,)
            ).fetchone()
        finally:
            conn.close()

        if not task_row:
            return JSONResponse({"error": "not_found"}, status_code=404)

        my_db_id = my_row[0] if my_row else 0
        t_created_by = task_row[6]
        if not is_admin and my_db_id != t_created_by:
            return JSONResponse({"error": "forbidden"}, status_code=403)

        new_db_val = None

        if field == "title":
            cleaned = value.strip()[:200]
            if not cleaned:
                return JSONResponse({"error": "empty_title"}, status_code=400)
            old_val = task_row[1]
            new_db_val = cleaned
            conn2 = db.get_connection()
            try:
                conn2.execute(
                    "UPDATE tasks SET title=?, updated_at=datetime('now') WHERE id=?",
                    (cleaned, task_id)
                )
                conn2.commit()
            finally:
                conn2.close()
            db.add_task_history(task_id, my_db_id, "title_changed", old_val, cleaned)

        elif field == "priority":
            if value not in {"low", "normal", "high", "urgent"}:
                return JSONResponse({"error": "invalid_value"}, status_code=400)
            old_val = task_row[2]
            new_db_val = value
            conn2 = db.get_connection()
            try:
                conn2.execute(
                    "UPDATE tasks SET priority=?, updated_at=datetime('now') WHERE id=?",
                    (value, task_id)
                )
                conn2.commit()
            finally:
                conn2.close()
            db.add_task_history(task_id, my_db_id, "priority_changed", old_val, value)

        elif field == "deadline":
            dl = value.strip() or None
            if dl:
                try:
                    from datetime import date as _date
                    _date.fromisoformat(dl[:10])
                    dl = dl[:10]
                except ValueError:
                    return JSONResponse({"error": "invalid_date"}, status_code=400)
            old_val = task_row[3]
            new_db_val = dl
            conn2 = db.get_connection()
            try:
                conn2.execute(
                    "UPDATE tasks SET deadline=?, updated_at=datetime('now') WHERE id=?",
                    (dl, task_id)
                )
                conn2.commit()
            finally:
                conn2.close()
            db.add_task_history(task_id, my_db_id, "deadline_changed", old_val, dl)

        elif field == "assigned_to" and is_admin:
            try:
                new_uid = int(value) if value and value.strip() != "0" else None
            except ValueError:
                return JSONResponse({"error": "invalid_value"}, status_code=400)
            old_val = str(task_row[4] or "")
            new_db_val = str(new_uid or "")
            conn2 = db.get_connection()
            try:
                conn2.execute(
                    "UPDATE tasks SET assigned_to=?, updated_at=datetime('now') WHERE id=?",
                    (new_uid, task_id)
                )
                conn2.commit()
            finally:
                conn2.close()
            db.add_task_history(task_id, my_db_id, "assignee_changed", old_val, new_db_val)

        return JSONResponse({"ok": True, "field": field, "value": new_db_val})

    except Exception as e:
        logger.error("task_inline_edit %s: %s", task_id, e)
        return JSONResponse({"error": "server_error"}, status_code=500)


# ─── BLOCKERS ────────────────────────────────────────────────────────────────

@router.post("/tasks/{task_id}/add-blocker")
def task_add_blocker(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
    blocker_id: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=forbidden", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        bid = int(blocker_id)
        if bid == task_id:
            raise ValueError("self-reference")
    except (ValueError, TypeError):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)

    try:
        db = get_web_db(telegram_id, org_db)
        db.add_task_blocker(bid, task_id)
    except Exception as e:
        logger.error("task_add_blocker: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)
    return RedirectResponse(url=f"/tasks/{task_id}?msg=blocker_added", status_code=303)


@router.post("/tasks/{task_id}/remove-blocker")
def task_remove_blocker(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
    blocker_id: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=forbidden", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        bid = int(blocker_id)
    except (ValueError, TypeError):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)

    try:
        db = get_web_db(telegram_id, org_db)
        db.remove_task_blocker(bid, task_id)
    except Exception as e:
        logger.error("task_remove_blocker: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)
    return RedirectResponse(url=f"/tasks/{task_id}?msg=blocker_removed", status_code=303)


# ─── NEW FORM ────────────────────────────────────────────────────────────────

def _chat_available(telegram_id: int) -> bool:
    """Доступен ли модуль чата (глобально не отключён + есть в биллинге владельца).

    Зеркалит логику chat_page: иначе задача могла создать тему в отключённом/
    неоплаченном чате, а ссылка «Обсудить в чате» вела на заблокированный модуль.
    """
    try:
        from web.routes.chat import _get_chat_min_plan, _chat_access_ok
        if _get_chat_min_plan() == "Отключён":
            return False
        return _chat_access_ok(telegram_id)
    except Exception:
        return False


# ─── TEMPLATES ───────────────────────────────────────────────────────────────

# ─── POOL (самоназначение) ────────────────────────────────────────────────────

# ─── AI DECOMPOSE ────────────────────────────────────────────────────────────

@router.get("/tasks/decompose")
def tasks_decompose_form(request: Request, msg: str = ""):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module as _has_module, has_extension as _has_ext

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    if not _has_module(telegram_id, 'tasks_pro') or not _has_ext(telegram_id, 'tasks_ai'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    ctx = {
        "request": request, "user": user, "is_admin": True,
        "topics": [], "priority_labels": PRIORITY_LABELS,
        "csrf_token": get_csrf_token(request),
        "msg": msg, "error": None,
    }
    try:
        db = get_web_db(telegram_id, org_db)
        ctx["topics"] = db.get_task_topics()
    except Exception as e:
        logger.error("tasks_decompose_form: %s", e)

    return request.app.state.templates.TemplateResponse(
        request, "tasks/decompose.html", ctx
    )


@router.post("/tasks/decompose/create")
def tasks_decompose_create(
    request: Request,
    csrf_token: str = Form(""),
    tasks_json: str = Form(""),
    topic_id: str = Form(""),
    deadline: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module as _has_module, has_extension as _has_ext
    import json as _json

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/decompose?msg=csrf_error", status_code=303)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    if not _has_module(telegram_id, 'tasks_pro') or not _has_ext(telegram_id, 'tasks_ai'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=303)

    try:
        tasks_data = _json.loads(tasks_json or "[]")
        if not tasks_data:
            return RedirectResponse(url="/tasks/decompose?msg=no_tasks", status_code=303)

        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute("SELECT id FROM users WHERE telegram_id=?", (telegram_id,)).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0

        tid = int(topic_id) if topic_id and topic_id.isdigit() else None
        dl = deadline.strip() or None

        created_count = 0
        for t in tasks_data[:10]:
            if isinstance(t, dict) and t.get('title'):
                db.create_task(
                    title=t['title'][:200],
                    description=t.get('description', '')[:2000],
                    topic_id=tid,
                    created_by=my_db_id,
                    priority=t.get('priority', 'normal') if t.get('priority') in ('normal', 'high', 'urgent', 'low') else 'normal',
                    deadline=dl,
                )
                created_count += 1

        return RedirectResponse(url=f"/tasks?msg=created_{created_count}", status_code=303)
    except Exception as e:
        logger.error("tasks_decompose_create: %s", e)
        return RedirectResponse(url="/tasks/decompose?msg=error", status_code=303)


@router.get("/tasks/pool")
def tasks_pool(request: Request, topic_id: str = "", q: str = "", msg: str = ""):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    ctx = {
        "request": request, "user": user,
        "is_admin": user.get("role") in ("owner", "admin", "super_admin"),
        "tasks": [], "topics": [],
        "priority_labels": PRIORITY_LABELS,
        "csrf_token": get_csrf_token(request),
        "topic_filter": topic_id, "q": q, "msg": msg,
        "error": None,
    }
    try:
        db = get_web_db(telegram_id, org_db)
        ctx["topics"] = db.get_task_topics()
        tid = int(topic_id) if topic_id and topic_id.isdigit() else None
        ctx["tasks"] = db.get_unassigned_tasks(topic_id=tid, q=q.strip() or None)
    except Exception as e:
        logger.error("tasks_pool: %s", e)
        ctx["error"] = "Ошибка загрузки пула задач."

    return request.app.state.templates.TemplateResponse(
        request, "tasks/pool.html", ctx
    )


@router.post("/tasks/{task_id}/self_assign")
def task_self_assign(request: Request, task_id: int, csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/pool?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=303)

    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks/pool?msg=not_found", status_code=303)
        # Only pool tasks (assign_all or shop-assigned, without an assignee) may be self-assigned
        if not (task.get('assign_all') or task.get('assigned_shop')):
            return RedirectResponse(url="/tasks/pool?msg=forbidden", status_code=303)
        if task.get('assigned_to'):
            return RedirectResponse(url="/tasks/pool?msg=already_taken", status_code=303)
        conn = db.get_connection()
        try:
            my_row = conn.execute("SELECT id FROM users WHERE telegram_id=?", (telegram_id,)).fetchone()
        finally:
            conn.close()
        if not my_row:
            return RedirectResponse(url="/tasks/pool?msg=error", status_code=303)
        my_db_id = my_row[0]
        ok = db.self_assign_task(task_id, my_db_id)
        if ok:
            try:
                conn2 = db.get_connection()
                try:
                    _urow = conn2.execute(
                        "SELECT first_name, last_name, username FROM users WHERE id = ?", (my_db_id,)
                    ).fetchone()
                finally:
                    conn2.close()
                _display = (f"{_urow[0] or ''} {_urow[1] or ''}".strip() or _urow[2] or str(my_db_id)) if _urow else str(my_db_id)
                db.add_task_history(task_id, my_db_id, 'assigned', None, f"Взял в работу: {_display}")
            except Exception as _he:
                logger.error("task_self_assign history: %s", _he)
            return RedirectResponse(url=f"/tasks/{task_id}?msg=status_updated", status_code=303)
        else:
            return RedirectResponse(url="/tasks/pool?msg=already_taken", status_code=303)
    except Exception as e:
        logger.error("task_self_assign: %s", e)
        return RedirectResponse(url="/tasks/pool?msg=error", status_code=303)


@router.get("/tasks/templates")
def tasks_templates_list(request: Request, msg: str = ""):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    ctx = {
        "request": request, "user": user, "is_admin": True,
        "templates": [], "topics": [],
        "priority_labels": PRIORITY_LABELS,
        "csrf_token": get_csrf_token(request),
        "msg": msg, "error": None,
    }
    try:
        db = get_web_db(telegram_id, org_db)
        ctx["templates"] = db.get_task_templates()
        ctx["topics"] = db.get_task_topics()
    except Exception as e:
        logger.error("tasks_templates_list: %s", e)
        ctx["error"] = "Ошибка загрузки шаблонов."

    return request.app.state.templates.TemplateResponse(
        request, "tasks/templates.html", ctx
    )


@router.post("/tasks/templates/new")
def tasks_templates_new(request: Request, csrf_token: str = Form(""),
                        title: str = Form(""), description: str = Form(""),
                        priority: str = Form("normal"), topic_id: str = Form(""),
                        checklist_items: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    import json as _json

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/templates?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=303)

    title = title.strip()
    if not title:
        return RedirectResponse(url="/tasks/templates?msg=no_title", status_code=303)

    checklist = []
    for line in (checklist_items or "").splitlines():
        line = line.strip()
        if line:
            checklist.append({"text": line, "done": False})
    checklist_json = _json.dumps(checklist, ensure_ascii=False)

    tid = int(topic_id) if topic_id and topic_id.isdigit() else None

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id=?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else None
        db.create_task_template(
            title=title, description=description.strip(),
            priority=priority, checklist_json=checklist_json,
            topic_id=tid, created_by=my_db_id
        )
        return RedirectResponse(url="/tasks/templates?msg=created", status_code=303)
    except Exception as e:
        logger.error("tasks_templates_new: %s", e)
        return RedirectResponse(url="/tasks/templates?msg=error", status_code=303)


@router.post("/tasks/templates/{tid}/delete")
def tasks_template_delete(request: Request, tid: int, csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/templates?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        db.delete_task_template(tid)
    except Exception as e:
        logger.error("tasks_template_delete: %s", e)

    return RedirectResponse(url="/tasks/templates?msg=deleted", status_code=303)


@router.post("/tasks/{task_id}/save_as_template")
def task_save_as_template(request: Request, task_id: int, csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    import json as _json

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=pro_required", status_code=303)

    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=302)

        # Build checklist from existing checklist items
        checklist_items = []
        conn = db.get_connection()
        try:
            rows = conn.execute(
                "SELECT text FROM task_checklist WHERE task_id=? ORDER BY sort_order, id",
                (task_id,)
            ).fetchall()
            checklist_items = [{"text": r[0], "done": False} for r in rows]
        finally:
            conn.close()

        conn = db.get_connection()
        try:
            my_row = conn.execute("SELECT id FROM users WHERE telegram_id=?", (telegram_id,)).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else None

        db.create_task_template(
            title=task['title'],
            description=task.get('description', ''),
            priority=task.get('priority', 'normal'),
            checklist_json=_json.dumps(checklist_items, ensure_ascii=False),
            topic_id=task.get('topic_id'),
            created_by=my_db_id,
        )
        return RedirectResponse(url=f"/tasks/{task_id}?msg=template_saved", status_code=303)
    except Exception as e:
        logger.error("task_save_as_template: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)


@router.get("/tasks/new")
def tasks_new_form(request: Request, template_id: int = 0, parent_id: int = 0):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module, has_extension as _has_ext
    tasks_pro = _has_module(telegram_id, 'tasks_pro')
    tasks_ai = tasks_pro and _has_ext(telegram_id, 'tasks_ai')

    ctx = {
        "request": request, "user": user, "is_admin": True,
        "topics": [], "staff_list": [], "shops_list": [],
        "priority_labels": PRIORITY_LABELS,
        "csrf_token": get_csrf_token(request),
        "edit_task": None, "error": None,
        "chat_available": _chat_available(telegram_id),
        "search_items_json": "[]", "init_selected_json": "[]", "init_assign_all": "false",
        "tasks_pro": tasks_pro, "tasks_ai": tasks_ai, "templates": [],
        "prefill_title": "", "prefill_description": "",
        "prefill_priority": "normal", "prefill_topic_id": 0,
        "prefill_checklist": "",
        "prefill_parent_task_id": parent_id or 0,
        "prefill_parent_task": None,
        "msg": None,
    }
    try:
        db = get_web_db(telegram_id, org_db)
        ctx["topics"] = db.get_task_topics()
        ctx["staff_list"] = _get_staff_list(db)
        ctx["shops_list"] = _get_shops_list(db)
        ctx["search_items_json"] = _build_search_items_json(ctx["staff_list"], ctx["shops_list"])
        ctx["init_selected_json"] = "[]"
        ctx["init_assign_all"] = "false"
        if parent_id:
            try:
                ctx["prefill_parent_task"] = db.get_task(parent_id)
            except Exception:
                pass
        if tasks_pro:
            ctx["templates"] = db.get_task_templates()
            if template_id:
                tmpl = db.get_task_template(template_id)
                if tmpl:
                    ctx["prefill_title"] = tmpl['title']
                    ctx["prefill_description"] = tmpl['description']
                    ctx["prefill_priority"] = tmpl['priority']
                    ctx["prefill_topic_id"] = tmpl['topic_id'] or 0
                    ctx["prefill_checklist"] = "\n".join(
                        item['text'] for item in tmpl.get('checklist', [])
                    )
    except Exception as e:
        logger.error("tasks_new_form: %s", e)

    return request.app.state.templates.TemplateResponse(
        request, "tasks/form.html", ctx
    )


@router.post("/tasks/new")
async def tasks_new_post(
    request: Request,
    csrf_token: str = Form(""),
    title: str = Form(""),
    description: str = Form(""),
    topic_id: str = Form(""),
    assign_mode: str = Form("person"),
    assigned_to: str = Form(""),
    assigned_shop: str = Form(""),
    priority: str = Form("normal"),
    deadline: str = Form(""),
    recurrence: str = Form("none"),
    create_chat_topic: str = Form(""),
    checklist_items: str = Form(""),
    recipients_json: str = Form(""),
    estimated_hours: str = Form(""),
    parent_task_id: str = Form(""),
    files: List[UploadFile] = File(default=[]),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/new?msg=csrf_error", status_code=303)

    title = title.strip()
    if not title:
        return RedirectResponse(url="/tasks/new?msg=no_title", status_code=303)
    _VALID_ASSIGN_MODES = ('person', 'shop', 'all', 'none', '')
    if assign_mode not in _VALID_ASSIGN_MODES:
        return RedirectResponse(url="/tasks/new?msg=bad_assign", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    if any(f and f.filename for f in files):
        try:
            from web.rate_store import check_rate_limit
            if not check_rate_limit(f"task_upload:{telegram_id}", max_requests=20, window_seconds=60):
                return RedirectResponse(url="/tasks/new?msg=rate_limit", status_code=303)
        except Exception:
            pass

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0

        _topic_id = int(topic_id) if topic_id.isdigit() else None
        _deadline = deadline.strip() or None
        _parent_task_id = int(parent_task_id) if str(parent_task_id).strip().isdigit() else None
        if priority not in PRIORITY_LABELS:
            priority = 'normal'

        # Resolve assignment based on mode
        _assigned_to = None
        _assigned_shop = None
        _assign_all = 0
        if assign_mode == "person":
            _assigned_to = int(assigned_to) if assigned_to.isdigit() else None
        elif assign_mode == "shop":
            _assigned_shop = assigned_shop.strip() or None
        elif assign_mode == "all":
            _assign_all = 1
        elif assign_mode == "multi":
            # Multiple recipients — create one task per recipient, return early
            _recip_list = []
            if recipients_json:
                try:
                    _recip_list = json.loads(recipients_json)
                except Exception:
                    pass
            if _recip_list:
                _items = [s.strip() for s in checklist_items.split("\n") if s.strip()]
                for _rec in _recip_list:
                    _r_to = _rec.get('id') if _rec.get('type') == 'user' else None
                    _r_shop = _rec.get('name') if _rec.get('type') == 'shop' else None
                    _rtid = db.create_task(
                        title=title,
                        description=description.strip(),
                        topic_id=_topic_id,
                        created_by=my_db_id,
                        assigned_to=_r_to,
                        assigned_shop=_r_shop,
                        assign_all=0,
                        priority=priority,
                        deadline=_deadline,
                        linked_chat_topic_id=None,
                        checklist=_items,
                        recurrence=recurrence if recurrence not in ('none', '') else None,
                    )
                    _prio_label = PRIORITY_LABELS.get(priority, "")
                    _dl_str = f"\n📅 Срок: {_fmt_deadline(_deadline)}" if _deadline else ""
                    _multi_tg_text = (
                        f"📋 <b>Вам назначена задача</b>\n\n"
                        f"<b>{_html.escape(title)}</b>\n"
                        f"{_prio_label}{_dl_str}\n\n"
                        f"🌐 Откройте веб-кабинет для подробностей."
                    )
                    _multi_msg = f"📋 Назначена задача: {title}"
                    if _r_to and _rtid:
                        try:
                            db.add_notification_to_history(_r_to, "task_assigned", _multi_msg)
                        except Exception:
                            pass
                        _r_tg = _get_user_tg_id(db, _r_to)
                        if _r_tg:
                            _send_tg_task_notify(_r_tg, _multi_tg_text)
                            try:
                                from web.push_utils import send_web_push
                                send_web_push(_r_tg, "📋 Новая задача", title, "/tasks")
                            except Exception:
                                pass
                    elif _r_shop and _rtid:
                        for _sm_uid, _sm_tg in _get_shop_members_tg_ids(db, _r_shop):
                            try:
                                db.add_notification_to_history(_sm_uid, "task_assigned", _multi_msg)
                            except Exception:
                                pass
                            if _sm_tg:
                                _send_tg_task_notify(_sm_tg, _multi_tg_text)
                                try:
                                    from web.push_utils import send_web_push
                                    send_web_push(_sm_tg, "📋 Новая задача", title, "/tasks")
                                except Exception:
                                    pass
                return RedirectResponse(url="/tasks?msg=created", status_code=303)

        _linked_chat_topic_id = None
        if create_chat_topic == "1" and title and _chat_available(telegram_id):
            try:
                chat_topic_id = db.add_chat_topic(title, my_db_id)
                _linked_chat_topic_id = chat_topic_id
            except Exception as ce:
                logger.warning("tasks: add_chat_topic failed: %s", ce)

        items = [s.strip() for s in checklist_items.split("\n") if s.strip()]

        task_id = db.create_task(
            title=title,
            description=description.strip(),
            topic_id=_topic_id,
            created_by=my_db_id,
            assigned_to=_assigned_to,
            assigned_shop=_assigned_shop,
            assign_all=_assign_all,
            priority=priority,
            deadline=_deadline,
            linked_chat_topic_id=_linked_chat_topic_id,
            checklist=items,
            recurrence=recurrence if recurrence not in ('none', '') else None,
            parent_task_id=_parent_task_id,
        )
        if task_id:
            try:
                db.add_task_history(task_id, my_db_id, 'created', None, title)
            except Exception:
                pass
            # Phase 4.3: estimated_hours
            try:
                _eh = float(estimated_hours.strip().replace(',', '.')) if estimated_hours.strip() else None
                if _eh and _eh > 0:
                    _c2 = db.get_connection()
                    try:
                        _c2.execute("UPDATE tasks SET estimated_hours=? WHERE id=?", (_eh, task_id))
                        _c2.commit()
                    finally:
                        _c2.close()
            except Exception:
                pass

        # Send notifications
        deadline_str = f"\n📅 Срок: {_fmt_deadline(_deadline)}" if _deadline else ""
        prio = PRIORITY_LABELS.get(priority, "")
        notify_text = (
            f"📋 <b>Вам назначена задача</b>\n\n"
            f"<b>{_html.escape(title)}</b>\n"
            f"{prio}{deadline_str}\n\n"
            f"🌐 Откройте веб-кабинет для подробностей."
        )
        notif_msg = f"📋 Назначена задача: {title}"

        def _safe_add_notif(uid: int, ntype: str, msg: str):
            try:
                db.add_notification_to_history(uid, ntype, msg)
            except Exception as _ne:
                logger.warning("tasks: add_notification_to_history uid=%s: %s", uid, _ne)

        if assign_mode == "person" and _assigned_to:
            tg_id = _get_user_tg_id(db, _assigned_to)
            _safe_add_notif(_assigned_to, "task_assigned", notif_msg)
            if tg_id:
                _send_tg_task_notify(tg_id, notify_text)
                try:
                    from web.push_utils import apush
                    await apush(tg_id, "📋 Новая задача", notif_msg, "/tasks")
                except Exception:
                    pass
            # Уведомление создателю в историю (если он не тот же, кто исполнитель)
            if my_db_id and my_db_id != _assigned_to:
                assignee_name = next(
                    (s["name"] for s in _get_staff_list(db) if s["id"] == _assigned_to),
                    f"id={_assigned_to}"
                )
                creator_msg = f"📋 Задача создана: «{title}» → {assignee_name}"
                _safe_add_notif(my_db_id, "task_assigned", creator_msg)
        elif assign_mode == "shop" and _assigned_shop:
            members = _get_shop_members_tg_ids(db, _assigned_shop)
            for uid, tg_id in members:
                _safe_add_notif(uid, "task_assigned", notif_msg)
                if tg_id:
                    _send_tg_task_notify(tg_id, notify_text)
                    try:
                        from web.push_utils import apush
                        await apush(tg_id, "📋 Новая задача", notif_msg, "/tasks")
                    except Exception:
                        pass
            # Notify creator even if they're not in that shop
            if my_db_id and not any(uid == my_db_id for uid, _ in members):
                _safe_add_notif(my_db_id, "task_assigned",
                                f"📋 Задача создана: «{title}» → магазин {_assigned_shop}")
        elif assign_mode == "all":
            members = _get_all_members_tg_ids(db)
            for uid, tg_id in members:
                _safe_add_notif(uid, "task_assigned", notif_msg)
                if tg_id:
                    _send_tg_task_notify(tg_id, notify_text)
                    try:
                        from web.push_utils import apush
                        await apush(tg_id, "📋 Новая задача", notif_msg, "/tasks")
                    except Exception:
                        pass
            # Notify creator if somehow not in members list
            if my_db_id and not any(uid == my_db_id for uid, _ in members):
                _safe_add_notif(my_db_id, "task_assigned",
                                f"📋 Задача создана: «{title}» → вся команда")

        # Automation rules: task created / assigned
        try:
            from task_automation import run_rules as _run_rules
            _new_task = db.get_task(task_id) if task_id else None
            if _new_task:
                _run_rules(db, 'task_created', dict(_new_task), my_db_id)
                if _new_task.get('assigned_to') or _new_task.get('assigned_shop') or _new_task.get('assign_all'):
                    _run_rules(db, 'task_assigned', dict(_new_task), my_db_id)
        except Exception as _are:
            logger.warning("tasks_new_post automation: %s", _are)

        # Вложения при создании задачи
        if files:
            uploads_dir = _uploads_dir_tasks(org_db)
            month_dir = datetime.now().strftime("%Y-%m")
            month_path = os.path.join(uploads_dir, month_dir)
            os.makedirs(month_path, exist_ok=True)
            saved = []
            for f in files:
                if not f or not f.filename:
                    continue
                if len(saved) >= MAX_TASK_FILES:
                    break
                try:
                    raw_data = await f.read()
                    if len(raw_data) == 0 or len(raw_data) > MAX_TASK_FILE_SIZE:
                        continue
                    safe_name = _safe_filename_tasks(f.filename)
                    mime = mimetypes.guess_type(safe_name)[0] or "application/octet-stream"
                    uid = uuid.uuid4().hex[:12]
                    dest = os.path.join(month_path, f"{uid}_{safe_name}")
                    import anyio as _anyio
                    _dest_cap, _data_cap = dest, raw_data
                    await _anyio.to_thread.run_sync(
                        lambda: open(_dest_cap, "wb").write(_data_cap)
                    )
                    saved.append({
                        "file_path": dest, "file_name": f.filename[:255],
                        "file_type": mime, "file_size": len(raw_data),
                        "uploaded_by": my_db_id,
                    })
                except Exception as fe:
                    logger.warning("tasks_new_post: file skip: %s", fe)
            if saved:
                db.add_task_attachments(task_id, my_db_id, saved)

        return RedirectResponse(url=f"/tasks/{task_id}?msg=created", status_code=303)
    except Exception as e:
        logger.error("tasks_new_post: %s", e)
        return RedirectResponse(url="/tasks/new?msg=error", status_code=303)


# ─── TOPICS (must be BEFORE /{task_id} to avoid route shadowing) ─────────────

@router.get("/tasks/kanban")
def tasks_kanban(request: Request, topic_id: int = 0, shop_filter: str = "",
                 assigned_filter: int = 0, group_by: str = "none",
                 wip_in_progress: int = 5, wip_review: int = 5):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    if group_by not in ("none", "assignee", "priority"):
        group_by = "none"

    # WIP-лимиты применяются только к рабочим колонкам
    wip_limits = {
        "in_progress": max(0, wip_in_progress),
        "review": max(0, wip_review),
    }

    ctx = {
        "request": request, "user": user, "is_admin": is_admin,
        "columns": [], "swimlanes": [],
        "topics": [], "staff_list": [], "shops_list": [],
        "topic_filter": topic_id, "shop_filter": shop_filter,
        "assigned_filter": assigned_filter,
        "group_by": group_by, "wip_limits": wip_limits,
        "status_labels": STATUS_LABELS, "status_css": STATUS_CSS,
        "priority_labels": PRIORITY_LABELS, "priority_css": PRIORITY_CSS,
        "topic_colors": TOPIC_COLORS,
        "csrf_token": get_csrf_token(request),
        "fmt_deadline": _fmt_deadline, "is_overdue": _is_overdue,
        "error": None,
    }

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or "") if my_row else ""

        ctx["topics"] = db.get_task_topics()
        if is_admin:
            ctx["staff_list"] = _get_staff_list(db)
            ctx["shops_list"] = _get_shops_list(db)

        all_tasks = db.get_tasks(
            topic_id=topic_id or None,
            assigned_to=assigned_filter if assigned_filter else None,
            shop_filter=shop_filter or None,
            is_admin=is_admin,
            my_user_id=my_db_id,
            my_shop=my_shop or None,
        )
        ctx["my_db_id"] = my_db_id

        col_order = ['new', 'in_progress', 'review', 'done', 'cancelled']
        col_icons = {'new': '🆕', 'in_progress': '🔄', 'review': '🔍',
                     'done': '✅', 'cancelled': '🚫'}
        col_names = {'new': 'Новые', 'in_progress': 'В работе',
                     'review': 'На проверке', 'done': 'Выполнены',
                     'cancelled': 'Отменены'}

        def _build_columns(tasks):
            by_status = {s: [] for s in col_order}
            for t in tasks:
                s = t.get('status', 'new')
                if s in by_status:
                    by_status[s].append(t)
            cols = []
            for s in col_order:
                lim = wip_limits.get(s)
                cnt = len(by_status[s])
                cols.append({
                    "key": s, "icon": col_icons[s], "name": col_names[s],
                    "tasks": by_status[s], "count": cnt,
                    "wip": lim, "over_wip": (lim is not None and lim > 0 and cnt > lim),
                })
            return cols

        # Плоская доска (без swimlane)
        ctx["columns"] = _build_columns(all_tasks)

        # Swimlane-группировка
        if group_by == "assignee":
            groups: dict = {}
            order: list = []
            for t in all_tasks:
                name = t.get('assigned_name') or '— Не назначено'
                key = t.get('assigned_to') or 0
                if key not in groups:
                    groups[key] = {"label": name, "tasks": []}
                    order.append(key)
                groups[key]["tasks"].append(t)
            ctx["swimlanes"] = [
                {"label": groups[k]["label"], "columns": _build_columns(groups[k]["tasks"])}
                for k in order
            ]
        elif group_by == "priority":
            prio_order = ['urgent', 'high', 'normal', 'medium', 'low']
            present = [p for p in prio_order if any(
                (t.get('priority') or 'normal') == p for t in all_tasks)]
            extra = sorted({(t.get('priority') or 'normal') for t in all_tasks}
                           - set(prio_order))
            for p in extra:
                present.append(p)
            for p in present:
                lane_tasks = [t for t in all_tasks
                              if (t.get('priority') or 'normal') == p]
                ctx["swimlanes"].append({
                    "label": PRIORITY_LABELS.get(p, p),
                    "columns": _build_columns(lane_tasks),
                })
    except Exception as e:
        logger.error("tasks_kanban: %s", e)
        ctx["error"] = "Ошибка загрузки канбан-доски."

    return request.app.state.templates.TemplateResponse(
        request, "tasks/kanban.html", ctx
    )


@router.post("/tasks/kanban/move")
def tasks_kanban_move(request: Request,
                      task_id: int = Form(...), status: str = Form(...),
                      csrf_token: str = Form(...)):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from fastapi.responses import JSONResponse

    user = get_session_user(request)
    if not user:
        return JSONResponse({"error": "not_auth"}, status_code=401)
    if not verify_csrf_token(request, csrf_token):
        return JSONResponse({"error": "csrf"}, status_code=403)
    if status not in STATUS_LABELS:
        return JSONResponse({"error": "bad_status"}, status_code=400)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return JSONResponse({"error": "not_found"}, status_code=404)

        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0

        can_edit = (
            is_admin
            or task.get('created_by') == my_db_id
            or task.get('assigned_to') == my_db_id
            or task.get('assign_all')
        )
        if not can_edit:
            return JSONResponse({"error": "forbidden"}, status_code=403)

        if not is_admin and status == 'cancelled':
            return JSONResponse({"error": "forbidden"}, status_code=403)

        old_status_kb = task.get('status', '')
        db.update_task_status(task_id, status)
        try:
            db.add_task_history(task_id, my_db_id, 'status',
                                STATUS_LABELS.get(old_status_kb, old_status_kb),
                                STATUS_LABELS.get(status, status))
        except Exception:
            pass
        try:
            from task_automation import run_rules as _run_rules
            _ev_task = dict(task)
            _ev_task['_old_status'] = old_status_kb
            _ev_task['status'] = status
            _run_rules(db, 'status_changed', _ev_task, my_db_id)
        except Exception as _are:
            logger.warning("tasks_kanban_move automation: %s", _are)
        # Spawn next recurring task when dragged to 'done' (unified + idempotent)
        try:
            from task_automation import spawn_recurring_if_done
            _na = spawn_recurring_if_done(db, task, old_status_kb, status)
            if _na:
                _notify_recurring_spawn(db, _na, task['title'])
        except Exception as _re:
            logger.warning("tasks_kanban_move spawn recurring: %s", _re)
        return JSONResponse({"ok": True, "new_status": status,
                             "label": STATUS_LABELS[status]})
    except Exception as e:
        logger.error("tasks_kanban_move: %s", e)
        return JSONResponse({"error": "server_error"}, status_code=500)


@router.get("/tasks/topics")
def tasks_topics(request: Request, msg: str = ""):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    ctx = {
        "request": request, "user": user, "is_admin": True,
        "topics": [], "topic_colors": TOPIC_COLORS,
        "csrf_token": get_csrf_token(request),
        "msg": msg, "error": None,
        "colors": [("blue","bg-blue-500","Синий"), ("green","bg-emerald-500","Зелёный"),
                   ("amber","bg-amber-500","Жёлтый"), ("red","bg-red-500","Красный"),
                   ("purple","bg-purple-500","Фиолетовый"), ("slate","bg-slate-400","Серый")],
    }
    try:
        db = get_web_db(telegram_id, org_db)
        ctx["topics"] = db.get_task_topics()
    except Exception as e:
        logger.error("tasks_topics: %s", e)

    return request.app.state.templates.TemplateResponse(
        request, "tasks/topics.html", ctx
    )


@router.post("/tasks/topics/new")
def tasks_topics_new(
    request: Request,
    csrf_token: str = Form(""),
    name: str = Form(""),
    color: str = Form("blue"),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks/topics", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/topics?msg=csrf_error", status_code=303)

    name = name.strip()
    if not name:
        return RedirectResponse(url="/tasks/topics?msg=no_name", status_code=303)
    if color not in TOPIC_COLORS:
        color = "blue"

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        db.create_task_topic(name, color, my_db_id)
    except Exception as e:
        logger.error("tasks_topics_new: %s", e)
        return RedirectResponse(url="/tasks/topics?msg=error", status_code=303)

    return RedirectResponse(url="/tasks/topics?msg=created", status_code=303)


@router.post("/tasks/topics/{tid}/edit")
def tasks_topics_edit(
    request: Request,
    tid: int,
    csrf_token: str = Form(""),
    name: str = Form(""),
    color: str = Form("blue"),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks/topics", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/topics?msg=csrf_error", status_code=303)

    name = name.strip()
    if not name:
        return RedirectResponse(url="/tasks/topics?msg=no_name", status_code=303)
    if color not in TOPIC_COLORS:
        color = "blue"

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            _topic = conn.execute("SELECT id FROM task_topics WHERE id=?", (tid,)).fetchone()
        finally:
            conn.close()
        if not _topic:
            return RedirectResponse(url="/tasks/topics?msg=not_found", status_code=303)
        db.update_task_topic(tid, name, color)
    except Exception as e:
        logger.error("tasks_topics_edit: %s", e)
        return RedirectResponse(url="/tasks/topics?msg=error", status_code=303)

    return RedirectResponse(url="/tasks/topics?msg=saved", status_code=303)


@router.post("/tasks/topics/{tid}/delete")
def tasks_topics_delete(
    request: Request,
    tid: int,
    csrf_token: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks/topics", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/topics?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            _topic = conn.execute("SELECT id FROM task_topics WHERE id=?", (tid,)).fetchone()
        finally:
            conn.close()
        if not _topic:
            return RedirectResponse(url="/tasks/topics?msg=not_found", status_code=303)
        db.delete_task_topic(tid)
    except Exception as e:
        logger.error("tasks_topics_delete: %s", e)
        return RedirectResponse(url="/tasks/topics?msg=error", status_code=303)

    return RedirectResponse(url="/tasks/topics?msg=deleted", status_code=303)


# ─── AUTOMATION & SLA (Phase 2, tasks_pro) ──────────────────────────────────

_SLA_PRIORITIES = [
    ("urgent", "🔴 Срочный"), ("high", "🟠 Высокий"),
    ("normal", "🔵 Обычный"), ("low", "🟢 Низкий"),
]
_RULE_EVENTS = [
    ("status_changed", "Смена статуса"),
    ("task_created", "Создание задачи"),
    ("task_assigned", "Назначение исполнителя"),
    ("deadline_approaching", "Приближается дедлайн"),
    ("deadline_passed", "Дедлайн прошёл"),
]
_RULE_ACTIONS = [
    ("notify", "Уведомить"),
    ("set_status", "Сменить статус"),
    ("set_priority", "Сменить приоритет"),
    ("add_comment", "Добавить комментарий"),
]
_RULE_TARGETS = [
    ("assignee", "Исполнителя"), ("creator", "Постановщика"),
    ("admins", "Руководителя"), ("team", "Команду"),
]


def _automation_ctx(request, telegram_id, org_db):
    from web.auth import get_csrf_token
    from web.deps import get_web_db
    ctx = {
        "request": request, "is_admin": True,
        "csrf_token": get_csrf_token(request),
        "rules": [], "sla_policies": {}, "topics": [],
        "sla_priorities": _SLA_PRIORITIES, "rule_events": _RULE_EVENTS,
        "rule_actions": _RULE_ACTIONS, "rule_targets": _RULE_TARGETS,
        "statuses": [("new", "Новая"), ("in_progress", "В работе"),
                     ("review", "На проверке"), ("done", "Выполнена"),
                     ("cancelled", "Отменена")],
        "priorities": _SLA_PRIORITIES,
        "event_labels": dict(_RULE_EVENTS), "action_labels": dict(_RULE_ACTIONS),
    }
    try:
        db = get_web_db(telegram_id, org_db)
        ctx["topics"] = db.get_task_topics()
        ctx["rules"] = db.get_automation_rules()
        ctx["sla_policies"] = {p["priority"]: p for p in db.get_sla_policies()}
    except Exception as e:
        logger.error("_automation_ctx: %s", e)
    return ctx


@router.get("/tasks/automation")
def tasks_automation(request: Request, msg: str = ""):
    from web.auth import get_session_user
    from billing_utils import has_module as _has_module

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    ctx = _automation_ctx(request, telegram_id, org_db)
    ctx["user"] = user
    ctx["msg"] = msg
    return request.app.state.templates.TemplateResponse(
        request, "tasks/automation.html", ctx
    )


@router.post("/tasks/automation/sla")
async def tasks_automation_sla(request: Request):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module as _has_module

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks", status_code=302)

    form = await request.form()
    if not verify_csrf_token(request, form.get("csrf_token", "")):
        return RedirectResponse(url="/tasks/automation?msg=csrf_error", status_code=303)

    def _num(v):
        try:
            f = float(str(v).strip().replace(",", "."))
            return f if f > 0 else None
        except Exception:
            return None

    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        for prio, _ in _SLA_PRIORITIES:
            react = _num(form.get(f"react_{prio}", ""))
            resolve = _num(form.get(f"resolve_{prio}", ""))
            is_active = 1 if form.get(f"active_{prio}") else 0
            db.upsert_sla_policy(prio, react, resolve, is_active)
    except Exception as e:
        logger.error("tasks_automation_sla: %s", e)
        return RedirectResponse(url="/tasks/automation?msg=error", status_code=303)

    return RedirectResponse(url="/tasks/automation?msg=sla_saved", status_code=303)


def _build_rule_payload(form):
    """Build (conditions_json, actions_json) from the rule form."""
    import json as _json
    from task_automation import (
        VALID_EVENTS, VALID_ACTIONS, _PRIORITIES, _STATUSES,
    )
    conditions = {}
    if form.get("cond_topic_id"):
        try:
            conditions["topic_id"] = int(form.get("cond_topic_id"))
        except Exception:
            pass
    if form.get("cond_priority") in _PRIORITIES:
        conditions["priority"] = form.get("cond_priority")
    if form.get("cond_to_status") in _STATUSES:
        conditions["to_status"] = form.get("cond_to_status")

    atype = form.get("action_type", "notify")
    if atype not in VALID_ACTIONS:
        atype = "notify"
    action = {"type": atype}
    if atype == "notify":
        target = form.get("action_target", "assignee")
        action["target"] = target if target in dict(_RULE_TARGETS) else "assignee"
        if (form.get("action_text") or "").strip():
            action["text"] = form.get("action_text").strip()
    elif atype == "set_status":
        action["value"] = form.get("action_status") if form.get("action_status") in _STATUSES else "in_progress"
    elif atype == "set_priority":
        action["value"] = form.get("action_priority") if form.get("action_priority") in _PRIORITIES else "high"
    elif atype == "add_comment":
        action["text"] = (form.get("action_text") or "").strip()
    return _json.dumps(conditions, ensure_ascii=False), _json.dumps([action], ensure_ascii=False)


@router.post("/tasks/automation/rules/new")
async def tasks_automation_rule_new(request: Request):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module as _has_module
    from task_automation import VALID_EVENTS

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks", status_code=302)

    form = await request.form()
    if not verify_csrf_token(request, form.get("csrf_token", "")):
        return RedirectResponse(url="/tasks/automation?msg=csrf_error", status_code=303)

    name = (form.get("name") or "").strip()
    event = form.get("event", "")
    if not name:
        return RedirectResponse(url="/tasks/automation?msg=no_name", status_code=303)
    if event not in VALID_EVENTS:
        return RedirectResponse(url="/tasks/automation?msg=bad_event", status_code=303)

    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        conditions_json, actions_json = _build_rule_payload(form)
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        db.create_automation_rule(name, event, conditions_json, actions_json,
                                  1, my_db_id)
    except Exception as e:
        logger.error("tasks_automation_rule_new: %s", e)
        return RedirectResponse(url="/tasks/automation?msg=error", status_code=303)

    return RedirectResponse(url="/tasks/automation?msg=rule_created", status_code=303)


@router.post("/tasks/automation/rules/{rid}/toggle")
def tasks_automation_rule_toggle(request: Request, rid: int,
                                 csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module as _has_module

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/automation?msg=csrf_error", status_code=303)

    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        rules = {r["id"]: r for r in db.get_automation_rules()}
        r = rules.get(rid)
        if not r:
            return RedirectResponse(url="/tasks/automation?msg=not_found", status_code=303)
        db.update_automation_rule(rid, r["name"], r["event"],
                                  r["conditions_json"], r["actions_json"],
                                  0 if r["is_active"] else 1)
    except Exception as e:
        logger.error("tasks_automation_rule_toggle: %s", e)
        return RedirectResponse(url="/tasks/automation?msg=error", status_code=303)

    return RedirectResponse(url="/tasks/automation?msg=rule_saved", status_code=303)


@router.post("/tasks/automation/rules/{rid}/delete")
def tasks_automation_rule_delete(request: Request, rid: int,
                                 csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module as _has_module

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks/automation?msg=csrf_error", status_code=303)

    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    try:
        db = get_web_db(telegram_id, org_db)
        db.delete_automation_rule(rid)
    except Exception as e:
        logger.error("tasks_automation_rule_delete: %s", e)
        return RedirectResponse(url="/tasks/automation?msg=error", status_code=303)

    return RedirectResponse(url="/tasks/automation?msg=rule_deleted", status_code=303)


# ─── BULK OPERATIONS ─────────────────────────────────────────────────────────

@router.post("/tasks/bulk")
def tasks_bulk(request: Request, action: str = Form(""),
               task_ids: list[int] = Form(default=[]),
               csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks?msg=csrf_error", status_code=303)
    if not task_ids or not action:
        return RedirectResponse(url="/tasks?msg=error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=303)

    STATUS_MAP = {
        "status_new": "new",
        "status_in_progress": "in_progress",
        "status_review": "review",
        "status_done": "done",
        "status_cancelled": "cancelled",
    }

    try:
        db = get_web_db(telegram_id, org_db)

        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else None

        ok = 0
        if action in STATUS_MAP:
            new_status = STATUS_MAP[action]
            for tid in task_ids:
                try:
                    _btask = db.get_task(tid)
                    if not _btask:
                        continue
                    _b_old = dict(_btask).get('status', '')
                    db.update_task_status(tid, new_status)
                    try:
                        db.add_task_history(tid, my_db_id, 'status', None, new_status)
                    except Exception:
                        pass
                    if _b_old != new_status:
                        try:
                            from task_automation import run_rules as _run_rules
                            _ev_task = dict(_btask)
                            _ev_task['_old_status'] = _b_old
                            _ev_task['status'] = new_status
                            _run_rules(db, 'status_changed', _ev_task, my_db_id)
                        except Exception as _are:
                            logger.warning("tasks_bulk automation: %s", _are)
                        # Spawn recurring task on bulk close to 'done' (idempotent)
                        try:
                            from task_automation import spawn_recurring_if_done
                            _na = spawn_recurring_if_done(db, dict(_btask), _b_old, new_status)
                            if _na:
                                _notify_recurring_spawn(db, _na, dict(_btask).get('title', ''))
                        except Exception as _re:
                            logger.warning("tasks_bulk spawn recurring: %s", _re)
                    ok += 1
                except Exception:
                    pass
        elif action == "delete":
            for tid in task_ids:
                try:
                    if not db.get_task(tid):
                        continue
                    try:
                        _atts = db.get_task_attachments(tid)
                        for _att in _atts:
                            _att_id = _att["id"] if isinstance(_att, dict) else _att[0]
                            _ok_a, _fpath = db.delete_task_attachment(_att_id, telegram_id, is_admin=True)
                            if _ok_a and _fpath:
                                try:
                                    os.remove(_fpath)
                                except OSError:
                                    pass
                    except Exception as _ae:
                        logger.warning("tasks_bulk delete attachments: %s", _ae)
                    db.delete_task(tid)
                    ok += 1
                except Exception:
                    pass
        else:
            return RedirectResponse(url="/tasks?msg=error", status_code=303)

        return RedirectResponse(url=f"/tasks?msg=bulk_ok_{ok}", status_code=303)
    except Exception as e:
        logger.error("tasks_bulk: %s", e)
        return RedirectResponse(url="/tasks?msg=error", status_code=303)


# ─── EXCEL EXPORT ────────────────────────────────────────────────────────────

@router.get("/tasks/export")
def tasks_export_excel(request: Request, status: str = "", topic_id: int = 0,
                       assigned_filter: int = 0, shop_filter: str = "", q: str = ""):
    from web.auth import get_session_user
    from web.deps import get_web_db
    from fastapi.responses import StreamingResponse
    import io, datetime as _dt

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url="/tasks", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or "") if my_row else ""

        tasks = db.get_tasks(
            status=status or None,
            topic_id=topic_id or None,
            assigned_to=assigned_filter if assigned_filter else None,
            shop_filter=shop_filter or None,
            is_admin=True,
            my_user_id=my_db_id,
            my_shop=my_shop or None,
            q=q.strip() or None,
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Задачи"

        header_fill = PatternFill("solid", fgColor="4F46E5")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        headers = [
            ("ID", 6), ("Название", 40), ("Статус", 14), ("Приоритет", 12),
            ("Тема", 18), ("Исполнитель", 20), ("Магазин", 16),
            ("Срок", 14), ("Создана", 14), ("Описание", 50),
        ]
        for col, (h, w) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

        ws.row_dimensions[1].height = 22

        status_ru = {"new": "Новая", "in_progress": "В работе", "review": "Проверка",
                     "done": "Выполнена", "cancelled": "Отменена"}
        priority_ru = {"low": "Низкий", "normal": "Обычный", "high": "Высокий", "urgent": "Срочный"}

        for i, t in enumerate(tasks, 2):
            assigned = ""
            if t.get("assign_all"):
                assigned = "Вся команда"
            elif t.get("assigned_shop"):
                assigned = t["assigned_shop"]
            elif t.get("assigned_name"):
                assigned = t["assigned_name"]

            deadline_str = t.get("deadline", "") or ""
            created_str = (t.get("created_at") or "")[:10]

            ws.cell(row=i, column=1, value=t.get("id"))
            ws.cell(row=i, column=2, value=t.get("title", ""))
            ws.cell(row=i, column=3, value=status_ru.get(t.get("status", ""), t.get("status", "")))
            ws.cell(row=i, column=4, value=priority_ru.get(t.get("priority", ""), t.get("priority", "")))
            ws.cell(row=i, column=5, value=t.get("topic_name") or "")
            ws.cell(row=i, column=6, value=assigned)
            ws.cell(row=i, column=7, value=t.get("assigned_shop") or "")
            ws.cell(row=i, column=8, value=deadline_str[:10] if deadline_str else "")
            ws.cell(row=i, column=9, value=created_str)
            ws.cell(row=i, column=10, value=(t.get("description") or "")[:500])

            if i % 2 == 0:
                row_fill = PatternFill("solid", fgColor="F8F7FF")
                for col in range(1, 11):
                    ws.cell(row=i, column=col).fill = row_fill

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        fname = f"tasks_{_dt.date.today().isoformat()}.xlsx"
        from web.response_utils import content_disposition as _cd
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": _cd(fname)},
        )
    except Exception as e:
        logger.error("tasks_export_excel: %s", e)
        return RedirectResponse(url="/tasks?msg=error", status_code=303)


# ─── ANALYTICS ───────────────────────────────────────────────────────────────

@router.get("/tasks/analytics")
def tasks_analytics(request: Request):
    from web.auth import get_session_user
    from web.deps import get_web_db
    import json as _json

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    ctx = {
        "request": request, "user": user, "is_admin": is_admin,
        "stats": {}, "error": None,
        "chart_status": "{}", "chart_status_keys": "[]",
        "chart_topics_labels": [], "chart_topics_data": "[]", "chart_topics_ids": "[]",
        "chart_assignee_labels": [], "chart_assignee_data": "[]", "chart_assignee_ids": "[]",
        "chart_priority_labels": "[]", "chart_priority_data": "[]", "chart_priority_keys": "[]",
        "chart_trend_labels": "[]",
        "chart_created_data": "[]", "chart_completed_data": "[]",
    }

    try:
        db = get_web_db(telegram_id, org_db)
        stats = db.get_tasks_analytics()
        if not is_admin:
            stats['by_assignee'] = []
        ctx["stats"] = stats

        STATUS_RU = {
            'new': 'Новые', 'in_progress': 'В работе',
            'review': 'На проверке', 'done': 'Выполнены', 'cancelled': 'Отменены'
        }
        _by_status = stats.get('by_status', {})
        ctx["chart_status"] = _json.dumps(
            {STATUS_RU.get(k, k): v for k, v in _by_status.items()}
        )
        ctx["chart_status_keys"] = _json.dumps(list(_by_status.keys()))

        by_topic = stats.get('by_topic', [])
        ctx["chart_topics_labels"] = [r[0] for r in by_topic]
        ctx["chart_topics_data"] = _json.dumps([r[1] for r in by_topic])
        ctx["chart_topics_ids"] = _json.dumps([(r[2] if len(r) > 2 else 0) for r in by_topic])

        by_assignee = stats.get('by_assignee', [])
        ctx["chart_assignee_labels"] = [r[0] for r in by_assignee]
        ctx["chart_assignee_data"] = _json.dumps([r[1] for r in by_assignee])
        ctx["chart_assignee_ids"] = _json.dumps([(r[2] if len(r) > 2 else 0) for r in by_assignee])

        PRIORITY_RU = {'urgent': '🔴 Критичный', 'high': '🟠 Высокий',
                       'medium': '🔵 Средний', 'low': '🟢 Низкий'}
        by_priority = stats.get('by_priority', [])
        ctx["chart_priority_labels"] = _json.dumps([PRIORITY_RU.get(r[0], r[0]) for r in by_priority])
        ctx["chart_priority_data"] = _json.dumps([r[1] for r in by_priority])
        ctx["chart_priority_keys"] = _json.dumps([r[0] for r in by_priority])

        # Trend: merge created + completed, fill gaps for last 30 days
        from datetime import date, timedelta
        all_days = [(date.today() - timedelta(days=i)).isoformat() for i in range(29, -1, -1)]
        created_map = dict(stats.get('created_daily', []))
        completed_map = dict(stats.get('completed_daily', []))
        ctx["chart_trend_labels"] = _json.dumps([d[5:] for d in all_days])  # MM-DD
        ctx["chart_created_data"] = _json.dumps([created_map.get(d, 0) for d in all_days])
        ctx["chart_completed_data"] = _json.dumps([completed_map.get(d, 0) for d in all_days])

    except Exception as e:
        logger.error("tasks_analytics: %s", e)
        ctx["error"] = "Ошибка загрузки аналитики."

    return request.app.state.templates.TemplateResponse(
        request, "tasks/analytics.html", ctx
    )


# ─── CALENDAR ────────────────────────────────────────────────────────────────

@router.get("/tasks/calendar")
def tasks_calendar(request: Request, year: int = 0, month: int = 0):
    from web.auth import get_session_user
    from web.deps import get_web_db
    from timezone_utils import get_user_time as _gut, DEFAULT_TZ as _DTZ
    import calendar as _cal

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    ctx = {
        "request": request, "user": user, "is_admin": is_admin,
        "status_labels": STATUS_LABELS, "status_css": STATUS_CSS,
        "priority_labels": PRIORITY_LABELS, "priority_css": PRIORITY_CSS,
        "topic_colors": TOPIC_COLORS,
        "weeks": [], "year": 0, "month": 0, "month_name": "",
        "prev_url": "", "next_url": "", "today_iso": "",
        "weekday_names": ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"],
        "fmt_deadline": _fmt_deadline, "is_overdue": _is_overdue,
        "error": None,
    }

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name, timezone FROM users WHERE telegram_id = ?",
                (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or "") if my_row else ""
        _tz = (my_row[2] or _DTZ) if my_row else _DTZ

        now_local = _gut(datetime.utcnow(), _tz) or datetime.utcnow()
        if not year or not month or month < 1 or month > 12:
            year, month = now_local.year, now_local.month

        from datetime import timedelta as _td
        cal = _cal.Calendar(firstweekday=0)  # 0 = Monday
        grid = cal.monthdatescalendar(year, month)
        grid_start = grid[0][0]
        grid_end = grid[-1][-1]
        # Fetch in UTC по дате; расширяем окно на ±1 день, т.к. перевод deadline
        # в локальную TZ может сдвинуть задачу на соседний день у границ грида.
        date_from = (grid_start - _td(days=1)).isoformat()
        date_to = (grid_end + _td(days=1)).isoformat()

        tasks = db.get_tasks_for_calendar(
            date_from, date_to, is_admin=is_admin,
            my_user_id=my_db_id, my_shop=my_shop or None,
        )

        buckets: dict = {}
        for t in tasks:
            dl = t.get("deadline")
            if not dl:
                continue
            s = str(dl)
            if len(s) >= 13 and (("T" in s) or (" " in s[10:])):
                loc = _gut(s, _tz)
                key = loc.date().isoformat() if loc else s[:10]
            else:
                key = s[:10]
            buckets.setdefault(key, []).append(t)

        today_iso = now_local.date().isoformat()
        weeks = []
        for week in grid:
            row = []
            for d in week:
                iso = d.isoformat()
                row.append({
                    "day": d.day, "iso": iso,
                    "in_month": (d.month == month),
                    "is_today": (iso == today_iso),
                    "tasks": buckets.get(iso, []),
                })
            weeks.append(row)

        _MONTHS_RU = ['', 'Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
                      'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь']
        pm, py = (12, year - 1) if month == 1 else (month - 1, year)
        nm, ny = (1, year + 1) if month == 12 else (month + 1, year)
        ctx.update({
            "weeks": weeks, "year": year, "month": month,
            "month_name": _MONTHS_RU[month],
            "prev_url": f"/tasks/calendar?year={py}&month={pm}",
            "next_url": f"/tasks/calendar?year={ny}&month={nm}",
            "today_iso": today_iso,
        })
    except Exception as e:
        logger.error("tasks_calendar: %s", e)
        ctx["error"] = "Ошибка загрузки календаря."

    return request.app.state.templates.TemplateResponse(
        request, "tasks/calendar.html", ctx
    )


# ─── GANTT / TIMELINE ────────────────────────────────────────────────────────

@router.get("/tasks/gantt")
def tasks_gantt(request: Request):
    from web.auth import get_session_user
    from web.deps import get_web_db
    import json as _json

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)

    ctx = {
        "request": request, "user": user, "is_admin": is_admin,
        "tasks_json": "[]", "deps_json": "[]", "has_data": False,
        "error": None,
    }

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or "") if my_row else ""

        tasks = db.get_tasks(
            is_admin=is_admin, my_user_id=my_db_id, my_shop=my_shop or None,
        )

        items = []
        ids_present = set()
        for t in tasks:
            dl = t.get("deadline")
            if not dl:
                continue  # Гант строится по дедлайну (конец отрезка)
            end_key = str(dl)[:10]
            start_raw = t.get("created_at") or dl
            start_key = str(start_raw)[:10]
            if start_key > end_key:
                start_key = end_key
            items.append({
                "id": t["id"], "title": t["title"],
                "start": start_key, "end": end_key,
                "status": t.get("status", "new"),
                "priority": t.get("priority", "normal"),
                "is_blocked": bool(t.get("is_blocked")),
                "assigned_name": t.get("assigned_name", ""),
            })
            ids_present.add(t["id"])

        deps = [
            {"blocker": b, "blocked": bd}
            for (b, bd) in db.get_all_task_dependencies()
            if b in ids_present and bd in ids_present
        ]
        ctx["tasks_json"] = _json.dumps(items, ensure_ascii=False)
        ctx["deps_json"] = _json.dumps(deps, ensure_ascii=False)
        ctx["has_data"] = bool(items)
    except Exception as e:
        logger.error("tasks_gantt: %s", e)
        ctx["error"] = "Ошибка загрузки таймлайна."

    return request.app.state.templates.TemplateResponse(
        request, "tasks/gantt.html", ctx
    )


# ─── WORKLOAD ────────────────────────────────────────────────────────────────

@router.get("/tasks/workload")
def tasks_workload(request: Request):
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    from billing_utils import has_module as _has_module
    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url="/tasks?msg=pro_required", status_code=302)
    if not is_admin:
        return RedirectResponse(url="/tasks?msg=forbidden", status_code=302)

    ctx = {
        "request": request, "user": user, "is_admin": is_admin,
        "rows": [], "max_estimate": 0, "max_active": 0, "error": None,
    }

    try:
        db = get_web_db(telegram_id, org_db)
        rows = db.get_team_workload()
        for r in rows:
            r["logged_hours"] = round((r.get("logged_minutes") or 0) / 60.0, 1)
        ctx["rows"] = rows
        ctx["max_estimate"] = max([(r["estimate_hours"] or 0) for r in rows] + [0])
        ctx["max_active"] = max([(r["active_count"] or 0) for r in rows] + [0])
    except Exception as e:
        logger.error("tasks_workload: %s", e)
        ctx["error"] = "Ошибка загрузки данных о загрузке команды."

    return request.app.state.templates.TemplateResponse(
        request, "tasks/workload.html", ctx
    )


# ─── DETAIL ──────────────────────────────────────────────────────────────────

@router.get("/tasks/{task_id}")
def task_detail(request: Request, task_id: int, msg: str = ""):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    from billing_utils import has_module as _has_module, has_extension as _has_ext
    from timezone_utils import format_user_datetime as _fmt_user_dt, DEFAULT_TZ as _DEFAULT_TZ
    tasks_pro = _has_module(telegram_id, 'tasks_pro')
    tasks_ai = tasks_pro and _has_ext(telegram_id, 'tasks_ai')

    ctx = {
        "request": request, "user": user, "is_admin": is_admin,
        "task": None, "comments": [], "my_db_id": 0,
        "subtasks": [], "subtask_progress": (0, 0),
        "task_blockers": [], "task_blocking": [], "parent_task": None,
        "status_labels": STATUS_LABELS, "status_css": STATUS_CSS,
        "priority_labels": PRIORITY_LABELS, "priority_css": PRIORITY_CSS,
        "topic_colors": TOPIC_COLORS,
        "csrf_token": get_csrf_token(request),
        "msg": msg, "error": None,
        "fmt_deadline": _fmt_deadline, "is_overdue": _is_overdue,
        "team_completions": [], "team_members_for_task": [], "completed_user_ids": [],
        "my_completion": None,
        "chat_available": _chat_available(telegram_id),
        "task_history": [],
        "my_reminders": [],
        "tasks_pro": tasks_pro, "tasks_ai": tasks_ai,
        "comments_text": "",
        "is_watching": False, "watcher_count": 0,
        "time_total_minutes": 0, "time_logs": [],
        "mention_users": [],
        "fmt_user_dt": lambda dt: _fmt_dt_with_offset(dt, _DEFAULT_TZ),
    }

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name, timezone FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or "") if my_row else ""
        _user_tz = (my_row[2] or _DEFAULT_TZ) if my_row else _DEFAULT_TZ
        ctx["my_db_id"] = my_db_id
        ctx["fmt_filesize"] = _fmt_filesize
        # Перезаписываем fmt_user_dt с реальной таймзоной пользователя (с UTC-меткой)
        ctx["fmt_user_dt"] = lambda dt: _fmt_dt_with_offset(dt, _user_tz)

        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=302)

        if not is_admin:
            allowed = (
                task.get("assigned_to") == my_db_id
                or task.get("created_by") == my_db_id
                or task.get("assign_all")
                or (task.get("assigned_shop") and my_shop and task["assigned_shop"] == my_shop)
            )
            if not allowed:
                return RedirectResponse(url="/tasks", status_code=302)

        ctx["task"] = task
        ctx["my_shop"] = my_shop
        ctx["comments"] = db.get_task_comments(task_id)
        try:
            ctx["attachments"] = db.get_task_attachments(task_id)
        except Exception:
            ctx["attachments"] = []
        try:
            ctx["subtasks"] = db.get_subtasks(task_id)
            ctx["subtask_progress"] = db.get_subtask_progress(task_id)
            ctx["task_blockers"] = db.get_task_blockers(task_id)
            ctx["task_blocking"] = db.get_task_blocking(task_id)
            pt_id = (task or {}).get("parent_task_id")
            ctx["parent_task"] = db.get_task(pt_id) if pt_id else None
        except Exception as _e:
            logger.error("task_detail subtasks/blockers: %s", _e)
            ctx["subtasks"] = []
            ctx["subtask_progress"] = (0, 0)
            ctx["task_blockers"] = []
            ctx["task_blocking"] = []
            ctx["parent_task"] = None

        # Team completion tracking (admin + assign_all/shop tasks)
        _assign_all = task.get('assign_all', False)
        _assigned_shop = task.get('assigned_shop', '')
        if is_admin and (_assign_all or _assigned_shop):
            try:
                completions = db.get_task_user_completions(task_id)
                completed_ids = [c['user_id'] for c in completions]
                all_staff = _get_staff_list(db)
                if _assign_all:
                    task_members = all_staff
                else:
                    task_members = [s for s in all_staff if s.get('shop') == _assigned_shop]
                ctx["team_completions"] = completions
                ctx["completions_map"] = {c['user_id']: c for c in completions}
                ctx["team_members_for_task"] = task_members
                ctx["completed_user_ids"] = completed_ids
            except Exception as _te:
                logger.error("task_detail team_completions: %s", _te)
                ctx["team_completions"] = []
                ctx["completions_map"] = {}
                ctx["team_members_for_task"] = []
                ctx["completed_user_ids"] = []
        else:
            ctx["team_completions"] = []
            ctx["completions_map"] = {}
            ctx["team_members_for_task"] = []
            ctx["completed_user_ids"] = []

        # Для сотрудников командных задач — показываем их коллег и общий прогресс (read-only)
        if not is_admin and (task.get('assign_all') or task.get('assigned_shop')):
            try:
                _assign_all_e = task.get('assign_all', False)
                _assigned_shop_e = task.get('assigned_shop', '')
                _completions_e = db.get_task_user_completions(task_id)
                _all_staff_e = _get_staff_list(db)
                _task_members_e = _all_staff_e if _assign_all_e else [
                    s for s in _all_staff_e if s.get('shop') == _assigned_shop_e
                ]
                ctx["team_members_for_task"] = _task_members_e
                ctx["completed_user_ids"] = [c['user_id'] for c in _completions_e]
                ctx["completions_map"] = {c['user_id']: c for c in _completions_e}
            except Exception:
                pass

        # My personal completion for team tasks (for employees)
        if not is_admin and (_assign_all or _assigned_shop):
            try:
                ctx["my_completion"] = db.get_task_user_completion(task_id, my_db_id)
            except Exception:
                ctx["my_completion"] = None
        else:
            ctx["my_completion"] = None

        # Task history (admin only or task creator)
        can_see_history = is_admin or (task.get('created_by') == my_db_id)
        if can_see_history:
            try:
                ctx["task_history"] = db.get_task_history(task_id, limit=30)
            except Exception:
                ctx["task_history"] = []
        else:
            ctx["task_history"] = []

        # My active reminders for this task
        if my_db_id and task.get('status') not in ('done', 'cancelled'):
            try:
                ctx["my_reminders"] = db.get_task_reminders_for_user(task_id, my_db_id)
            except Exception:
                ctx["my_reminders"] = []

        # Phase 4.2: Task watchers
        try:
            ctx["is_watching"] = db.is_watching_task(task_id, my_db_id) if my_db_id else False
            watcher_ids = db.get_task_watcher_db_ids(task_id)
            ctx["watcher_count"] = len(watcher_ids)
        except Exception:
            ctx["is_watching"] = False
            ctx["watcher_count"] = 0

        # Phase 4.3: Time tracking (tasks_pro gate)
        if tasks_pro:
            try:
                ctx["time_total_minutes"] = db.get_task_time_total(task_id)
                ctx["time_logs"] = db.get_task_time_logs(task_id)
            except Exception:
                ctx["time_total_minutes"] = 0
                ctx["time_logs"] = []
        else:
            ctx["time_total_minutes"] = 0
            ctx["time_logs"] = []

        # Build org user list for @mention autocomplete (4.1)
        # Admin: full list (up to 100); non-admin: only users related to this task
        try:
            _mention_conn = db.get_connection()
            try:
                if is_admin:
                    _mention_rows = _mention_conn.execute(
                        "SELECT id, first_name, last_name, username FROM users WHERE telegram_id>0 ORDER BY first_name LIMIT 100"
                    ).fetchall()
                else:
                    # Collect related user ids: assigned_to, created_by, watchers, shop members
                    _related_ids = set()
                    if task.get("assigned_to"):
                        _related_ids.add(task["assigned_to"])
                    if task.get("created_by"):
                        _related_ids.add(task["created_by"])
                    _w_rows = _mention_conn.execute(
                        "SELECT user_id FROM task_watchers WHERE task_id = ?", (task_id,)
                    ).fetchall()
                    for _wr in _w_rows:
                        _related_ids.add(_wr[0])
                    if task.get("assigned_shop"):
                        _sh_rows = _mention_conn.execute(
                            "SELECT id FROM users WHERE shop_name = ? AND telegram_id > 0",
                            (task["assigned_shop"],)
                        ).fetchall()
                        for _sr in _sh_rows:
                            _related_ids.add(_sr[0])
                    if task.get("assign_all"):
                        _all_rows = _mention_conn.execute(
                            "SELECT id FROM users WHERE telegram_id > 0 LIMIT 100"
                        ).fetchall()
                        for _ar in _all_rows:
                            _related_ids.add(_ar[0])
                    if _related_ids:
                        _ph = ",".join("?" * len(_related_ids))
                        _mention_rows = _mention_conn.execute(
                            f"SELECT id, first_name, last_name, username FROM users WHERE id IN ({_ph})",
                            list(_related_ids)
                        ).fetchall()
                    else:
                        _mention_rows = []
            finally:
                _mention_conn.close()
            ctx["mention_users"] = [
                {"db_id": r[0], "name": f"{r[1] or ''} {r[2] or ''}".strip() or r[3] or str(r[0]),
                 "username": r[3] or ""}
                for r in _mention_rows
            ]
        except Exception:
            ctx["mention_users"] = []

        # Build comments_text for AI context (existing)
        try:
            ctx["comments_text"] = " | ".join(
                c.get("text", "") for c in ctx["comments"] if c.get("text")
            )[:2000]
        except Exception:
            ctx["comments_text"] = ""

    except Exception as e:
        logger.error("task_detail: %s", e)
        ctx["error"] = "Ошибка загрузки задачи."

    return request.app.state.templates.TemplateResponse(
        request, "tasks/detail.html", ctx
    )


# ─── TASK ATTACHMENTS ─────────────────────────────────────────────────────────

@router.post("/tasks/{task_id}/upload")
async def task_upload_attachment(
    request: Request,
    task_id: int,
    csrf_token: str = Form(default=""),
    files: List[UploadFile] = File(default=[]),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db") or ""
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    if any(f and f.filename for f in files):
        try:
            from web.rate_store import check_rate_limit
            if not check_rate_limit(f"task_upload:{telegram_id}", max_requests=20, window_seconds=60):
                return RedirectResponse(url=f"/tasks/{task_id}?msg=rate_limit", status_code=303)
        except Exception:
            pass

    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0

        if not is_admin and task.get("assigned_to") != my_db_id and not task.get("assign_all"):
            return RedirectResponse(url=f"/tasks/{task_id}?msg=access_denied", status_code=303)

        existing_count = db.get_task_attachments_count(task_id)
        if existing_count >= MAX_TASK_FILES:
            return RedirectResponse(url=f"/tasks/{task_id}?msg=too_many_files", status_code=303)

        uploads_dir = _uploads_dir_tasks(org_db)
        month_dir = datetime.now().strftime("%Y-%m")
        month_path = os.path.join(uploads_dir, month_dir)
        os.makedirs(month_path, exist_ok=True)
        saved = []
        for f in files:
            if not f or not f.filename:
                continue
            if existing_count + len(saved) >= MAX_TASK_FILES:
                break
            try:
                raw_data = await f.read()
                if len(raw_data) == 0 or len(raw_data) > MAX_TASK_FILE_SIZE:
                    continue
                safe_name = _safe_filename_tasks(f.filename)
                mime = mimetypes.guess_type(safe_name)[0] or "application/octet-stream"
                uid = uuid.uuid4().hex[:12]
                dest = os.path.join(month_path, f"{uid}_{safe_name}")
                import anyio as _anyio
                _dest_cap2, _data_cap2 = dest, raw_data
                await _anyio.to_thread.run_sync(
                    lambda: open(_dest_cap2, "wb").write(_data_cap2)
                )
                saved.append({
                    "file_path": dest, "file_name": f.filename[:255],
                    "file_type": mime, "file_size": len(raw_data),
                    "uploaded_by": my_db_id,
                })
            except Exception as exc:
                logger.warning("task_upload skip: %s", exc)

        if saved:
            db.add_task_attachments(task_id, my_db_id, saved)

    except Exception as e:
        logger.error("task_upload_attachment: %s", e)

    return RedirectResponse(url=f"/tasks/{task_id}", status_code=303)


@router.get("/tasks/attachment/{att_id}")
def task_serve_attachment(request: Request, att_id: int):
    from web.auth import get_session_user
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db") or ""
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    try:
        db = get_web_db(telegram_id, org_db)
        row = db.get_task_attachment(att_id)
        if not row:
            return Response(content="Файл не найден", status_code=404)

        # row = (file_path, file_name, file_type, task_id, user_id)
        fpath, fname, ftype, task_id, _ = row

        if not is_admin:
            task = db.get_task(task_id)
            conn = db.get_connection()
            try:
                my_row = conn.execute(
                    "SELECT id, shop_name FROM users WHERE telegram_id = ?", (telegram_id,)
                ).fetchone()
            finally:
                conn.close()
            my_db_id = my_row[0] if my_row else 0
            my_shop = (my_row[1] or "") if my_row else ""
            shop_ok = bool(my_shop and task and task.get("assigned_shop") == my_shop)
            if not task or (
                task.get("assigned_to") != my_db_id
                and not task.get("assign_all")
                and not shop_ok
            ):
                return Response(content="Доступ запрещён", status_code=403)

        if not os.path.isfile(fpath):
            return Response(content="Файл не найден на диске", status_code=404)

        exp_uploads = os.path.abspath(_uploads_dir_tasks(org_db))
        real_fpath = os.path.abspath(fpath)
        if not real_fpath.startswith(exp_uploads):
            return Response(content="Доступ запрещён", status_code=403)

        return FileResponse(fpath, media_type=ftype or "application/octet-stream",
                            filename=fname or os.path.basename(fpath))
    except Exception as exc:
        logger.error(f"task_serve_attachment error: {exc}")
        return Response(content="Ошибка", status_code=500)


@router.post("/tasks/attachment/{att_id}/delete")
def task_delete_attachment(
    request: Request,
    att_id: int,
    csrf_token: str = Form(default=""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url="/tasks?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db") or ""
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")
    task_id = 0

    try:
        db = get_web_db(telegram_id, org_db)
        row = db.get_task_attachment(att_id)
        if not row:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        # row = (file_path, file_name, file_type, task_id, user_id)
        _, _, _, task_id, _ = row

        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0

        # delete_task_attachment returns (success, file_path)
        ok, fpath = db.delete_task_attachment(att_id, my_db_id, is_admin=is_admin)
        if ok and fpath and os.path.isfile(fpath):
            exp_uploads = os.path.abspath(_uploads_dir_tasks(org_db))
            real_fpath = os.path.abspath(fpath)
            if real_fpath.startswith(exp_uploads):
                try:
                    os.remove(fpath)
                except OSError as e:
                    logger.warning(f"task_delete_attachment fs: {e}")

        return RedirectResponse(url=f"/tasks/{task_id or ''}", status_code=303)
    except Exception as exc:
        logger.error(f"task_delete_attachment: {exc}")
        return RedirectResponse(url=f"/tasks/{task_id or ''}", status_code=303)


# ─── CHANGE STATUS ───────────────────────────────────────────────────────────

@router.post("/tasks/{task_id}/status")
def task_change_status(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
    status: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from fastapi.responses import JSONResponse as _JSONResponse

    wants_json = request.headers.get("x-requested-with") == "XMLHttpRequest"

    user = get_session_user(request)
    if not user:
        if wants_json:
            return _JSONResponse({"error": "not_auth"}, status_code=401)
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        if wants_json:
            return _JSONResponse({"error": "csrf"}, status_code=403)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    if status not in STATUS_LABELS:
        return RedirectResponse(url=f"/tasks/{task_id}?msg=bad_status", status_code=303)

    # Non-admin cannot cancel tasks
    if not is_admin and status == 'cancelled':
        return RedirectResponse(url=f"/tasks/{task_id}?msg=bad_status", status_code=303)

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0

        # Deny unresolved users
        if not my_db_id:
            return RedirectResponse(url="/tasks?msg=error", status_code=303)

        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        if not is_admin and task.get("assigned_to") != my_db_id:
            return RedirectResponse(url="/tasks", status_code=303)

        # Server-side enforce valid transition for non-admin
        _VALID_NEXT = {'new': 'in_progress', 'in_progress': 'review', 'review': 'done'}
        if not is_admin:
            allowed = _VALID_NEXT.get(task.get('status', ''))
            if status != allowed:
                return RedirectResponse(url=f"/tasks/{task_id}?msg=bad_status", status_code=303)

        old_status = task.get('status', '')
        db.update_task_status(task_id, status)
        try:
            db.add_task_history(task_id, my_db_id, 'status',
                                STATUS_LABELS.get(old_status, old_status),
                                STATUS_LABELS.get(status, status))
        except Exception:
            pass

        # Automation rules: status changed
        try:
            from task_automation import run_rules as _run_rules
            _ev_task = dict(task)
            _ev_task['_old_status'] = old_status
            _ev_task['status'] = status
            _run_rules(db, 'status_changed', _ev_task, my_db_id)
        except Exception as _are:
            logger.warning("task_change_status automation: %s", _are)

        # Record per-user completion for team tasks
        if status in ('done', 'review') and (task.get('assign_all') or task.get('assigned_shop')):
            try:
                db.record_task_user_completion(task_id, my_db_id, status)
            except Exception:
                pass

        # Spawn next recurring task when entering 'done' (unified + idempotent)
        try:
            from task_automation import spawn_recurring_if_done
            _new_assigned_to = spawn_recurring_if_done(db, task, old_status, status)
            if _new_assigned_to:
                _notify_recurring_spawn(db, _new_assigned_to, task['title'])
        except Exception as _re:
            logger.warning("task_change_status spawn recurring: %s", _re)

        creator_id = task.get("created_by")
        if creator_id and creator_id != my_db_id:
            tg_id = _get_user_tg_id(db, creator_id)
            if tg_id:
                status_text = STATUS_LABELS.get(status, status)
                db.add_notification_to_history(
                    creator_id, "task_status",
                    f"📋 Задача «{task['title']}»: {status_text}"
                )
                _send_tg_task_notify(
                    tg_id,
                    f"📋 <b>Статус задачи изменён</b>\n\n"
                    f"<b>{_html.escape(task['title'])}</b>\n"
                    f"Новый статус: {_html.escape(status_text)}"
                )
                try:
                    from web.push_utils import send_web_push
                    send_web_push(tg_id, "📋 Задача обновлена", f"«{task['title']}»: {status_text}", "/tasks")
                except Exception:
                    pass

        if status == 'done' and task.get("assigned_to") and task["assigned_to"] != my_db_id:
            assigned_tg_id = _get_user_tg_id(db, task["assigned_to"])
            if assigned_tg_id:
                db.add_notification_to_history(
                    task["assigned_to"], "task_status",
                    f"✅ Задача выполнена: {task['title']}"
                )
                _send_tg_task_notify(
                    assigned_tg_id,
                    f"✅ <b>Задача выполнена</b>\n\n<b>{_html.escape(task['title'])}</b>"
                )
                try:
                    from web.push_utils import send_web_push
                    send_web_push(assigned_tg_id, "✅ Задача выполнена", task['title'], "/tasks")
                except Exception:
                    pass

    except Exception as e:
        logger.error("task_change_status: %s", e)
        if wants_json:
            return _JSONResponse({"error": "server_error"}, status_code=500)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)

    if wants_json:
        return _JSONResponse({"ok": True, "new_status": status,
                              "label": STATUS_LABELS[status]})
    return RedirectResponse(url=f"/tasks/{task_id}?msg=status_updated", status_code=303)


# ─── COMMENT ─────────────────────────────────────────────────────────────────

@router.post("/tasks/{task_id}/comment")
def task_add_comment(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
    text: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    text = text.strip()[:2000]
    if not text:
        return RedirectResponse(url=f"/tasks/{task_id}?msg=no_text", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or "") if my_row and len(my_row) > 1 else ""

        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)
        if not is_admin:
            _can_comment = (
                task.get("assigned_to") == my_db_id
                or task.get("created_by") == my_db_id
                or task.get("assign_all")
                or (task.get("assigned_shop") and my_shop and task["assigned_shop"] == my_shop)
            )
            if not _can_comment:
                return RedirectResponse(url="/tasks", status_code=303)

        db.add_task_comment(task_id, my_db_id, text)

        # Уведомить создателя И исполнителя (кроме комментатора) — все 3 канала
        _c_notify: set = set()
        _c_creator = task.get("created_by")
        _c_assignee = task.get("assigned_to")
        if _c_creator and _c_creator != my_db_id:
            _c_notify.add(_c_creator)
        if _c_assignee and _c_assignee != my_db_id:
            _c_notify.add(_c_assignee)

        # Expand recipients for shared tasks (assign_all or assigned_shop)
        _task_assign_all = task.get("assign_all", False)
        _task_assigned_shop = task.get("assigned_shop") or None
        if _task_assign_all or _task_assigned_shop:
            try:
                _exp_conn = db.get_connection()
                try:
                    if _task_assign_all:
                        _exp_rows = _exp_conn.execute(
                            "SELECT id FROM users WHERE id != ?", (my_db_id,)
                        ).fetchall()
                    else:
                        _exp_rows = _exp_conn.execute(
                            "SELECT id FROM users WHERE shop_name = ? AND id != ?",
                            (_task_assigned_shop, my_db_id)
                        ).fetchall()
                finally:
                    _exp_conn.close()
                for _er in _exp_rows:
                    _c_notify.add(_er[0])
            except Exception:
                pass

        # Phase 4.2: также уведомить наблюдателей (watchers)
        try:
            _watcher_ids = db.get_task_watcher_db_ids(task_id)
            for _wid in _watcher_ids:
                if _wid != my_db_id:
                    _c_notify.add(_wid)
        except Exception:
            pass

        # Cap at 30 recipients to avoid flooding small-team orgs
        _c_notify = set(list(_c_notify)[:30])

        for _nid in _c_notify:
            _ntg = _get_user_tg_id(db, _nid)
            if _ntg:
                try:
                    db.add_notification_to_history(
                        _nid, "task_status",
                        f"💬 Новый комментарий к задаче «{task['title']}»"
                    )
                except Exception:
                    pass
                _send_tg_task_notify(
                    _ntg,
                    f"💬 <b>Новый комментарий</b>\n\nЗадача: <b>{_html.escape(task['title'])}</b>"
                )
                try:
                    from web.push_utils import send_web_push
                    send_web_push(_ntg, "💬 Новый комментарий",
                                  f"Задача: {task['title']}", f"/tasks/{task_id}")
                except Exception:
                    pass

        # Phase 4.1: @mentions — найти @username в тексте и уведомить
        import re as _re
        _mentions = _re.findall(r'@(\w+)', text)
        if _mentions:
            try:
                _conn_m = db.get_connection()
                try:
                    _m_rows = _conn_m.execute(
                        "SELECT id, username, telegram_id FROM users WHERE telegram_id>0"
                    ).fetchall()
                finally:
                    _conn_m.close()
                _username_map = {
                    (r[1] or "").lower(): (r[0], r[2]) for r in _m_rows if r[1]
                }
                for _mention in _mentions:
                    _ml = _mention.lower()
                    if _ml in _username_map:
                        _m_db_id, _m_tg = _username_map[_ml]
                        if _m_db_id != my_db_id and _m_tg:
                            try:
                                db.add_notification_to_history(
                                    _m_db_id, "task_status",
                                    f"🔔 Вас упомянули в задаче «{task['title']}»"
                                )
                            except Exception:
                                pass
                            _send_tg_task_notify(
                                _m_tg,
                                f"🔔 <b>Вас упомянули</b> в комментарии к задаче\n\n"
                                f"«{_html.escape(task['title'])}»\n"
                                f"<i>{_html.escape(text[:200])}</i>"
                            )
                            try:
                                from web.push_utils import send_web_push
                                send_web_push(_m_tg, "🔔 Вас упомянули",
                                              f"Задача: {task['title']}", f"/tasks/{task_id}")
                            except Exception:
                                pass
            except Exception as _me:
                logger.error("task_comment mentions: %s", _me)

    except Exception as e:
        logger.error("task_add_comment: %s", e)

    return RedirectResponse(url=f"/tasks/{task_id}?msg=commented", status_code=303)


# ─── CHECKLIST TOGGLE ────────────────────────────────────────────────────────

@router.post("/tasks/{task_id}/checklist/{item_id}/toggle")
def task_toggle_checklist(
    request: Request,
    task_id: int,
    item_id: int,
    csrf_token: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from fastapi.responses import JSONResponse as _JSONResponse

    wants_json = request.headers.get("x-requested-with") == "XMLHttpRequest"

    user = get_session_user(request)
    if not user:
        if wants_json:
            return _JSONResponse({"error": "not_auth"}, status_code=401)
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        if wants_json:
            return _JSONResponse({"error": "csrf"}, status_code=403)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0

        # Deny unresolved users
        if not my_db_id:
            return RedirectResponse(url="/tasks?msg=error", status_code=303)

        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)
        if not is_admin:
            _my_shop_cl = ""
            try:
                _cl_conn = db.get_connection()
                _cl_row = _cl_conn.execute(
                    "SELECT shop_name FROM users WHERE id = ?", (my_db_id,)
                ).fetchone()
                _cl_conn.close()
                _my_shop_cl = (_cl_row[0] or "") if _cl_row else ""
            except Exception:
                pass
            _can_toggle = (
                task.get("assigned_to") == my_db_id
                or task.get("assign_all")
                or (task.get("assigned_shop") and _my_shop_cl and task["assigned_shop"] == _my_shop_cl)
            )
            if not _can_toggle:
                return RedirectResponse(url="/tasks", status_code=303)

        # IDOR guard: verify item_id belongs to this task_id
        item_conn = db.get_connection()
        try:
            item_row = item_conn.execute(
                "SELECT task_id FROM task_checklist WHERE id = ?", (item_id,)
            ).fetchone()
        finally:
            item_conn.close()
        if not item_row or item_row[0] != task_id:
            return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)

        # Block checklist edits on closed tasks for non-admin
        if not is_admin and task.get("status") in ('done', 'cancelled'):
            return RedirectResponse(url=f"/tasks/{task_id}", status_code=303)

        db.toggle_task_checklist_item(item_id, my_db_id)
        if wants_json:
            _ic = db.get_connection()
            try:
                _ir = _ic.execute(
                    "SELECT is_done FROM task_checklist WHERE id = ?", (item_id,)
                ).fetchone()
            finally:
                _ic.close()
            new_is_done = bool(_ir[0]) if _ir else False
            return _JSONResponse({"ok": True, "is_done": new_is_done})
    except Exception as e:
        logger.error("task_toggle_checklist: %s", e)
        if wants_json:
            return _JSONResponse({"error": "server_error"}, status_code=500)

    return RedirectResponse(url=f"/tasks/{task_id}", status_code=303)


# ─── EDIT ────────────────────────────────────────────────────────────────────

@router.post("/tasks/{task_id}/rate")
def task_rate(request: Request, task_id: int,
              csrf_token: str = Form(""),
              rating: int = Form(0),
              rating_comment: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url=f"/tasks/{task_id}", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)
    if rating < 1 or rating > 5:
        return RedirectResponse(url=f"/tasks/{task_id}?msg=rate_invalid", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        if task.get('status') != 'done':
            return RedirectResponse(url=f"/tasks/{task_id}?msg=task_not_done", status_code=303)

        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else None

        # Запретить само-оценку: создатель не может оценить свою задачу
        if my_db_id and task.get('created_by') == my_db_id:
            return RedirectResponse(url=f"/tasks/{task_id}?msg=cannot_rate_own", status_code=303)

        db.rate_task(task_id, rating, rating_comment)
        try:
            db.add_task_history(task_id, my_db_id, 'rated', None,
                                f"{'⭐' * rating} ({rating}/5)" + (f": {rating_comment}" if rating_comment.strip() else ""))
        except Exception:
            pass
        return RedirectResponse(url=f"/tasks/{task_id}?msg=rated_ok", status_code=303)
    except Exception as e:
        logger.error("task_rate: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)


@router.post("/tasks/{task_id}/rate-inline")
def task_rate_inline(request: Request, task_id: int,
                     csrf_token: str = Form(""),
                     rating: int = Form(0)):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from fastapi.responses import JSONResponse

    user = get_session_user(request)
    if not user:
        return JSONResponse({"ok": False, "error": "auth"}, status_code=401)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return JSONResponse({"ok": False, "error": "forbidden"}, status_code=403)
    if not verify_csrf_token(request, csrf_token):
        return JSONResponse({"ok": False, "error": "csrf"}, status_code=403)
    if rating < 1 or rating > 5:
        return JSONResponse({"ok": False, "error": "invalid_rating"}, status_code=400)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return JSONResponse({"ok": False, "error": "not_found"}, status_code=404)

        if task.get('status') != 'done':
            return JSONResponse({"ok": False, "error": "task_not_done"}, status_code=400)

        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else None

        if my_db_id and task.get('created_by') == my_db_id:
            return JSONResponse({"ok": False, "error": "cannot_rate_own"}, status_code=400)

        db.rate_task(task_id, rating)
        try:
            db.add_task_history(task_id, my_db_id, 'rated', None,
                                f"{'⭐' * rating} ({rating}/5) — из списка задач")
        except Exception:
            pass
        return JSONResponse({"ok": True, "rating": rating})
    except Exception as e:
        logger.error("task_rate_inline: %s", e)
        return JSONResponse({"ok": False, "error": "server_error"}, status_code=500)


@router.post("/tasks/{task_id}/remind")
def task_set_reminder(request: Request, task_id: int,
                      csrf_token: str = Form(""),
                      remind_in: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    import datetime as _dt

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        hours = float(remind_in) if remind_in else 0
        if hours <= 0 or hours > 720:
            return RedirectResponse(url=f"/tasks/{task_id}?msg=remind_invalid", status_code=303)

        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        if not my_row:
            return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)
        my_db_id = my_row[0]

        remind_at = (_dt.datetime.utcnow() + _dt.timedelta(hours=hours)).strftime("%Y-%m-%d %H:%M:%S")
        db.add_task_reminder(task_id, my_db_id, remind_at)
        try:
            db.add_task_history(task_id, my_db_id, 'remind_set', None,
                                f"через {remind_in}ч ({remind_at} UTC)")
        except Exception:
            pass
        return RedirectResponse(url=f"/tasks/{task_id}?msg=remind_ok", status_code=303)
    except Exception as e:
        logger.error("task_set_reminder: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)


@router.post("/tasks/{task_id}/duplicate")
def task_duplicate(request: Request, task_id: int, csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url=f"/tasks/{task_id}", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0

        cl_items = [i.get('text', '') for i in (task.get('checklist') or []) if i.get('text', '').strip()]
        new_title = f"[Копия] {task['title']}"
        new_id = db.create_task(
            title=new_title,
            description=task.get('description', ''),
            topic_id=task.get('topic_id'),
            created_by=my_db_id,
            assigned_to=task.get('assigned_to'),
            assigned_shop=task.get('assigned_shop'),
            assign_all=task.get('assign_all', 0),
            priority=task.get('priority', 'normal'),
            deadline=None,
            checklist=cl_items,
        )
        if new_id:
            try:
                db.add_task_history(new_id, my_db_id, 'created', None, f"{new_title} (скопирована из #{task_id})")
            except Exception:
                pass
        return RedirectResponse(url=f"/tasks/{new_id}/edit", status_code=303)
    except Exception as e:
        logger.error("task_duplicate: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)


@router.get("/tasks/{task_id}/edit")
def task_edit_form(request: Request, task_id: int):
    from web.auth import get_session_user, get_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url=f"/tasks/{task_id}", status_code=302)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    from billing_utils import has_module as _has_module, has_extension as _has_ext
    _tpro = _has_module(telegram_id, 'tasks_pro')
    _tai  = _tpro and _has_ext(telegram_id, 'tasks_ai')

    ctx = {
        "request": request, "user": user, "is_admin": True,
        "topics": [], "staff_list": [], "shops_list": [],
        "priority_labels": PRIORITY_LABELS,
        "csrf_token": get_csrf_token(request),
        "edit_task": None, "error": None,
        "chat_available": _chat_available(telegram_id),
        "search_items_json": "[]", "init_selected_json": "[]", "init_assign_all": "false",
        "tasks_pro": _tpro, "tasks_ai": _tai,
    }
    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=302)
        ctx["edit_task"] = task
        ctx["topics"] = db.get_task_topics()
        ctx["staff_list"] = _get_staff_list(db)
        ctx["shops_list"] = _get_shops_list(db)
        ctx["search_items_json"] = _build_search_items_json(ctx["staff_list"], ctx["shops_list"])
        init_sel, init_all = _build_init_selected_json(task)
        ctx["init_selected_json"] = init_sel
        ctx["init_assign_all"] = init_all
    except Exception as e:
        logger.error("task_edit_form: %s", e)

    return request.app.state.templates.TemplateResponse(
        request, "tasks/form.html", ctx
    )


@router.post("/tasks/{task_id}/edit")
def task_edit_post(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
    title: str = Form(""),
    description: str = Form(""),
    topic_id: str = Form(""),
    assign_mode: str = Form("person"),
    assigned_to: str = Form(""),
    assigned_shop: str = Form(""),
    priority: str = Form("normal"),
    deadline: str = Form(""),
    recurrence: str = Form("none"),
    recipients_json: str = Form(""),
    estimated_hours: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url=f"/tasks/{task_id}", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}/edit?msg=csrf_error", status_code=303)

    title = title.strip()
    if not title:
        return RedirectResponse(url=f"/tasks/{task_id}/edit?msg=no_title", status_code=303)
    _VALID_ASSIGN_MODES = ('person', 'shop', 'all', 'none', '')
    if assign_mode not in _VALID_ASSIGN_MODES:
        return RedirectResponse(url=f"/tasks/{task_id}/edit?msg=bad_assign", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        old_task = db.get_task(task_id)
        if not old_task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        _topic_id = int(topic_id) if topic_id.isdigit() else None
        _deadline = deadline.strip() or None
        if priority not in PRIORITY_LABELS:
            priority = 'normal'

        _assigned_to = None
        _assigned_shop_val = None
        _assign_all = 0
        if assign_mode == "person":
            _assigned_to = int(assigned_to) if assigned_to.isdigit() else None
        elif assign_mode == "shop":
            _assigned_shop_val = assigned_shop.strip() or None
        elif assign_mode == "all":
            _assign_all = 1

        db.update_task(task_id, title, description.strip(), _topic_id,
                       _assigned_to, None, priority, _deadline,
                       assigned_shop=_assigned_shop_val, assign_all=_assign_all,
                       recurrence=recurrence if recurrence not in ('none', '') else None)
        # Phase 4.3: estimated_hours
        try:
            _eh_e = float(estimated_hours.strip().replace(',', '.')) if estimated_hours.strip() else None
            db.update_task_estimated_hours(task_id, _eh_e)
        except Exception:
            pass
        try:
            conn_ed = db.get_connection()
            ed_row = conn_ed.execute("SELECT id FROM users WHERE telegram_id = ?",
                                    (telegram_id,)).fetchone()
            conn_ed.close()
            _ed_uid = ed_row[0] if ed_row else None
            db.add_task_history(task_id, _ed_uid, 'edit', None, title)
            if old_task.get('priority') != priority:
                db.add_task_history(task_id, _ed_uid, 'priority',
                                    PRIORITY_LABELS.get(old_task.get('priority', ''), old_task.get('priority', '')),
                                    PRIORITY_LABELS.get(priority, priority))
        except Exception:
            pass

        # Уведомить исполнителя об изменениях задачи
        _deadline_str = f"\n📅 Срок: {_fmt_deadline(_deadline)}" if _deadline else ""
        if assign_mode == "person" and _assigned_to:
            tg_id = _get_user_tg_id(db, _assigned_to)
            if tg_id:
                if _assigned_to != old_task.get("assigned_to"):
                    # Новый исполнитель
                    try:
                        db.add_notification_to_history(
                            _assigned_to, "task_assigned", f"📋 Назначена задача: {title}"
                        )
                    except Exception:
                        pass
                    _send_tg_task_notify(
                        tg_id,
                        f"📋 <b>Вам назначена задача</b>\n\n"
                        f"<b>{_html.escape(title)}</b>"
                        f"{_deadline_str}\n\n"
                        f"🌐 Откройте веб-кабинет для подробностей."
                    )
                    try:
                        from web.push_utils import send_web_push
                        send_web_push(tg_id, "📋 Назначена задача", title, "/tasks")
                    except Exception:
                        pass
                else:
                    # Тот же исполнитель — задача обновлена
                    try:
                        db.add_notification_to_history(
                            _assigned_to, "task_assigned", f"📋 Задача обновлена: {title}"
                        )
                    except Exception:
                        pass
                    _send_tg_task_notify(
                        tg_id,
                        f"📋 <b>Задача обновлена</b>\n\n"
                        f"<b>{_html.escape(title)}</b>"
                        f"{_deadline_str}\n\n"
                        f"🌐 Откройте веб-кабинет для деталей."
                    )
                    try:
                        from web.push_utils import send_web_push
                        send_web_push(tg_id, "📋 Задача обновлена", title, "/tasks")
                    except Exception:
                        pass
        elif assign_mode == "shop" and _assigned_shop_val:
            _is_new_shop = _assigned_shop_val != old_task.get("assigned_shop")
            _sh_notif_msg = f"📋 Назначена задача: {title}" if _is_new_shop else f"📋 Задача обновлена: {title}"
            _sh_tg_hdr = "Вам назначена задача" if _is_new_shop else "Задача обновлена"
            for _sm_uid, _sm_tg in _get_shop_members_tg_ids(db, _assigned_shop_val):
                try:
                    db.add_notification_to_history(_sm_uid, "task_assigned", _sh_notif_msg)
                except Exception:
                    pass
                if _sm_tg:
                    _send_tg_task_notify(
                        _sm_tg,
                        f"📋 <b>{_html.escape(_sh_tg_hdr)}</b>\n\n<b>{_html.escape(title)}</b>{_deadline_str}\n\n"
                        f"🌐 Откройте веб-кабинет."
                    )
                    try:
                        from web.push_utils import send_web_push
                        send_web_push(_sm_tg, f"📋 {_sh_tg_hdr}", title, "/tasks")
                    except Exception:
                        pass

        # Automation rules: reassignment fires task_assigned
        try:
            _reassigned = (
                (_assigned_to and _assigned_to != old_task.get("assigned_to"))
                or (_assigned_shop_val and _assigned_shop_val != old_task.get("assigned_shop"))
                or (_assign_all and not old_task.get("assign_all"))
            )
            if _reassigned:
                from task_automation import run_rules as _run_rules
                _new_t = db.get_task(task_id)
                if _new_t:
                    _run_rules(db, 'task_assigned', dict(_new_t), _ed_uid if '_ed_uid' in dir() else None)
        except Exception as _are:
            logger.warning("task_edit_post automation: %s", _are)
    except Exception as e:
        logger.error("task_edit_post: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}/edit?msg=error", status_code=303)

    return RedirectResponse(url=f"/tasks/{task_id}?msg=saved", status_code=303)


# ─── DELETE ──────────────────────────────────────────────────────────────────

@router.post("/tasks/{task_id}/delete")
def task_delete(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if user.get("role") not in ("owner", "admin", "super_admin"):
        return RedirectResponse(url=f"/tasks/{task_id}", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        # Сначала удалить файлы вложений с диска
        try:
            attachments = db.get_task_attachments(task_id)
            for att in attachments:
                att_id = att["id"] if isinstance(att, dict) else att[0]
                ok, fpath = db.delete_task_attachment(att_id, telegram_id, is_admin=True)
                if ok and fpath:
                    import os as _os
                    try:
                        _os.remove(fpath)
                    except OSError:
                        pass
        except Exception as _ae:
            logger.warning("task_delete attachments cleanup: %s", _ae)
        db.delete_task(task_id)
    except Exception as e:
        logger.error("task_delete: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)

    return RedirectResponse(url="/tasks?msg=deleted", status_code=303)


@router.post("/tasks/{task_id}/my_complete")
def task_my_complete(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
    report: str = Form(""),
):
    """Сотрудник помечает командную задачу как выполненную со своей стороны."""
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    try:
        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id = ?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or '') if my_row else ''
        if not my_db_id:
            return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)

        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        # Only allowed on assign_all or shop tasks
        if not (task.get('assign_all') or task.get('assigned_shop')):
            return RedirectResponse(url=f"/tasks/{task_id}", status_code=303)

        # For shop-specific tasks: verify user belongs to the assigned shop
        if task.get('assigned_shop') and not task.get('assign_all'):
            if my_shop != (task.get('assigned_shop') or ''):
                return RedirectResponse(url=f"/tasks/{task_id}", status_code=303)

        db.record_task_user_completion(task_id, my_db_id, 'done')

        # Сохранить текстовый отчёт как комментарий
        report_text = (report or "").strip()[:1000]
        if report_text:
            try:
                db.add_task_comment(task_id, my_db_id, f"📋 Отчёт: {report_text}")
            except Exception as _re:
                logger.warning("task_my_complete save report: %s", _re)

        # Авто-переход в «На проверку» когда все участники отметили выполнение
        try:
            _assign_all = task.get('assign_all', False)
            _assigned_shop = task.get('assigned_shop', '')
            if (_assign_all or _assigned_shop) and task.get('status') not in ('done', 'cancelled', 'review'):
                all_staff = _get_staff_list(db)
                task_members = all_staff if _assign_all else [
                    s for s in all_staff if s.get('shop') == _assigned_shop
                ]
                completions_list = db.get_task_user_completions(task_id)
                completed_ids = {c['user_id'] for c in completions_list}
                member_ids = {m['id'] for m in task_members}
                if member_ids and member_ids.issubset(completed_ids):
                    _mc_old = task.get('status', '')
                    db.update_task_status(task_id, 'review')
                    if _mc_old != 'review':
                        try:
                            from task_automation import run_rules as _run_rules
                            _ev_task = dict(task)
                            _ev_task['_old_status'] = _mc_old
                            _ev_task['status'] = 'review'
                            _run_rules(db, 'status_changed', _ev_task, my_db_id)
                        except Exception as _are:
                            logger.warning("task_my_complete automation: %s", _are)
        except Exception as _ae:
            logger.error("task_my_complete auto-advance: %s", _ae)

        # Уведомить создателя задачи — все 3 канала
        _creator_id = task.get("created_by")
        if _creator_id and _creator_id != my_db_id:
            _creator_tg = _get_user_tg_id(db, _creator_id)
            if _creator_tg:
                try:
                    _nc = db.get_connection()
                    try:
                        _nr = _nc.execute(
                            "SELECT first_name FROM users WHERE telegram_id = ?", (telegram_id,)
                        ).fetchone()
                    finally:
                        _nc.close()
                    _uname = (_nr[0] or "Сотрудник") if _nr else "Сотрудник"
                except Exception:
                    _uname = "Сотрудник"
                _cmsg = f"✅ «{task['title']}» — {_uname} выполнил(а)"
                try:
                    db.add_notification_to_history(_creator_id, "task_status", _cmsg)
                except Exception:
                    pass
                _send_tg_task_notify(
                    _creator_tg,
                    f"✅ <b>Задача отмечена выполненной</b>\n\n"
                    f"<b>{_html.escape(task['title'])}</b>\n{_html.escape(_uname)}"
                )
                try:
                    from web.push_utils import send_web_push
                    send_web_push(_creator_tg, "✅ Задача выполнена",
                                  f"«{task['title']}»: {_uname}", f"/tasks/{task_id}")
                except Exception:
                    pass

    except Exception as e:
        logger.error("task_my_complete: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)

    return RedirectResponse(url=f"/tasks/{task_id}?msg=done", status_code=303)


# ═══════════════════════════════════════════════════════════════════════════════
# ── Phase 4.2: Наблюдатели задачи (watch / unwatch) ─────────────────────────
# ═══════════════════════════════════════════════════════════════════════════════

@router.post("/tasks/{task_id}/watch")
def task_watch(request: Request, task_id: int, csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")
    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id=?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or '') if my_row else ''
        can_view = (
            is_admin
            or task.get('created_by') == my_db_id
            or task.get('assigned_to') == my_db_id
            or task.get('assign_all')
            or (task.get('assigned_shop') and my_shop and task['assigned_shop'] == my_shop)
        )
        if my_db_id and can_view:
            db.add_task_watcher(task_id, my_db_id)
    except Exception as e:
        logger.error("task_watch: %s", e)
    return RedirectResponse(url=f"/tasks/{task_id}?msg=watching", status_code=303)


@router.post("/tasks/{task_id}/unwatch")
def task_unwatch(request: Request, task_id: int, csrf_token: str = Form("")):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")
    try:
        db = get_web_db(telegram_id, org_db)
        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id=?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or '') if my_row else ''
        can_view = (
            is_admin
            or task.get('created_by') == my_db_id
            or task.get('assigned_to') == my_db_id
            or task.get('assign_all')
            or (task.get('assigned_shop') and my_shop and task['assigned_shop'] == my_shop)
        )
        if my_db_id and can_view:
            db.remove_task_watcher(task_id, my_db_id)
    except Exception as e:
        logger.error("task_unwatch: %s", e)
    return RedirectResponse(url=f"/tasks/{task_id}?msg=unwatched", status_code=303)


# ═══════════════════════════════════════════════════════════════════════════════
# ── Phase 4.3: Учёт времени (log time / delete time log) ─────────────────────
# ═══════════════════════════════════════════════════════════════════════════════

@router.post("/tasks/{task_id}/log_time")
def task_log_time(
    request: Request,
    task_id: int,
    csrf_token: str = Form(""),
    hours: str = Form(""),
    minutes: str = Form(""),
    note: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db
    from billing_utils import has_module as _has_module

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)

    if not _has_module(telegram_id, 'tasks_pro'):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=pro_required", status_code=303)

    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    try:
        h = int(hours or 0)
        m = int(minutes or 0)
        total_minutes = h * 60 + m
        if total_minutes <= 0:
            return RedirectResponse(url=f"/tasks/{task_id}?msg=bad_time", status_code=303)

        db = get_web_db(telegram_id, org_db)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id, shop_name FROM users WHERE telegram_id=?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        my_shop = (my_row[1] or '') if my_row else ''

        task = db.get_task(task_id)
        if not task:
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)

        if not is_admin:
            can_act = (
                task.get('assigned_to') == my_db_id
                or task.get('created_by') == my_db_id
                or task.get('assign_all')
                or (task.get('assigned_shop') and my_shop and task['assigned_shop'] == my_shop)
            )
            if not can_act:
                return RedirectResponse(url=f"/tasks/{task_id}?msg=forbidden", status_code=303)

        db.log_task_time(task_id, my_db_id, total_minutes, note.strip()[:500])

    except Exception as e:
        logger.error("task_log_time: %s", e)
        return RedirectResponse(url=f"/tasks/{task_id}?msg=error", status_code=303)

    return RedirectResponse(url=f"/tasks/{task_id}?msg=time_logged", status_code=303)


@router.post("/tasks/{task_id}/log_time/{log_id}/delete")
def task_delete_time_log(
    request: Request,
    task_id: int,
    log_id: int,
    csrf_token: str = Form(""),
):
    from web.auth import get_session_user, verify_csrf_token
    from web.deps import get_web_db

    user = get_session_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)
    if not verify_csrf_token(request, csrf_token):
        return RedirectResponse(url=f"/tasks/{task_id}?msg=csrf_error", status_code=303)

    telegram_id = int(user["sub"])
    org_db = user.get("org_db")
    if not org_db:
        return RedirectResponse("/dashboard", status_code=302)
    is_admin = user.get("role") in ("owner", "admin", "super_admin")

    try:
        db = get_web_db(telegram_id, org_db)
        if not db.get_task(task_id):
            return RedirectResponse(url="/tasks?msg=not_found", status_code=303)
        conn = db.get_connection()
        try:
            my_row = conn.execute(
                "SELECT id FROM users WHERE telegram_id=?", (telegram_id,)
            ).fetchone()
        finally:
            conn.close()
        my_db_id = my_row[0] if my_row else 0
        db.delete_task_time_log(log_id, my_db_id, is_admin)
    except Exception as e:
        logger.error("task_delete_time_log: %s", e)

    return RedirectResponse(url=f"/tasks/{task_id}", status_code=303)
