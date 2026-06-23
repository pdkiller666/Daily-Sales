"""
Административная панель управления платежной системой и подписками
"""
import asyncio
import json
import logging
from aiogram import Router, F
from aiogram.types import CallbackQuery, Message
from aiogram.fsm.context import FSMContext
from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton

from database import Database
from keyboards import back_button, invalidate_web_url_cache
from states import PaymentSystemStates
from env_manager import env_manager
from db_utils import clear_state_keep_org
from utils import he
from message_utils import fsm_edit

# Создаем роутер
payment_system_router = Router()


def _get_db():
    """Локальная БД для платёжного/подписочного функционала"""
    return Database('data/shop_bot.db')

_CANCEL_PROMO_KB = InlineKeyboardMarkup(inline_keyboard=[
    [InlineKeyboardButton(text="❌ Отмена", callback_data="cancel_add_promo")]
])

@payment_system_router.callback_query(F.data == "payment_system_admin")
async def payment_system_admin_menu(callback: CallbackQuery):
    """Главное меню управления платежной системой"""
    db = _get_db()
    if not callback.message:
        await callback.answer("❌ Сообщение слишком старое.", show_alert=True)
        return
        
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    # Получаем текущие настройки
    payment_settings = db.get_payment_settings()
    
    text = "💳 <b>Управление платежной системой</b>\n\n"
    text += "🔧 <b>Текущие настройки:</b>\n"
    text += f"💳 Номер карты: {payment_settings.get('card_number', 'Не установлен')}\n"
    text += f"👤 Получатель: {payment_settings.get('recipient_name', 'Не установлен')}\n"
    text += f"🏦 Банк: {payment_settings.get('bank_name', 'Не установлен')}\n\n"
    
    # Статистика подписок
    subscriptions_stats = db.get_subscriptions_statistics()
    text += "📊 <b>Статистика подписок:</b>\n"
    text += f"👥 Всего подписчиков: {subscriptions_stats['total_subscribers']}\n"
    text += f"💰 Месячная выручка: {subscriptions_stats['monthly_revenue']:,.0f} ₽\n"
    text += f"📈 Конверсия: {subscriptions_stats['conversion_rate']:.1f}%\n\n"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="💳 Заявки на оплату", callback_data="pending_payments")],
        [InlineKeyboardButton(text="💳 Настройки оплаты", callback_data="payment_settings")],
        [InlineKeyboardButton(text="📏 Лимиты ресурсов", callback_data="manage_plans")],
        [InlineKeyboardButton(text="📊 Статистика платежей", callback_data="payment_statistics")],
        [InlineKeyboardButton(text="👥 Управление подписками", callback_data="manage_subscriptions")],
        [InlineKeyboardButton(text="🧩 Модули и пакеты", callback_data="billing_modules_admin")],
        [InlineKeyboardButton(text="🎁 Промокоды", callback_data="manage_promocodes")],
        [InlineKeyboardButton(text="🎫 Пробный период", callback_data="trial_settings")],
        [back_button("system_admin_panel")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data == "payment_settings")
async def payment_settings_menu(callback: CallbackQuery):
    """Настройки платежной системы"""
    db = _get_db()
    if not callback.message:
        await callback.answer("❌ Сообщение слишком старое.", show_alert=True)
        return
        
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    payment_settings = db.get_payment_settings()
    provider = db.get_payment_provider()
    from payment_provider import provider_label
    provider_display = provider_label(provider)

    text = "💳 <b>Настройки платежной системы</b>\n\n"
    text += f"🔀 <b>Активный провайдер:</b> {provider_display}\n\n"
    text += "Настройте реквизиты для получения платежей:\n\n"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"🔀 Провайдер оплаты: {provider_display}", callback_data="payment_provider_select")],
        [InlineKeyboardButton(text=f"💳 Номер карты: {payment_settings.get('card_number', 'Не установлен')}", callback_data="set_card_number")],
        [InlineKeyboardButton(text=f"👤 Получатель: {payment_settings.get('recipient_name', 'Не установлен')}", callback_data="set_recipient_name")],
        [InlineKeyboardButton(text=f"🏦 Банк: {payment_settings.get('bank_name', 'Не установлен')}", callback_data="set_bank_name")],
        [InlineKeyboardButton(text="💬 Инструкция для пользователей", callback_data="set_payment_instruction")],
        [InlineKeyboardButton(text="🧪 Тестовый платеж", callback_data="test_payment")],
        [back_button("payment_system_admin")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data == "manage_plans")
async def manage_plans_menu(callback: CallbackQuery):
    """Управление тарифными планами"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    plans = db.get_all_subscription_plans()

    text = "📏 <b>Лимиты ресурсов</b>\n\n"
    text += ("Базовая подписка задаёт только жёсткие лимиты ресурсов. "
             "Возможности (аналитика, команда, уведомления, интеграции) подключаются "
             "модулями в разделе «🧩 Модули и пакеты».\n\n")

    buttons = []
    for plan in plans:
        plan_id = plan[0]
        name = plan[1]
        max_products = plan[7] if len(plan) > 7 else None
        max_shops = plan[8] if len(plan) > 8 else None
        max_sales = plan[9] if len(plan) > 9 else None

        def _fmt(v):
            if v is None:
                return "—"
            return "∞" if v == -1 else str(v)

        text += (f"<b>{he(name)}</b>\n"
                 f"   📦 Товары: {_fmt(max_products)} · "
                 f"🏪 Магазины: {_fmt(max_shops)} · "
                 f"💰 Продажи/мес: {_fmt(max_sales)}\n\n")
        buttons.append([InlineKeyboardButton(
            text=f"✏️ {name}", callback_data=f"edit_plan_{plan_id}")])

    if not plans:
        text += "<i>Тарифов пока нет.</i>\n"

    buttons.append([back_button("payment_system_admin")])
    keyboard = InlineKeyboardMarkup(inline_keyboard=buttons)

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data == "payment_statistics")
async def payment_statistics_menu(callback: CallbackQuery):
    """Статистика платежей"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    stats = db.get_detailed_payment_statistics()
    
    text = "📊 <b>Детальная статистика платежей</b>\n\n"
    
    # Общая статистика
    text += "📈 <b>Общие показатели:</b>\n"
    text += f"💰 Общая выручка: {stats['total_revenue']:,.0f} ₽\n"
    text += f"📦 Всего платежей: {stats['total_payments']}\n"
    text += f"💳 Средний чек: {stats['average_payment']:,.0f} ₽\n"
    text += f"📅 За последний месяц: {stats['monthly_revenue']:,.0f} ₽\n\n"
    
    # По тарифам
    text += "💎 <b>По тарифным планам:</b>\n"
    for plan_name, plan_stats in stats['by_plans'].items():
        text += f"• {plan_name}: {plan_stats['count']} подписок ({plan_stats['revenue']:,.0f} ₽)\n"
    
    text += f"\n📊 <b>Конверсия:</b> {stats['conversion_rate']:.1f}%\n"
    text += f"🔄 <b>Продления:</b> {stats['renewal_rate']:.1f}%\n"

    # G5: модульный биллинг
    try:
        bstats = db.get_billing_stats()
        text += "\n🧩 <b>Модульный биллинг:</b>\n"
        text += f"📦 Активных модулей: {bstats['modules_active']} · пакетов: {bstats['bundles_active']}\n"
        text += f"👥 Клиентов с модулями: {bstats['clients_count']}\n"
        text += f"💰 Выручка модулей (30 дн.): {bstats['revenue_30d']:,.0f} ₽\n"
        if bstats['per_module']:
            text += "🔝 <b>Топ модулей:</b>\n"
            for pm in bstats['per_module'][:5]:
                text += f"   • {he(pm['key'])}: {pm['count']} ({pm['revenue']:,.0f} ₽)\n"
    except Exception:
        pass

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📅 За период", callback_data="stats_by_period")],
        [InlineKeyboardButton(text="📋 Экспорт отчета", callback_data="export_payment_report")],
        [InlineKeyboardButton(text="📈 Графики", callback_data="payment_charts")],
        [back_button("payment_system_admin")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data == "manage_subscriptions")
async def manage_subscriptions_menu(callback: CallbackQuery):
    """Управление подписками пользователей"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    active_subscriptions = db.get_all_active_subscriptions()

    text = "👥 <b>Управление подписками</b>\n\n"
    text += f"Активных подписок: {len(active_subscriptions)}\n\n"
    text += ("Откройте единую карточку клиента по Telegram ID: базовая подписка, "
             "лимиты, активные модули/расширения/пакеты — с выдачей, продлением и отзывом.\n\n")

    if active_subscriptions:
        text += "📋 <b>Последние подписки:</b>\n"
        for sub in active_subscriptions[:5]:
            plan_type = sub[2] if len(sub) > 2 else "Неизвестно"
            end_date = sub[4] if len(sub) > 4 else "Неизвестно"
            first_name = sub[5] if len(sub) > 5 else "Неизвестно"
            last_name = sub[6] if len(sub) > 6 else ""

            user_name = he(f"{first_name} {last_name}".strip())
            ed = end_date[:10] if isinstance(end_date, str) and len(end_date) >= 10 else end_date
            text += f"• {user_name} — {he(str(plan_type))} до {ed}\n"

        if len(active_subscriptions) > 5:
            text += f"... и ещё {len(active_subscriptions) - 5} подписок\n"

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔍 Карточка клиента по Telegram ID", callback_data="find_customer")],
        [back_button("payment_system_admin")]
    ])

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data == "manage_promocodes")
async def manage_promocodes_menu(callback: CallbackQuery):
    """Управление промокодами"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    promocodes = db.get_all_promocodes()
    
    text = "🎁 <b>Управление промокодами</b>\n\n"

    if promocodes:
        text += "📋 <b>Промокоды:</b>\n"
        for promo in promocodes:
            promo_id, code, discount, disc_type, usage_count, max_usage, is_active, expires_at, allowed_plans, last_used, created_at = promo
            status = "✅" if is_active else "❌"
            disc_label = f"{discount}₽" if disc_type == 'fixed' else f"{discount}%"
            exp_label = f" до {expires_at}" if expires_at else ""
            text += f"{status} <code>{code}</code> — {disc_label} скидка ({usage_count}/{max_usage}){exp_label}\n"
    else:
        text += "Промокоды не созданы\n"

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Создать промокод", callback_data="create_promocode")],
        [InlineKeyboardButton(text="📦 Пакетное создание", callback_data="batch_promocode_start")],
        [InlineKeyboardButton(text="✏️ Редактировать", callback_data="edit_promocode")],
        [InlineKeyboardButton(text="📊 Статистика использования", callback_data="promocode_stats")],
        [InlineKeyboardButton(text="🗑 Удалить промокод", callback_data="delete_promocode")],
        [back_button("payment_system_admin")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

# Обработчики для настройки реквизитов
@payment_system_router.callback_query(F.data == "set_card_number")
async def set_card_number_start(callback: CallbackQuery, state: FSMContext):
    """Начало установки номера карты"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    text = "💳 <b>Установка номера карты</b>\n\n"
    text += "Введите номер карты для получения платежей:\n"
    text += "Формат: 1234 5678 9012 3456\n\n"
    text += "⚠️ Убедитесь, что номер указан корректно!"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button("payment_settings")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.waiting_card_number)

@payment_system_router.message(PaymentSystemStates.waiting_card_number)
async def process_card_number(message: Message, state: FSMContext):
    """Обработка номера карты"""
    db = _get_db()
    card_number = message.text.strip().replace(" ", "")
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⚙️ Настройки оплаты", callback_data="payment_settings")]])

    if not card_number.isdigit() or len(card_number) != 16:
        await fsm_edit(state, message, "❌ Неверный формат номера карты. Введите 16 цифр:", reply_markup=_back_kb)
        return

    formatted_number = f"{card_number[:4]} {card_number[4:8]} {card_number[8:12]} {card_number[12:]}"

    try:
        db.update_payment_setting('card_number', formatted_number)
        await fsm_edit(state, message, f"✅ Номер карты обновлён: {formatted_number}", reply_markup=_back_kb)
    except Exception as e:
        logging.error(f"process_card_number: DB error: {e}")
        await fsm_edit(state, message, "❌ Ошибка при сохранении номера карты. Попробуйте позже.", reply_markup=_back_kb)
    await clear_state_keep_org(state)

@payment_system_router.callback_query(F.data == "set_recipient_name")
async def set_recipient_name_start(callback: CallbackQuery, state: FSMContext):
    """Начало установки имени получателя"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    text = "👤 <b>Установка имени получателя</b>\n\n"
    text += "Введите ФИО получателя платежей:\n"
    text += "Пример: Иванов Иван Иванович"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button("payment_settings")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.waiting_recipient_name)

@payment_system_router.message(PaymentSystemStates.waiting_recipient_name)
async def process_recipient_name(message: Message, state: FSMContext):
    """Обработка имени получателя"""
    db = _get_db()
    recipient_name = message.text.strip()
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⚙️ Настройки оплаты", callback_data="payment_settings")]])

    if len(recipient_name) < 5:
        await fsm_edit(state, message, "❌ Имя получателя слишком короткое. Введите заново:", reply_markup=_back_kb)
        return

    try:
        db.update_payment_setting('recipient_name', recipient_name)
        await fsm_edit(state, message, f"✅ Имя получателя обновлено: {he(recipient_name)}", reply_markup=_back_kb)
    except Exception as e:
        logging.error(f"process_recipient_name: DB error: {e}")
        await fsm_edit(state, message, "❌ Ошибка при сохранении имени получателя. Попробуйте позже.", reply_markup=_back_kb)
    await clear_state_keep_org(state)

@payment_system_router.callback_query(F.data == "set_bank_name")
async def set_bank_name_start(callback: CallbackQuery, state: FSMContext):
    """Начало установки названия банка"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    text = "🏦 <b>Установка названия банка</b>\n\n"
    text += "Введите название банка:\n"
    text += "Пример: Сбербанк, ВТБ, Альфа-Банк"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button("payment_settings")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.waiting_bank_name)

@payment_system_router.message(PaymentSystemStates.waiting_bank_name)
async def process_bank_name(message: Message, state: FSMContext):
    """Обработка названия банка"""
    db = _get_db()
    bank_name = message.text.strip()
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⚙️ Настройки оплаты", callback_data="payment_settings")]])

    if len(bank_name) < 3:
        await fsm_edit(state, message, "❌ Название банка слишком короткое. Введите заново:", reply_markup=_back_kb)
        return

    try:
        db.update_payment_setting('bank_name', bank_name)
        await fsm_edit(state, message, f"✅ Название банка обновлено: {he(bank_name)}", reply_markup=_back_kb)
    except Exception as e:
        logging.error(f"process_bank_name: DB error: {e}")
        await fsm_edit(state, message, "❌ Ошибка при сохранении названия банка. Попробуйте позже.", reply_markup=_back_kb)
    await clear_state_keep_org(state)

@payment_system_router.callback_query(F.data == "set_web_interface_url")
async def set_web_interface_url_start(callback: CallbackQuery, state: FSMContext):
    """Показать текущий URL и предложить изменить/удалить"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    current_url = db.get_web_interface_url()

    text = "🌐 <b>URL веб-интерфейса</b>\n\n"
    if current_url:
        text += f"Текущий URL:\n<code>{he(current_url)}</code>\n\n"
        text += "Кнопка «Веб-интерфейс» показывается всем пользователям в главном меню.\n\n"
        text += "Нажмите <b>Изменить</b> чтобы ввести новый URL, или <b>Удалить</b> чтобы скрыть кнопку."
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(text="✏️ Изменить", callback_data="edit_web_interface_url"),
                InlineKeyboardButton(text="🗑 Удалить", callback_data="clear_web_interface_url"),
            ],
            [back_button("system_admin_panel")],
        ])
    else:
        text += "URL не настроен. Кнопка «Веб-интерфейс» скрыта для всех пользователей.\n\n"
        text += "Введите URL вашего веб-интерфейса (должен начинаться с <code>https://</code>):"
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [back_button("system_admin_panel")],
        ])
        await state.update_data(anchor_msg_id=callback.message.message_id)
        await state.set_state(PaymentSystemStates.waiting_web_interface_url)

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")


@payment_system_router.callback_query(F.data == "edit_web_interface_url")
async def edit_web_interface_url_start(callback: CallbackQuery, state: FSMContext):
    """Запросить новый URL"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    text = (
        "🌐 <b>Изменение URL веб-интерфейса</b>\n\n"
        "Введите новый URL (должен начинаться с <code>https://</code>):\n\n"
        "Например: <code>https://yourapp.replit.app</code>"
    )
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button("system_admin_panel")],
    ])
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.waiting_web_interface_url)


@payment_system_router.message(PaymentSystemStates.waiting_web_interface_url)
async def process_web_interface_url(message: Message, state: FSMContext):
    """Сохранить новый URL веб-интерфейса"""
    db = _get_db()
    url = message.text.strip() if message.text else ""
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="🔧 Системная панель", callback_data="system_admin_panel")]])

    if not url.startswith("https://") and not url.startswith("http://"):
        await fsm_edit(state, message,
            "❌ Неверный формат. URL должен начинаться с <code>https://</code>\n\nПопробуйте ещё раз:",
            reply_markup=_back_kb)
        return

    try:
        db.set_web_interface_url(url)
        invalidate_web_url_cache()
        await fsm_edit(state, message,
            f"✅ URL веб-интерфейса обновлён:\n<code>{he(url)}</code>\n\n"
            "Кнопка «🌐 Веб-интерфейс» появится в главном меню всех пользователей.",
            reply_markup=_back_kb)
    except Exception as e:
        logging.error(f"process_web_interface_url: DB error: {e}")
        await fsm_edit(state, message, "❌ Ошибка при сохранении URL. Попробуйте позже.", reply_markup=_back_kb)
    await clear_state_keep_org(state)


@payment_system_router.callback_query(F.data == "clear_web_interface_url")
async def clear_web_interface_url_handler(callback: CallbackQuery):
    """Удалить URL — скрыть кнопку для всех пользователей"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    try:
        db.set_web_interface_url(None)
        invalidate_web_url_cache()
        text = (
            "🌐 <b>URL веб-интерфейса удалён</b>\n\n"
            "Кнопка «Веб-интерфейс» скрыта для всех пользователей."
        )
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔧 Системная панель", callback_data="system_admin_panel")],
        ])
        await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    except Exception as e:
        logging.error(f"clear_web_interface_url: DB error: {e}")
        await callback.answer("❌ Ошибка при удалении URL", show_alert=True)


@payment_system_router.callback_query(F.data.startswith("edit_plan_"))
async def edit_plan_details(callback: CallbackQuery, state: FSMContext):
    """Детали редактирования конкретного плана"""
    db = _get_db()
    plan_id = int(callback.data.split("_")[2])
    plan_details = db.get_subscription_plan_details(plan_id)
    
    if not plan_details:
        await callback.answer("❌ План не найден")
        return

    await callback.answer()
    await state.update_data(editing_plan_id=plan_id)

    def _lim(v):
        return "∞ Безлимит" if v == -1 else (str(v) if v is not None else "—")

    text = f"📏 <b>Лимиты тарифа: {he(plan_details['name'])}</b>\n\n"
    text += f"📦 Товары: {_lim(plan_details['max_products'])}\n"
    text += f"🏪 Магазины: {_lim(plan_details['max_shops'])}\n"
    text += f"💰 Продажи/месяц: {_lim(plan_details['max_sales_per_month'])}\n\n"
    text += ("🧩 Возможности (аналитика, команда, уведомления, интеграции) "
             "определяются модулями, а не тарифом.\n\n")
    text += "Выберите лимит для изменения (-1 = безлимит):"

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📦 Лимит товаров", callback_data="edit_field_max_products")],
        [InlineKeyboardButton(text="🏪 Лимит магазинов", callback_data="edit_field_max_shops")],
        [InlineKeyboardButton(text="💰 Лимит продаж/мес", callback_data="edit_field_max_sales")],
        [back_button("manage_plans")]
    ])

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

# Обработчики редактирования полей планов
@payment_system_router.callback_query(F.data.startswith("edit_field_"))
async def edit_plan_field_start(callback: CallbackQuery, state: FSMContext):
    """Начало редактирования поля плана"""
    db = _get_db()
    field = callback.data.split("_", 2)[2]
    data = await state.get_data()
    plan_id = data.get('editing_plan_id')
    
    if not plan_id:
        await callback.answer("❌ Ошибка: план не выбран")
        return
    
    field_names = {
        'max_products': '📦 лимит товаров',
        'max_shops': '🏪 лимит магазинов',
        'max_sales': '💰 лимит продаж/мес',
    }

    if field not in field_names:
        await callback.answer("❌ Неизвестное поле")
        return

    await state.update_data(editing_field=field)

    # Ввод нового значения
    await callback.answer()
    text = f"✏️ <b>Редактирование: {field_names[field]}</b>\n\n"
    text += "Введите новый лимит (-1 для безлимита):"

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button(f"edit_plan_{plan_id}")]
    ])

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.editing_plan_field)

@payment_system_router.message(PaymentSystemStates.editing_plan_field)
async def process_plan_field_edit(message: Message, state: FSMContext):
    """Обработка редактирования поля плана"""
    db = _get_db()
    data = await state.get_data()
    plan_id = data.get('editing_plan_id')
    field = data.get('editing_field')

    if not plan_id or not field:
        await fsm_edit(state, message, "❌ Ошибка: данные не найдены")
        return
    
    new_value = message.text.strip()
    
    try:
        # Валидация и преобразование значений (только лимиты ресурсов)
        if field in ['max_products', 'max_shops', 'max_sales']:
            new_value = int(new_value)
            if new_value < -1 or new_value == 0:
                raise ValueError("Значение должно быть -1 (безлимит) или больше 0")
            db_field = 'max_sales_per_month' if field == 'max_sales' else field
        else:
            raise ValueError("Неизвестное поле")

        # Обновляем в базе данных
        success = db.update_subscription_plan_field(plan_id, db_field, new_value)

        if success:
            limit_text = "∞ Безлимит" if new_value == -1 else str(new_value)
            text = f"✅ <b>Лимит обновлён!</b>\n\n🎯 Новое значение: {limit_text}"

            keyboard = InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="✏️ Продолжить редактирование", callback_data=f"edit_plan_{plan_id}")],
                [InlineKeyboardButton(text="📏 К лимитам ресурсов", callback_data="manage_plans")]
            ])
            await fsm_edit(state, message, text, reply_markup=keyboard)
        else:
            await fsm_edit(state, message, "❌ Ошибка при обновлении плана")

    except ValueError as e:
        await fsm_edit(state, message, f"❌ Ошибка валидации: {e}")
    except Exception as e:
        await fsm_edit(state, message, f"❌ Произошла ошибка: {e}")

    await clear_state_keep_org(state)

# ─────────────────── CREATION WIZARD ───────────────────

_SKIP_EXPIRES_KB = InlineKeyboardMarkup(inline_keyboard=[
    [InlineKeyboardButton(text="⏩ Без срока (пропустить)", callback_data="promo_skip_expires")],
    [InlineKeyboardButton(text="❌ Отмена", callback_data="cancel_add_promo")],
])
_SKIP_PLANS_KB = InlineKeyboardMarkup(inline_keyboard=[
    [InlineKeyboardButton(text="⏩ Все тарифы (пропустить)", callback_data="promo_skip_plans")],
    [InlineKeyboardButton(text="❌ Отмена", callback_data="cancel_add_promo")],
])

@payment_system_router.callback_query(F.data == "create_promocode")
async def create_promocode_start(callback: CallbackQuery, state: FSMContext):
    """Шаг 1: Ввод кода промокода или генерация случайного."""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    await callback.answer()
    text = (
        "🎁 <b>Создание промокода — шаг 1/5</b>\n\n"
        "Введите код промокода:\n"
        "<i>Пример: SALE20, NEWCLIENT, VIP50</i>\n"
        "Только латиница, цифры и знак подчёркивания, минимум 3 символа.\n\n"
        "Или нажмите «🎲 Случайный» для автоматической генерации."
    )
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🎲 Случайный код", callback_data="gen_promo_code")],
        [back_button("manage_promocodes")],
    ])
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.waiting_promocode)


@payment_system_router.callback_query(F.data == "gen_promo_code")
async def gen_promo_code(callback: CallbackQuery, state: FSMContext):
    """Генерировать случайный код и перейти к шагу 2."""
    import secrets, string
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    await callback.answer()
    code = ''.join(secrets.choice(string.ascii_uppercase + string.digits) for _ in range(8))
    await state.update_data(promocode=code)
    await _ask_discount_type(callback.message, code)
    await state.set_state(PaymentSystemStates.waiting_discount_type)


@payment_system_router.message(PaymentSystemStates.waiting_promocode)
async def process_promocode_code(message: Message, state: FSMContext):
    """Принять введённый код и перейти к шагу 2."""
    code = message.text.strip().upper()
    if not code.replace('_', '').isalnum() or len(code) < 3:
        await fsm_edit(state, message,
            "❌ Код должен содержать только латинские буквы, цифры и подчёркивания (минимум 3 символа)",
            reply_markup=_CANCEL_PROMO_KB)
        return
    await state.update_data(promocode=code)
    await _ask_discount_type(message, code, state=state)
    await state.set_state(PaymentSystemStates.waiting_discount_type)


async def _ask_discount_type(target, code: str, state=None):
    """Шаг 2: выбор типа скидки."""
    text = (
        f"🎁 <b>Создание промокода — шаг 2/5</b>\n\n"
        f"Код: <code>{code}</code>\n\n"
        f"Выберите тип скидки:"
    )
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📊 Процент (%)", callback_data="promo_disc_type_percent")],
        [InlineKeyboardButton(text="💵 Фиксированная сумма (₽)", callback_data="promo_disc_type_fixed")],
        [InlineKeyboardButton(text="❌ Отмена", callback_data="cancel_add_promo")],
    ])
    if hasattr(target, 'edit_text'):
        await target.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    elif state is not None:
        await fsm_edit(state, target, text, reply_markup=keyboard)
    else:
        await target.answer(text, reply_markup=keyboard, parse_mode="HTML")


@payment_system_router.callback_query(F.data.in_({"promo_disc_type_percent", "promo_disc_type_fixed"}))
async def promo_disc_type_chosen(callback: CallbackQuery, state: FSMContext):
    """Шаг 3: ввод размера скидки."""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    await callback.answer()
    disc_type = 'fixed' if 'fixed' in callback.data else 'percent'
    await state.update_data(discount_type=disc_type)
    data = await state.get_data()
    code = data.get('promocode', '')
    if disc_type == 'fixed':
        hint = "Введите сумму скидки в рублях:\n<i>Пример: 100, 250, 500</i>"
    else:
        hint = "Введите размер скидки в процентах (1–90):\n<i>Пример: 10, 20, 50</i>"
    text = (
        f"🎁 <b>Создание промокода — шаг 3/5</b>\n\n"
        f"Код: <code>{code}</code>\n"
        f"Тип: {'Фиксированная (₽)' if disc_type == 'fixed' else 'Процент (%)'}\n\n"
        f"{hint}"
    )
    await callback.message.edit_text(text, reply_markup=_CANCEL_PROMO_KB, parse_mode="HTML")
    await state.set_state(PaymentSystemStates.waiting_discount)


@payment_system_router.message(PaymentSystemStates.waiting_discount)
async def process_discount(message: Message, state: FSMContext):
    """Шаг 4 (одиночный) или финал (batch): ввод скидки."""
    data = await state.get_data()
    disc_type = data.get('discount_type', 'percent')
    try:
        discount = int(message.text.strip())
        if disc_type == 'percent' and not (1 <= discount <= 90):
            raise ValueError()
        if disc_type == 'fixed' and discount <= 0:
            raise ValueError()
    except ValueError:
        hint = "от 1 до 90" if disc_type == 'percent' else "больше 0"
        await fsm_edit(state, message, f"❌ Неверный формат. Введите число {hint}", reply_markup=_CANCEL_PROMO_KB)
        return
    await state.update_data(discount=discount)

    # Batch mode: сразу сохраняем набор кодов (max_usage=1 у каждого)
    if data.get('batch_mode'):
        await _save_batch_promocodes(message, state)
        return

    code = data.get('promocode', '')
    disc_label = f"{discount}₽" if disc_type == 'fixed' else f"{discount}%"
    text = (
        f"🎁 <b>Создание промокода — шаг 4/5</b>\n\n"
        f"Код: <code>{code}</code> · Скидка: {disc_label}\n\n"
        f"Введите максимальное количество использований:\n"
        f"<i>Пример: 10, 50, 100</i>"
    )
    await fsm_edit(state, message, text, reply_markup=_CANCEL_PROMO_KB)
    await state.set_state(PaymentSystemStates.waiting_max_usage)


@payment_system_router.message(PaymentSystemStates.waiting_max_usage)
async def process_max_usage(message: Message, state: FSMContext):
    """Шаг 5: срок действия."""
    try:
        max_usage = int(message.text.strip())
        if max_usage <= 0:
            raise ValueError()
    except ValueError:
        await fsm_edit(state, message, "❌ Неверный формат. Введите число больше 0", reply_markup=_CANCEL_PROMO_KB)
        return
    await state.update_data(max_usage=max_usage)
    data = await state.get_data()
    code = data.get('promocode', '')
    disc_type = data.get('discount_type', 'percent')
    discount = data.get('discount', 0)
    disc_label = f"{discount}₽" if disc_type == 'fixed' else f"{discount}%"
    text = (
        f"🎁 <b>Создание промокода — шаг 5/5</b>\n\n"
        f"Код: <code>{code}</code> · Скидка: {disc_label} · Лимит: {max_usage}\n\n"
        f"Введите дату окончания действия промокода в формате <b>ГГГГ-ММ-ДД</b>:\n"
        f"<i>Пример: 2025-12-31</i>\n\n"
        f"Или нажмите «Пропустить» — промокод будет бессрочным."
    )
    await fsm_edit(state, message, text, reply_markup=_SKIP_EXPIRES_KB)
    await state.set_state(PaymentSystemStates.waiting_expires_at)


@payment_system_router.callback_query(F.data == "promo_skip_expires")
async def promo_skip_expires(callback: CallbackQuery, state: FSMContext):
    """Пропустить срок действия → спросить о тарифах."""
    await callback.answer()
    await state.update_data(expires_at=None)
    await _ask_allowed_plans(callback.message, state)


@payment_system_router.message(PaymentSystemStates.waiting_expires_at)
async def process_expires_at(message: Message, state: FSMContext):
    """Обработать дату окончания действия."""
    from datetime import datetime
    raw = message.text.strip()
    try:
        dt = datetime.strptime(raw, '%Y-%m-%d')
        if dt.date() <= datetime.utcnow().date():
            await fsm_edit(state, message, "❌ Дата должна быть в будущем. Введите в формате ГГГГ-ММ-ДД.", reply_markup=_SKIP_EXPIRES_KB)
            return
        expires_at = raw
    except ValueError:
        await fsm_edit(state, message, "❌ Неверный формат. Введите дату как ГГГГ-ММ-ДД (например, 2025-12-31).", reply_markup=_SKIP_EXPIRES_KB)
        return
    await state.update_data(expires_at=expires_at)
    await _ask_allowed_plans(message, state)


async def _ask_allowed_plans(target, state: FSMContext):
    """Спросить об ограничении по тарифным планам."""
    from subscription_handlers import get_current_subscription_plans
    plans = get_current_subscription_plans()
    plan_names = [v['name'] for v in plans.values()]
    plan_list = "\n".join(f"• <code>{n}</code>" for n in plan_names) if plan_names else "(нет активных планов)"
    text = (
        f"🎁 <b>Ограничение по тарифу (необязательно)</b>\n\n"
        f"Доступные тарифы:\n{plan_list}\n\n"
        f"Введите названия тарифов через запятую, чтобы промокод работал только для них:\n"
        f"<i>Пример: Стандарт, Премиум</i>\n\n"
        f"Или нажмите «Пропустить» — промокод будет работать для всех тарифов."
    )
    if hasattr(target, 'edit_text'):
        await target.edit_text(text, reply_markup=_SKIP_PLANS_KB, parse_mode="HTML")
    else:
        await fsm_edit(state, target, text, reply_markup=_SKIP_PLANS_KB)
    await state.set_state(PaymentSystemStates.waiting_allowed_plans)


@payment_system_router.callback_query(F.data == "promo_skip_plans")
async def promo_skip_plans(callback: CallbackQuery, state: FSMContext):
    """Пропустить ограничение по тарифам → сохранить промокод."""
    await callback.answer()
    await state.update_data(allowed_plans=None)
    await _save_new_promocode(callback.message, state)


@payment_system_router.message(PaymentSystemStates.waiting_allowed_plans)
async def process_allowed_plans(message: Message, state: FSMContext):
    """Обработать список разрешённых тарифов → сохранить промокод."""
    import json
    raw = message.text.strip()
    plan_names = [p.strip() for p in raw.split(',') if p.strip()]
    allowed_plans = json.dumps(plan_names, ensure_ascii=False) if plan_names else None
    await state.update_data(allowed_plans=allowed_plans)
    await _save_new_promocode(message, state)


async def _save_new_promocode(target, state: FSMContext):
    """Финальное сохранение промокода после всех шагов wizard'а."""
    db = _get_db()
    data = await state.get_data()
    code = data.get('promocode', '')
    discount = data.get('discount', 0)
    disc_type = data.get('discount_type', 'percent')
    max_usage = data.get('max_usage', 1)
    expires_at = data.get('expires_at')
    allowed_plans = data.get('allowed_plans')

    disc_label = f"{discount}₽" if disc_type == 'fixed' else f"{discount}%"
    exp_label = expires_at or "бессрочный"
    plans_label = allowed_plans or "все тарифы"

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🎁 Управление промокодами", callback_data="manage_promocodes")]
    ])

    try:
        db.create_promocode(
            code=code,
            discount_percent=discount,
            max_usage=max_usage,
            discount_type=disc_type,
            expires_at=expires_at,
            allowed_plans=allowed_plans,
        )
        text = (
            f"✅ <b>Промокод создан!</b>\n\n"
            f"🎁 Код: <code>{code}</code>\n"
            f"💸 Скидка: {disc_label}\n"
            f"📊 Лимит: {max_usage} использований\n"
            f"⏳ Срок: {exp_label}\n"
            f"📋 Тарифы: {plans_label}"
        )
    except Exception as e:
        text = f"❌ Ошибка создания промокода: {he(str(e))}"

    if hasattr(target, 'edit_text'):
        await target.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    else:
        await fsm_edit(state, target, text, reply_markup=keyboard)
    await clear_state_keep_org(state)


@payment_system_router.callback_query(F.data == "cancel_add_promo")
async def cancel_add_promo(callback: CallbackQuery, state: FSMContext):
    """Отмена создания промокода."""
    await clear_state_keep_org(state)
    await callback.answer()
    await manage_promocodes_menu(callback)


# ─────────────────── BATCH CREATION ───────────────────

@payment_system_router.callback_query(F.data == "batch_promocode_start")
async def batch_promocode_start(callback: CallbackQuery, state: FSMContext):
    """Начало пакетного создания промокодов."""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    await callback.answer()
    text = (
        "📦 <b>Пакетное создание промокодов</b>\n\n"
        "Создаёт набор уникальных одноразовых кодов (max_usage=1).\n"
        "Каждый код = префикс + 6 случайных символов.\n\n"
        "Введите количество промокодов (1–100):"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data="manage_promocodes")]
    ])
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.waiting_batch_count)


@payment_system_router.message(PaymentSystemStates.waiting_batch_count)
async def process_batch_count(message: Message, state: FSMContext):
    """Принять количество → спросить префикс."""
    _cancel_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="❌ Отмена", callback_data="manage_promocodes")]])
    try:
        count = int(message.text.strip())
        if not (1 <= count <= 100):
            raise ValueError()
    except ValueError:
        await fsm_edit(state, message, "❌ Введите число от 1 до 100.", reply_markup=_cancel_kb)
        return
    await state.update_data(batch_count=count)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⏩ Без префикса", callback_data="batch_promo_no_prefix")],
        [InlineKeyboardButton(text="❌ Отмена", callback_data="manage_promocodes")],
    ])
    await fsm_edit(state, message,
        f"📦 <b>Количество:</b> {count}\n\n"
        "Введите префикс для кодов (лат. буквы/цифры, мин. 2):\n"
        "<i>Пример: PARTNER → PARTNER_A3K9X2</i>\n\n"
        "Или нажмите «Без префикса».",
        reply_markup=kb)
    await state.set_state(PaymentSystemStates.waiting_batch_prefix)


@payment_system_router.callback_query(F.data == "batch_promo_no_prefix")
async def batch_promo_no_prefix(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    await state.update_data(batch_prefix='')
    await _ask_batch_discount(callback.message, state)


@payment_system_router.message(PaymentSystemStates.waiting_batch_prefix)
async def process_batch_prefix(message: Message, state: FSMContext):
    prefix = message.text.strip().upper()
    _cancel_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="❌ Отмена", callback_data="manage_promocodes")]])
    if not prefix.replace('_', '').isalnum() or len(prefix) < 2:
        await fsm_edit(state, message, "❌ Префикс: только лат. буквы/цифры, минимум 2 символа.", reply_markup=_cancel_kb)
        return
    await state.update_data(batch_prefix=prefix + '_')
    await _ask_batch_discount(message, state)


async def _ask_batch_discount(target, state: FSMContext):
    data = await state.get_data()
    count = data.get('batch_count', 1)
    prefix = data.get('batch_prefix', '')
    text = (
        f"📦 <b>Пакет: {count} кодов · префикс: {prefix or 'нет'}</b>\n\n"
        "Введите размер скидки в процентах (1–90):\n"
        "<i>Все коды получат одинаковую скидку. Каждый код — одноразовый.</i>"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data="manage_promocodes")]
    ])
    await state.update_data(batch_mode=True, discount_type='percent')
    await state.set_state(PaymentSystemStates.waiting_discount)
    if hasattr(target, 'edit_text'):
        await target.edit_text(text, reply_markup=kb, parse_mode="HTML")
    else:
        await fsm_edit(state, target, text, reply_markup=kb)


async def _save_batch_promocodes(message: Message, state: FSMContext):
    """Создать все коды пакета и отчитаться."""
    db = _get_db()
    data = await state.get_data()
    count = data.get('batch_count', 1)
    prefix = data.get('batch_prefix', '')
    discount = data.get('discount', 10)
    disc_type = data.get('discount_type', 'percent')

    created = db.create_promocodes_batch(
        prefix=prefix,
        discount_percent=discount,
        count=count,
        discount_type=disc_type,
    )
    disc_label = f"{discount}₽" if disc_type == 'fixed' else f"{discount}%"
    lines = "\n".join(f"• <code>{c}</code>" for c in created[:10])
    more = f"\n<i>...и ещё {len(created) - 10}</i>" if len(created) > 10 else ""
    skipped = count - len(created)
    skip_note = f"\n⚠️ Пропущено (дубликаты): {skipped}" if skipped > 0 else ""
    text = (
        f"✅ <b>Пакет создан!</b>\n\n"
        f"📦 Создано кодов: <b>{len(created)}</b> · скидка: {disc_label}\n"
        f"{skip_note}\n\n"
        f"{lines}{more}"
    )
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🎁 Управление промокодами", callback_data="manage_promocodes")]
    ])
    await fsm_edit(state, message, text, reply_markup=keyboard)
    await clear_state_keep_org(state)


# ─────────────────── EDIT: EXPIRES_AT ───────────────────

@payment_system_router.callback_query(F.data.startswith("edit_promo_expires_"))
async def edit_promo_expires_start(callback: CallbackQuery, state: FSMContext):
    """Начало редактирования срока действия промокода."""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    promo_id = int(callback.data.split("_")[3])
    db = _get_db()
    promo = db.get_promocode_by_id(promo_id)
    if not promo:
        await callback.answer("❌ Промокод не найден")
        return
    await callback.answer()
    code = promo[1]
    current_exp = promo[7] or "не задан"
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🗑 Убрать срок", callback_data=f"promo_clear_exp_{promo_id}")],
        [InlineKeyboardButton(text="❌ Отмена", callback_data=f"edit_promo_{promo_id}")],
    ])
    await state.set_data({'promo_id': promo_id})
    await callback.message.edit_text(
        f"⏳ <b>Срок действия промокода {code}</b>\n\n"
        f"Текущий срок: {current_exp}\n\n"
        "Введите новую дату в формате <b>ГГГГ-ММ-ДД</b>:\n"
        "<i>Пример: 2025-12-31</i>",
        reply_markup=kb, parse_mode="HTML"
    )
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.editing_promocode_expires)


@payment_system_router.callback_query(F.data.startswith("promo_clear_exp_"))
async def promo_clear_expires(callback: CallbackQuery, state: FSMContext):
    """Убрать срок действия → сделать промокод бессрочным."""
    promo_id = int(callback.data.split("_")[3])
    db = _get_db()
    db.update_promocode(promo_id, expires_at=None)
    await callback.answer("✅ Срок действия убран — промокод стал бессрочным")
    await clear_state_keep_org(state)
    data_cb = type('obj', (object,), {'data': f"edit_promo_{promo_id}", 'from_user': callback.from_user, 'message': callback.message, 'answer': callback.answer})()
    await edit_specific_promocode(data_cb)


@payment_system_router.message(PaymentSystemStates.editing_promocode_expires)
async def process_promo_expires_edit(message: Message, state: FSMContext):
    """Сохранить новую дату окончания."""
    from datetime import datetime
    data = await state.get_data()
    promo_id = data.get('promo_id')
    raw = message.text.strip()
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="❌ Отмена", callback_data=f"edit_promo_{promo_id}")]])
    try:
        dt = datetime.strptime(raw, '%Y-%m-%d')
        if dt.date() <= datetime.utcnow().date():
            raise ValueError("past")
    except ValueError:
        await fsm_edit(state, message, "❌ Дата должна быть будущей, формат ГГГГ-ММ-ДД.", reply_markup=_back_kb)
        return
    db = _get_db()
    db.update_promocode(promo_id, expires_at=raw)
    _done_kb = InlineKeyboardMarkup(inline_keyboard=[[back_button(f"edit_promo_{promo_id}", "⬅️ К промокоду")]])
    await fsm_edit(state, message, f"✅ Срок действия обновлён: {raw}", reply_markup=_done_kb)
    await clear_state_keep_org(state)


# Тестовый платеж
@payment_system_router.callback_query(F.data == "test_payment")
async def test_payment(callback: CallbackQuery):
    """Тестовый платеж"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    await callback.answer()
    settings = db.get_payment_settings()
    
    text = "🧪 <b>Тестовый платеж</b>\n\n"
    text += "💳 <b>Реквизиты для перевода:</b>\n"
    text += f"Карта: <code>{settings['card_number']}</code>\n"
    text += f"Получатель: {settings['recipient_name']}\n"
    text += f"Банк: {settings['bank_name']}\n\n"
    text += "💰 Сумма: 990 ₽ (тест базового тарифа)\n\n"
    text += "📋 После перевода прикрепите скриншот чека"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button("payment_settings")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

# Экспорт отчета по платежам
@payment_system_router.callback_query(F.data == "export_payment_report")
async def export_payment_report(callback: CallbackQuery):
    """Экспорт отчета по платежам"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора", show_alert=True)
        return
    
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from datetime import datetime
        
        # Создаем книгу Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчет по платежам"
        
        # Заголовок
        ws['A1'] = f"Отчет по платежной системе от {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:F1')
        
        # Получаем статистику
        stats = db.get_detailed_payment_statistics()
        subscriptions = db.get_all_active_subscriptions()
        
        # Общая статистика
        row = 3
        ws[f'A{row}'] = "ОБЩАЯ СТАТИСТИКА"
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        
        ws[f'A{row}'] = f"Общая выручка: {stats['total_revenue']:,.0f} ₽"
        row += 1
        ws[f'A{row}'] = f"Всего платежей: {stats['total_payments']}"
        row += 1
        ws[f'A{row}'] = f"Средний чек: {stats['average_payment']:,.0f} ₽"
        row += 1
        ws[f'A{row}'] = f"Месячная выручка: {stats['monthly_revenue']:,.0f} ₽"
        row += 2
        
        # Активные подписки
        ws[f'A{row}'] = "АКТИВНЫЕ ПОДПИСКИ"
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        
        headers = ['Пользователь', 'Тариф', 'Магазин', 'Начало', 'Окончание']
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header).font = Font(bold=True)
        row += 1
        
        for sub in subscriptions:
            ws.cell(row=row, column=1, value=f"{sub[7]} {sub[8]}")  # first_name, last_name
            ws.cell(row=row, column=2, value=sub[2])  # plan_type
            ws.cell(row=row, column=3, value=sub[10])  # shop_name
            ws.cell(row=row, column=4, value=sub[3][:10])  # start_date
            ws.cell(row=row, column=5, value=sub[4][:10])  # end_date
            row += 1
        
        # Сохраняем файл
        filename = f"payment_report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        filepath = f"data/{filename}"
        wb.save(filepath)
        
        # Отправляем файл
        from aiogram.types import FSInputFile
        document = FSInputFile(filepath)
        
        await callback.message.answer_document(
            document=document,
            caption="📊 Отчет по платежной системе"
        )
        
        # Удаляем временный файл
        import os
        os.remove(filepath)
        
        await callback.answer("Отчет экспортирован")
        
    except Exception as e:
        await callback.answer(f"Ошибка экспорта: {str(e)}", show_alert=True)

# Дополнительные обработчики платежной системы

@payment_system_router.callback_query(F.data == "set_payment_instruction")
async def set_payment_instruction(callback: CallbackQuery, state: FSMContext):
    """Установка инструкции для пользователей"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return
    
    await callback.answer()
    text = "📝 <b>Настройка инструкции для пользователей</b>\n\n" \
           "Введите новую инструкцию, которую будут видеть пользователи при оплате:"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button("payment_settings")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.waiting_payment_instruction)

@payment_system_router.message(PaymentSystemStates.waiting_payment_instruction)
async def process_payment_instruction(message: Message, state: FSMContext):
    """Сохранение инструкции для пользователей"""
    db = _get_db()
    instruction = message.text.strip()
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⚙️ Настройки оплаты", callback_data="payment_settings")]])
    if len(instruction) < 5:
        await fsm_edit(state, message, "❌ Инструкция слишком короткая. Введите более подробный текст:", reply_markup=_back_kb)
        return
    try:
        db.update_payment_setting('payment_instruction', instruction)
        await fsm_edit(state, message, "✅ Инструкция для пользователей обновлена.", reply_markup=_back_kb)
    except Exception as e:
        logging.error(f"process_payment_instruction: DB error: {e}")
        await fsm_edit(state, message, "❌ Ошибка при сохранении инструкции. Попробуйте позже.", reply_markup=_back_kb)
    await clear_state_keep_org(state)

@payment_system_router.callback_query(F.data == "payment_charts")
async def payment_charts(callback: CallbackQuery):
    """G5: текстовая визуализация выручки по модулям."""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return

    await callback.answer()
    text = "📈 <b>Выручка по модулям</b>\n\n"
    try:
        bstats = db.get_billing_stats()
        per_module = bstats.get('per_module', [])
        if not per_module:
            text += "📭 Пока нет продаж модулей."
        else:
            max_rev = max((pm['revenue'] or 0) for pm in per_module) or 1
            text += f"💰 Всего за 30 дней: {bstats['revenue_30d']:,.0f} ₽\n"
            text += f"👥 Клиентов: {bstats['clients_count']}\n\n"
            for pm in per_module:
                rev = pm['revenue'] or 0
                bar_len = int(round((rev / max_rev) * 12)) if max_rev else 0
                bar = "█" * bar_len + "░" * (12 - bar_len)
                text += f"{he(pm['key'])}\n  {bar} {rev:,.0f} ₽ ({pm['count']})\n"
    except Exception:
        text += "❌ Не удалось получить данные."

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📊 Статистика", callback_data="payment_statistics")],
        [InlineKeyboardButton(text="🧩 Модули и пакеты", callback_data="billing_modules_admin")],
        [back_button("payment_system_admin")]
    ])

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data == "delete_promocode")
async def delete_promocode(callback: CallbackQuery):
    """Удаление промокодов"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return

    await callback.answer()
    promocodes = db.get_all_promocodes()
    
    text = "🗑 <b>Удаление промокодов</b>\n\n"
    
    if not promocodes:
        text += "📭 Нет доступных промокодов для удаления"
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="➕ Создать промокод", callback_data="create_promocode")],
            [back_button("manage_promocodes")]
        ])
    else:
        text += "Выберите промокод для удаления:\n\n"
        keyboard_buttons = []

        for promo in promocodes:
            promo_id, code, discount, disc_type, usage_count, max_usage, is_active, expires_at, allowed_plans, last_used, created_at = promo
            disc_label = f"{discount}₽" if disc_type == 'fixed' else f"{discount}%"
            text += f"🎁 <b>{code}</b> — {disc_label} (использован {usage_count}/{max_usage})\n"
            keyboard_buttons.append([InlineKeyboardButton(
                text=f"🗑 {code}",
                callback_data=f"delete_promo_{promo_id}"
            )])
        
        keyboard_buttons.append([back_button("manage_promocodes")])
        keyboard = InlineKeyboardMarkup(inline_keyboard=keyboard_buttons)
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data.startswith("delete_promo_"))
async def confirm_delete_promocode(callback: CallbackQuery):
    """Подтверждение удаления промокода"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return
    
    promo_id = int(callback.data.split("_")[2])
    promo = db.get_promocode_by_id(promo_id)
    
    if not promo:
        await callback.answer("❌ Промокод не найден")
        return
    
    await callback.answer()
    code = promo[1]  # code находится во втором поле
    text = f"🗑 <b>Подтверждение удаления</b>\n\n" \
           f"Вы действительно хотите удалить промокод:\n" \
           f"🎁 <b>{code}</b>\n\n" \
           f"❗️ Это действие нельзя отменить!"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Да, удалить", callback_data=f"confirm_delete_promo_{promo_id}")],
        [InlineKeyboardButton(text="❌ Отмена", callback_data="delete_promocode")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data.startswith("confirm_delete_promo_"))
async def execute_delete_promocode(callback: CallbackQuery):
    """Выполнение удаления промокода"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return
    
    promo_id = int(callback.data.split("_")[3])
    promo = db.get_promocode_by_id(promo_id)
    
    if not promo:
        await callback.answer("❌ Промокод не найден")
        return
    
    code = promo[1]
    success = db.delete_promocode(promo_id)
    
    if success:
        text = f"✅ <b>Промокод удален</b>\n\n" \
               f"Промокод <b>{code}</b> был успешно удален из системы."
        await callback.answer("✅ Промокод удален")
    else:
        text = f"❌ <b>Ошибка удаления</b>\n\n" \
               f"Не удалось удалить промокод <b>{code}</b>."
        await callback.answer("❌ Ошибка удаления")
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🗑 Удалить еще", callback_data="delete_promocode")],
        [back_button("manage_promocodes", "⬅️ К управлению промокодами")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data == "edit_promocode")
async def edit_promocode(callback: CallbackQuery):
    """Редактирование промокодов"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return

    await callback.answer()
    promocodes = db.get_all_promocodes()
    
    text = "✏️ <b>Редактирование промокодов</b>\n\n"
    
    if not promocodes:
        text += "📭 Нет доступных промокодов для редактирования"
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="➕ Создать промокод", callback_data="create_promocode")],
            [back_button("manage_promocodes")]
        ])
    else:
        text += "Выберите промокод для редактирования:\n\n"
        keyboard_buttons = []
        
        for promo in promocodes:
            promo_id, code, discount, disc_type, usage_count, max_usage, is_active, expires_at, allowed_plans, last_used, created_at = promo
            status = "✅" if is_active else "❌"
            disc_label = f"{discount}₽" if disc_type == 'fixed' else f"{discount}%"
            exp_label = f" до {expires_at}" if expires_at else ""
            text += f"{status} <b>{code}</b> — {disc_label}{exp_label}\n"
            text += f"   Использован: {usage_count}/{max_usage}\n\n"
            keyboard_buttons.append([InlineKeyboardButton(
                text=f"✏️ {code}",
                callback_data=f"edit_promo_{promo_id}"
            )])
        
        keyboard_buttons.append([back_button("manage_promocodes")])
        keyboard = InlineKeyboardMarkup(inline_keyboard=keyboard_buttons)
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data.startswith("edit_promo_"))
async def edit_specific_promocode(callback: CallbackQuery):
    """Редактирование конкретного промокода"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return
    
    promo_id = int(callback.data.split("_")[2])
    promo = db.get_promocode_by_id(promo_id)
    
    if not promo:
        await callback.answer("❌ Промокод не найден")
        return
    
    await callback.answer()
    promo_id, code, discount_percent, disc_type, current_usage, max_usage, is_active, expires_at, allowed_plans, last_used, created_at = promo
    status = "✅ Активен" if is_active else "❌ Неактивен"
    disc_label = f"{discount_percent}₽" if disc_type == 'fixed' else f"{discount_percent}%"
    exp_label = expires_at or "не задан"
    plans_label = allowed_plans or "все тарифы"

    text = (
        f"✏️ <b>Редактирование промокода</b>\n\n"
        f"🎁 <b>Код:</b> {code}\n"
        f"💰 <b>Скидка:</b> {disc_label} ({disc_type})\n"
        f"📊 <b>Использований:</b> {current_usage}/{max_usage}\n"
        f"⏳ <b>Действует до:</b> {exp_label}\n"
        f"📋 <b>Тарифы:</b> {plans_label}\n"
        f"📈 <b>Статус:</b> {status}\n"
        f"📅 <b>Создан:</b> {created_at[:10]}\n\n"
        f"Что изменить?"
    )

    keyboard_buttons = []

    if is_active:
        keyboard_buttons.append([InlineKeyboardButton(text="❌ Деактивировать", callback_data=f"toggle_promo_status_{promo_id}")])
    else:
        keyboard_buttons.append([InlineKeyboardButton(text="✅ Активировать", callback_data=f"toggle_promo_status_{promo_id}")])

    keyboard_buttons.append([InlineKeyboardButton(text="✏️ Изменить скидку", callback_data=f"edit_promo_discount_{promo_id}")])
    keyboard_buttons.append([InlineKeyboardButton(text="🔢 Изменить лимит использований", callback_data=f"edit_promo_max_usage_{promo_id}")])
    keyboard_buttons.append([InlineKeyboardButton(text="⏳ Срок действия", callback_data=f"edit_promo_expires_{promo_id}")])

    keyboard_buttons.append([InlineKeyboardButton(text="🗑 Удалить промокод", callback_data=f"delete_promo_{promo_id}")])
    keyboard_buttons.append([back_button("edit_promocode", "⬅️ К списку промокодов")])
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=keyboard_buttons)
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data.startswith("toggle_promo_status_"))
async def toggle_promocode_status(callback: CallbackQuery):
    """Переключение статуса промокода"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return
    
    promo_id = int(callback.data.split("_")[3])
    promo = db.get_promocode_by_id(promo_id)
    
    if not promo:
        await callback.answer("❌ Промокод не найден")
        return
    
    current_status = bool(promo[6])  # is_active поле (индекс 6 в новой схеме)
    new_status = not current_status
    
    success = db.update_promocode(promo_id, is_active=new_status)
    
    if success:
        status_text = "активирован" if new_status else "деактивирован"
        await callback.answer(f"✅ Промокод {status_text}")
        # Обновляем информацию о промокоде
        await edit_specific_promocode(callback)
    else:
        await callback.answer("❌ Ошибка изменения статуса")

@payment_system_router.callback_query(F.data.startswith("edit_promo_discount_"))
async def edit_promocode_discount_start(callback: CallbackQuery, state: FSMContext):
    """Начало редактирования скидки промокода"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return
    
    promo_id = int(callback.data.split("_")[3])
    promo = db.get_promocode_by_id(promo_id)
    
    if not promo:
        await callback.answer("❌ Промокод не найден")
        return
    
    await callback.answer()
    code = promo[1]
    current_discount = promo[2]
    
    text = f"✏️ <b>Изменение скидки промокода</b>\n\n" \
           f"🎁 <b>Код:</b> {code}\n" \
           f"💰 <b>Текущая скидка:</b> {current_discount}%\n\n" \
           f"Введите новый размер скидки (от 1 до 99):"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data=f"edit_promo_{promo_id}")]
    ])
    
    await state.set_data({"promo_id": promo_id})
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.editing_promocode_discount)
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.message(PaymentSystemStates.editing_promocode_discount)
async def process_promocode_discount_edit(message: Message, state: FSMContext):
    """Обработка изменения скидки промокода"""
    db = _get_db()
    data = await state.get_data()
    promo_id = data.get("promo_id")
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="❌ Отмена", callback_data=f"edit_promo_{promo_id}")]])

    try:
        new_discount = int(message.text.strip())
        if not (1 <= new_discount <= 99):
            raise ValueError("Скидка должна быть от 1 до 99 процентов")
    except ValueError:
        await fsm_edit(state, message, "❌ Некорректное значение! Введите число от 1 до 99:", reply_markup=_back_kb)
        return

    success = db.update_promocode(promo_id, discount_percent=new_discount)
    if success:
        promo = db.get_promocode_by_id(promo_id)
        code = promo[1]
        await fsm_edit(state, message,
            f"✅ <b>Скидка обновлена</b>\n\nСкидка промокода <b>{code}</b> изменена на {new_discount}%",
            reply_markup=_back_kb)
    else:
        await fsm_edit(state, message, "❌ Ошибка обновления скидки", reply_markup=_back_kb)

    await clear_state_keep_org(state)

@payment_system_router.callback_query(F.data.startswith("edit_promo_max_usage_"))
async def edit_promocode_max_usage_start(callback: CallbackQuery, state: FSMContext):
    """Начало редактирования лимита использований промокода"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return
    
    promo_id = int(callback.data.split("_")[4])
    promo = db.get_promocode_by_id(promo_id)
    
    if not promo:
        await callback.answer("❌ Промокод не найден")
        return
    
    await callback.answer()
    code = promo[1]
    current_usage = promo[3]
    current_max = promo[4]
    
    text = f"✏️ <b>Изменение лимита использований</b>\n\n" \
           f"🎁 <b>Код:</b> {code}\n" \
           f"📊 <b>Использований:</b> {current_usage}/{current_max}\n\n" \
           f"Введите новый лимит использований (минимум {current_usage}):"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data=f"edit_promo_{promo_id}")]
    ])
    
    await state.set_data({"promo_id": promo_id, "current_usage": current_usage})
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.editing_promocode_max_usage)
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.message(PaymentSystemStates.editing_promocode_max_usage)
async def process_promocode_max_usage_edit(message: Message, state: FSMContext):
    """Обработка изменения лимита использований промокода"""
    db = _get_db()
    data = await state.get_data()
    promo_id = data.get("promo_id")
    current_usage = data.get("current_usage")
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="❌ Отмена", callback_data=f"edit_promo_{promo_id}")]])

    try:
        new_max_usage = int(message.text.strip())
        if new_max_usage < current_usage:
            raise ValueError(f"Лимит не может быть меньше текущего количества использований ({current_usage})")
        if new_max_usage <= 0:
            raise ValueError("Лимит должен быть больше 0")
    except ValueError as e:
        await fsm_edit(state, message, f"❌ {str(e)}\nВведите корректное число:", reply_markup=_back_kb)
        return

    success = db.update_promocode(promo_id, max_usage=new_max_usage)
    if success:
        promo = db.get_promocode_by_id(promo_id)
        code = promo[1]
        await fsm_edit(state, message,
            f"✅ <b>Лимит обновлён</b>\n\nЛимит использований промокода <b>{code}</b> изменён на {new_max_usage}",
            reply_markup=_back_kb)
    else:
        await fsm_edit(state, message, "❌ Ошибка обновления лимита", reply_markup=_back_kb)

    await clear_state_keep_org(state)

@payment_system_router.callback_query(F.data == "promocode_stats")
async def promocode_stats(callback: CallbackQuery):
    """Расширенная статистика использования промокодов."""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return

    await callback.answer()
    stats = db.get_promocode_stats_detailed()

    text = "📊 <b>Статистика промокодов</b>\n\n"

    if not stats:
        text += "📭 Нет промокодов для анализа"
    else:
        total_usage = 0
        total_discount_rub = 0.0
        for p in stats:
            usage_pct = (p['usage_count'] / p['max_usage'] * 100) if p['max_usage'] > 0 else 0
            filled = int(usage_pct / 10)
            bar = "█" * filled + "░" * (10 - filled)
            disc_label = f"{p['discount_percent']}₽" if p['discount_type'] == 'fixed' else f"{p['discount_percent']}%"
            status = "✅" if p['is_active'] else "❌"
            last_used = (p['last_used_at'] or '')[:10] or "—"
            exp = f" · до {p['expires_at']}" if p['expires_at'] else ""
            text += (
                f"{status} <b>{p['code']}</b> — {disc_label}{exp}\n"
                f"   [{bar}] {p['usage_count']}/{p['max_usage']} ({usage_pct:.0f}%)\n"
                f"   💸 Скидок выдано: {p['total_discount_rub']:.0f}₽ · последнее: {last_used}\n\n"
            )
            total_usage += p['usage_count']
            total_discount_rub += p['total_discount_rub']

        text += (
            f"──────────────\n"
            f"📊 Итого использований: <b>{total_usage}</b>\n"
            f"💸 Итого скидок: <b>{total_discount_rub:.0f}₽</b>"
        )

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button("manage_promocodes")]
    ])

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

@payment_system_router.callback_query(F.data == "stats_by_period")
async def stats_by_period(callback: CallbackQuery):
    """Статистика за период"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return
    
    await callback.answer()
    text = "📅 <b>Статистика за период</b>\n\n" \
           "Выберите период для анализа:"
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📅 За сегодня", callback_data="stats_today")],
        [InlineKeyboardButton(text="📅 За неделю", callback_data="stats_week")],
        [InlineKeyboardButton(text="📅 За месяц", callback_data="stats_month")],
        [InlineKeyboardButton(text="📅 За год", callback_data="stats_year")],
        [back_button("payment_statistics")]
    ])
    
    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")

# ── Пробный период ────────────────────────────────────────────────────────────

@payment_system_router.callback_query(F.data == "trial_settings")
async def trial_settings_menu(callback: CallbackQuery):
    """Настройки пробного периода для новых пользователей"""
    db = _get_db()
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return

    await callback.answer()
    settings = db.get_payment_settings()
    trial_days = settings.get('trial_days', '14')
    trial_plan = settings.get('trial_plan', 'Премиум')

    # Получаем доступные планы для отображения
    plans = db.get_subscription_plans()
    plans_str = ", ".join(p[1] for p in plans) if plans else "—"

    text = (
        "🎫 <b>Настройки пробного периода</b>\n\n"
        f"⏱ Длительность: <b>{trial_days} дней</b>\n"
        f"💎 Тариф при регистрации: <b>{he(trial_plan)}</b>\n\n"
        f"📋 Доступные тарифы: {he(plans_str)}\n\n"
        "Пробный период активируется автоматически при регистрации "
        "(личный и корпоративный режимы). Установите 0 дней чтобы отключить."
    )

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⏱ Изменить длительность (дней)", callback_data="trial_edit_days")],
        [InlineKeyboardButton(text="💎 Изменить тариф", callback_data="trial_edit_plan")],
        [back_button("payment_system_admin")]
    ])

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")


@payment_system_router.callback_query(F.data == "trial_edit_days")
async def trial_edit_days_start(callback: CallbackQuery, state: FSMContext):
    """Редактирование длительности пробного периода"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return

    await callback.answer()
    db = _get_db()
    settings = db.get_payment_settings()
    current = settings.get('trial_days', '14')

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [back_button("trial_settings")]
    ])
    await callback.message.edit_text(
        f"⏱ <b>Длительность пробного периода</b>\n\n"
        f"Текущее значение: <b>{current} дней</b>\n\n"
        "Введите новое количество дней (0 — отключить пробный период):",
        reply_markup=keyboard,
        parse_mode="HTML"
    )
    await state.update_data(anchor_msg_id=callback.message.message_id)
    await state.set_state(PaymentSystemStates.waiting_trial_days)


@payment_system_router.message(PaymentSystemStates.waiting_trial_days)
async def process_trial_days(message: Message, state: FSMContext):
    """Сохранение нового значения длительности пробного периода"""
    db = _get_db()
    raw = message.text.strip()
    _back_kb = InlineKeyboardMarkup(inline_keyboard=[[back_button("trial_settings", "⬅️ К настройкам пробного периода")]])
    if not raw.isdigit():
        await fsm_edit(state, message, "❌ Введите целое число (количество дней):", reply_markup=_back_kb)
        return
    days = int(raw)
    if days > 365:
        await fsm_edit(state, message, "❌ Максимум 365 дней:", reply_markup=_back_kb)
        return
    db.update_payment_setting('trial_days', str(days))
    status = f"отключён (0 дней)" if days == 0 else f"<b>{days} дней</b>"
    await fsm_edit(state, message, f"✅ Длительность пробного периода обновлена: {status}", reply_markup=_back_kb)
    await clear_state_keep_org(state)


@payment_system_router.callback_query(F.data == "trial_edit_plan")
async def trial_edit_plan_start(callback: CallbackQuery, state: FSMContext):
    """Выбор тарифа для пробного периода"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return

    await callback.answer()
    db = _get_db()
    plans = db.get_subscription_plans()

    if not plans:
        await callback.answer("❌ Нет доступных тарифных планов.", show_alert=True)
        return

    # Кнопки с названиями планов
    plan_buttons = [
        [InlineKeyboardButton(text=p[1], callback_data=f"trial_plan_select_{p[1]}")]
        for p in plans
        if p[1] != 'Бесплатный'
    ]
    plan_buttons.append([back_button("trial_settings")])

    settings = db.get_payment_settings()
    current_plan = settings.get('trial_plan', 'Премиум')

    await callback.message.edit_text(
        f"💎 <b>Тариф для пробного периода</b>\n\n"
        f"Текущий: <b>{he(current_plan)}</b>\n\n"
        "Выберите тариф, который получают новые пользователи на время пробного периода:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=plan_buttons),
        parse_mode="HTML"
    )


@payment_system_router.callback_query(F.data.startswith("trial_plan_select_"))
async def trial_plan_selected(callback: CallbackQuery):
    """Сохранение выбранного тарифа для пробного периода"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Доступ только для супер-администратора")
        return

    plan_name = callback.data.removeprefix("trial_plan_select_")
    db = _get_db()
    db.update_payment_setting('trial_plan', plan_name)

    await callback.answer(f"✅ Тариф пробного периода: {plan_name}")
    await trial_settings_menu(callback)


# ===========================================================================
# Выбор провайдера оплаты (СБП / ЮKassa)
# ===========================================================================

@payment_system_router.callback_query(F.data == "payment_provider_select")
async def payment_provider_select(callback: CallbackQuery):
    """Экран выбора провайдера оплаты"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    db = _get_db()
    current = db.get_payment_provider()

    sbp_mark = "✅ " if current == "sbp" else ""
    yk_mark = "✅ " if current == "yookassa" else ""

    text = (
        "🔀 <b>Провайдер оплаты</b>\n\n"
        "Выберите, как пользователи будут оплачивать подписку:\n\n"
        "💳 <b>СБП</b> — пользователь переводит деньги вручную и присылает скриншот.\n"
        "   Вы проверяете и подтверждаете оплату вручную.\n\n"
        "🏦 <b>ЮKassa</b> — пользователь получает ссылку на оплату, подписка\n"
        "   активируется автоматически после успешного платежа.\n\n"
        "⚠️ <i>Для ЮKassa необходимо: ИП/ООО, договор с ЮKassa, Shop ID и секретный ключ.</i>"
    )

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"{sbp_mark}💳 СБП (ручная проверка)", callback_data="set_provider_sbp")],
        [InlineKeyboardButton(text=f"{yk_mark}🏦 ЮKassa (автоматически)", callback_data="set_provider_yookassa")],
        [InlineKeyboardButton(text="⚙️ Настройки ЮKassa", callback_data="yookassa_config_menu")],
        [back_button("payment_settings")],
    ])

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")


@payment_system_router.callback_query(F.data == "set_provider_sbp")
async def set_provider_sbp(callback: CallbackQuery):
    """Переключить на СБП"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Только для супер-администратора", show_alert=True)
        return

    db = _get_db()
    db.set_payment_provider("sbp")
    await callback.answer("✅ Провайдер переключён на СБП")
    await payment_provider_select(callback)


@payment_system_router.callback_query(F.data == "set_provider_yookassa")
async def set_provider_yookassa(callback: CallbackQuery):
    """Переключить на ЮKassa"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Только для супер-администратора", show_alert=True)
        return

    db = _get_db()
    cfg = db.get_yookassa_config()
    if not cfg['shop_id'] or not cfg['secret_key']:
        await callback.answer(
            "⚠️ Сначала заполните Shop ID и секретный ключ в настройках ЮKassa",
            show_alert=True,
        )
        return  # пользователь видит предупреждение и может нажать «⚙️ Настройки ЮKassa»

    db.set_payment_provider("yookassa")
    await callback.answer("✅ Провайдер переключён на ЮKassa")
    await payment_provider_select(callback)


# ---------------------------------------------------------------------------
# Настройки ЮKassa
# ---------------------------------------------------------------------------

@payment_system_router.callback_query(F.data == "yookassa_config_menu")
async def yookassa_config_menu(callback: CallbackQuery):
    """Экран настройки ЮKassa"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    db = _get_db()
    cfg = db.get_yookassa_config()

    shop_id_display = cfg['shop_id'] or "Не установлен"
    # Маскируем секретный ключ — показываем только первые 8 и последние 4 символа
    sk = cfg['secret_key']
    if sk and len(sk) > 12:
        secret_display = sk[:8] + "…" + sk[-4:]
    elif sk:
        secret_display = "••••••••"
    else:
        secret_display = "Не установлен"

    return_url_display = cfg['return_url'] or "Не установлен"

    text = (
        "⚙️ <b>Настройки ЮKassa</b>\n\n"
        f"🏪 <b>Shop ID:</b> {he(shop_id_display)}\n"
        f"🔑 <b>Секретный ключ:</b> {he(secret_display)}\n"
        f"🔗 <b>Return URL:</b> {he(return_url_display)}\n\n"
        "<i>Данные берутся из личного кабинета ЮKassa.\n"
        "Секретный ключ хранится в БД — не передавайте его третьим лицам.</i>"
    )

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🏪 Изменить Shop ID", callback_data="yk_set_shop_id")],
        [InlineKeyboardButton(text="🔑 Изменить секретный ключ", callback_data="yk_set_secret_key")],
        [InlineKeyboardButton(text="🔗 Изменить Return URL", callback_data="yk_set_return_url")],
        [back_button("payment_provider_select")],
    ])

    await callback.message.edit_text(text, reply_markup=keyboard, parse_mode="HTML")


@payment_system_router.callback_query(F.data == "yk_set_shop_id")
async def yk_set_shop_id_start(callback: CallbackQuery, state: FSMContext):
    """Начать ввод Shop ID"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    await state.set_state(PaymentSystemStates.waiting_yookassa_shop_id)
    await state.update_data(anchor_msg_id=callback.message.message_id)

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data="yookassa_config_menu")]
    ])
    await callback.message.edit_text(
        "🏪 <b>Введите Shop ID ЮKassa</b>\n\n"
        "Это числовой идентификатор вашего магазина в личном кабинете ЮKassa.\n"
        "<i>Пример: 123456</i>",
        reply_markup=keyboard,
        parse_mode="HTML",
    )


@payment_system_router.message(PaymentSystemStates.waiting_yookassa_shop_id)
async def yk_set_shop_id_save(message: Message, state: FSMContext):
    """Сохранить Shop ID"""
    from message_utils import fsm_edit
    db = _get_db()
    shop_id = message.text.strip() if message.text else ""
    if not shop_id:
        await fsm_edit(state, message, "❌ Пустое значение. Попробуйте ещё раз.")
        return

    db.set_yookassa_config(shop_id=shop_id)
    await clear_state_keep_org(state)

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⚙️ Настройки ЮKassa", callback_data="yookassa_config_menu")]
    ])
    await fsm_edit(state, message,
                   f"✅ <b>Shop ID сохранён:</b> {he(shop_id)}",
                   reply_markup=keyboard)


@payment_system_router.callback_query(F.data == "yk_set_secret_key")
async def yk_set_secret_key_start(callback: CallbackQuery, state: FSMContext):
    """Начать ввод секретного ключа"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    await state.set_state(PaymentSystemStates.waiting_yookassa_secret_key)
    await state.update_data(anchor_msg_id=callback.message.message_id)

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data="yookassa_config_menu")]
    ])
    await callback.message.edit_text(
        "🔑 <b>Введите секретный ключ ЮKassa</b>\n\n"
        "Находится в личном кабинете ЮKassa → Настройки → Ключи API.\n"
        "<i>Начинается с live_ или test_</i>\n\n"
        "⚠️ <b>Не передавайте ключ посторонним!</b>",
        reply_markup=keyboard,
        parse_mode="HTML",
    )


@payment_system_router.message(PaymentSystemStates.waiting_yookassa_secret_key)
async def yk_set_secret_key_save(message: Message, state: FSMContext):
    """Сохранить секретный ключ"""
    from message_utils import fsm_edit
    db = _get_db()
    secret_key = message.text.strip() if message.text else ""
    if not secret_key:
        await fsm_edit(state, message, "❌ Пустое значение. Попробуйте ещё раз.")
        return

    db.set_yookassa_config(secret_key=secret_key)
    await clear_state_keep_org(state)

    sk = secret_key
    masked = (sk[:8] + "…" + sk[-4:]) if len(sk) > 12 else "••••••••"
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⚙️ Настройки ЮKassa", callback_data="yookassa_config_menu")]
    ])
    await fsm_edit(state, message,
                   f"✅ <b>Секретный ключ сохранён:</b> {he(masked)}",
                   reply_markup=keyboard)


@payment_system_router.callback_query(F.data == "yk_set_return_url")
async def yk_set_return_url_start(callback: CallbackQuery, state: FSMContext):
    """Начать ввод Return URL"""
    if not env_manager.is_super_admin(callback.from_user.id):
        await callback.answer("❌ Только для супер-администратора", show_alert=True)
        return

    await callback.answer()
    await state.set_state(PaymentSystemStates.waiting_yookassa_return_url)
    await state.update_data(anchor_msg_id=callback.message.message_id)

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data="yookassa_config_menu")]
    ])
    await callback.message.edit_text(
        "🔗 <b>Введите Return URL</b>\n\n"
        "Страница, на которую ЮKassa перенаправит пользователя после оплаты.\n"
        "Если бота открывают через t.me — можно указать ссылку на бот.\n"
        "<i>Пример: https://t.me/YourBotName</i>",
        reply_markup=keyboard,
        parse_mode="HTML",
    )


@payment_system_router.message(PaymentSystemStates.waiting_yookassa_return_url)
async def yk_set_return_url_save(message: Message, state: FSMContext):
    """Сохранить Return URL"""
    db = _get_db()
    return_url = message.text.strip() if message.text else ""
    if not return_url:
        await fsm_edit(state, message, "❌ Пустое значение. Попробуйте ещё раз.")
        return

    db.set_yookassa_config(return_url=return_url)
    await clear_state_keep_org(state)

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⚙️ Настройки ЮKassa", callback_data="yookassa_config_menu")]
    ])
    await fsm_edit(state, message,
                   f"✅ <b>Return URL сохранён:</b> {he(return_url)}",
                   reply_markup=keyboard)


# ════════════════════════════════════════════════════════════════════
#  Единая карточка клиента (подписка + модули) и CRUD модульного биллинга
#  Зеркалит web /admin/billing. Источник истины — web/routes/admin_billing.py
# ════════════════════════════════════════════════════════════════════

_DENY = "❌ Доступ только для супер-администратора"
_DEFAULT_ICON = {'module': '📦', 'extension': '🧩', 'bundle': '🎁'}
_WIZ_TITLE = {'module': 'модуль', 'extension': 'расширение', 'bundle': 'пакет'}
_WIZ_FIELDS = ['key', 'name', 'icon', 'description', 'price']
_WIZ_PROMPTS = {
    'key': "ключ (латиница без пробелов, напр. analytics)",
    'name': "название",
    'icon': "иконку (emoji), или «-» для значка по умолчанию",
    'description': "описание, или «-» чтобы пропустить",
    'price': "цену в ₽/мес (число)",
}
_FIELD_KEYMAP = {'name': 'name', 'icon': 'icon', 'description': 'description',
                 'price': 'price_monthly', 'sort': 'sort_order'}


def _sa(uid) -> bool:
    return env_manager.is_super_admin(uid)


def _fmt_date(s) -> str:
    if not s:
        return "бессрочно"
    s = str(s)
    if s[:4] == '9999':
        return "бессрочно"
    return s[:10] if len(s) >= 10 else s


def _fmt_limit(v) -> str:
    if v is None:
        return "—"
    return "∞" if v == -1 else str(v)


def _invalidate(tg_id) -> None:
    try:
        from subscription_utils import invalidate_plan_cache
        invalidate_plan_cache(int(tg_id))
    except Exception as exc:
        logging.warning("invalidate_plan_cache(%s) failed: %s", tg_id, exc)


async def _notify_user(bot, tg_id, text: str) -> None:
    """Уведомить пользователя. Безопасно для email-only (tg_id<0) и ошибок."""
    try:
        if not tg_id or int(tg_id) < 0:
            return
        from notif_utils import add_read_btn
        await bot.send_message(int(tg_id), text, reply_markup=add_read_btn(),
                               parse_mode="HTML")
    except Exception as exc:
        logging.warning("notify_user(%s) failed: %s", tg_id, exc)


# ── Карточка клиента ────────────────────────────────────────────
def _build_customer_card(db, tg_id: int):
    user_id = db.get_user_id(tg_id)
    plan_type, end_date = "Бесплатный", None
    if user_id:
        sub = db.get_user_subscription(user_id)
        if sub:
            plan_type = sub[2] if len(sub) > 2 else "Бесплатный"
            end_date = sub[4] if len(sub) > 4 else None

    limits = None
    for p in db.get_all_subscription_plans():
        if p[1] == plan_type:
            limits = (p[7] if len(p) > 7 else None,
                      p[8] if len(p) > 8 else None,
                      p[9] if len(p) > 9 else None)
            break

    text = "👤 <b>Карточка клиента</b>\n"
    text += f"Telegram ID: <code>{tg_id}</code>\n"
    if not user_id:
        text += "⚠️ Нет личного аккаунта в боте (базовую подписку выдать нельзя).\n"
    text += f"\n📦 <b>Базовая подписка:</b> {he(str(plan_type))}"
    if end_date and str(end_date)[:4] != '9999':
        text += f" (до {_fmt_date(end_date)})"
    text += "\n"
    if limits:
        text += (f"   Лимиты: 📦 {_fmt_limit(limits[0])} · "
                 f"🏪 {_fmt_limit(limits[1])} · "
                 f"💰 {_fmt_limit(limits[2])}/мес\n")

    subs = db.get_billing_module_subs(user_telegram_id=tg_id, active_only=True, limit=100)
    mods = {m['key']: m for m in db.get_all_billing_modules()}
    exts = {e['key']: e for e in db.get_all_billing_extensions()}
    bnds = {b['key']: b for b in db.get_all_billing_bundles()}
    type_ru = {'module': 'модуль', 'extension': 'расширение', 'bundle': 'пакет'}

    text += "\n🧩 <b>Активные модули/расширения/пакеты:</b>\n"
    revoke_buttons = []
    if subs:
        for s in subs:
            it, key = s['item_type'], s['item_key']
            meta = (mods.get(key) if it == 'module'
                    else exts.get(key) if it == 'extension' else bnds.get(key))
            label = f"{meta['icon']} {meta['name']}" if meta else key
            text += f"• {he(label)} ({type_ru.get(it, it)}) — {_fmt_date(s['end_date'])}\n"
            blabel = ("❌ " + (meta['name'] if meta else key))[:40]
            revoke_buttons.append([InlineKeyboardButton(
                text=blabel, callback_data=f"cust_revoke_{s['id']}_{tg_id}")])
    else:
        text += "<i>нет</i>\n"

    kb = [
        [InlineKeyboardButton(text="📦 Базовая подписка", callback_data=f"cust_base_{tg_id}")],
        [InlineKeyboardButton(text="➕ Выдать модуль/пакет", callback_data=f"cust_grant_{tg_id}")],
    ]
    kb += revoke_buttons
    kb += [
        [InlineKeyboardButton(text="🔄 Обновить", callback_data=f"cust_card_{tg_id}")],
        [back_button("manage_subscriptions")],
    ]
    return text, InlineKeyboardMarkup(inline_keyboard=kb)


async def _show_card(callback: CallbackQuery, tg_id: int) -> None:
    text, kb = _build_customer_card(_get_db(), tg_id)
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data == "find_customer")
async def find_customer_start(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer()
    await state.set_state(PaymentSystemStates.waiting_user_telegram_id)
    await state.update_data(anchor_msg_id=callback.message.message_id)
    kb = InlineKeyboardMarkup(inline_keyboard=[[back_button("manage_subscriptions")]])
    await callback.message.edit_text(
        "🔍 <b>Карточка клиента</b>\n\nВведите Telegram ID клиента:",
        reply_markup=kb, parse_mode="HTML")


@payment_system_router.message(PaymentSystemStates.waiting_user_telegram_id)
async def process_find_customer(message: Message, state: FSMContext):
    if not _sa(message.from_user.id):
        return
    raw = (message.text or "").strip()
    try:
        tg_id = int(raw)
    except ValueError:
        await fsm_edit(state, message, "❌ Telegram ID должен быть числом. Попробуйте ещё раз.")
        return
    text, kb = _build_customer_card(_get_db(), tg_id)
    await fsm_edit(state, message, text, reply_markup=kb)
    await clear_state_keep_org(state)


@payment_system_router.callback_query(F.data.startswith("cust_card_"))
async def cust_card_refresh(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer("Обновлено")
    await _show_card(callback, int(callback.data.split("_")[2]))


# ── Базовая подписка клиента ────────────────────────────────────
@payment_system_router.callback_query(F.data.startswith("cust_base_"))
async def cust_base_menu(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer()
    tg_id = int(callback.data.split("_")[2])
    db = _get_db()
    rows = []
    for p in db.get_all_subscription_plans():
        ml = _fmt_limit(p[7] if len(p) > 7 else None)
        sl = _fmt_limit(p[8] if len(p) > 8 else None)
        vl = _fmt_limit(p[9] if len(p) > 9 else None)
        rows.append([InlineKeyboardButton(
            text=f"{p[1]} ({ml}/{sl}/{vl})",
            callback_data=f"cust_setbase_{tg_id}_{p[0]}")])
    rows.append([InlineKeyboardButton(text="🆓 Сбросить на Бесплатный",
                                      callback_data=f"cust_resetbase_{tg_id}")])
    rows.append([back_button(f"cust_card_{tg_id}")])
    await callback.message.edit_text(
        "📦 <b>Базовая подписка</b>\n\nВыберите тариф (лимиты 📦/🏪/💰):",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=rows), parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("cust_setbase_"))
async def cust_set_base(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    parts = callback.data.split("_")
    tg_id, plan_id = int(parts[2]), int(parts[3])
    db = _get_db()
    user_id = db.get_user_id(tg_id)
    if not user_id:
        await callback.answer("❌ У клиента нет личного аккаунта в боте", show_alert=True)
        return
    plan = db.get_subscription_plan_details(plan_id)
    if not plan:
        await callback.answer("❌ Тариф не найден", show_alert=True)
        return
    db.create_subscription(user_id, plan['name'])
    _invalidate(tg_id)
    await _notify_user(callback.bot, tg_id,
                       f"🎉 Вам назначена подписка «{he(plan['name'])}».")
    await callback.answer("✅ Подписка назначена")
    await _show_card(callback, tg_id)


@payment_system_router.callback_query(F.data.startswith("cust_resetbase_"))
async def cust_reset_base(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    tg_id = int(callback.data.split("_")[2])
    db = _get_db()
    user_id = db.get_user_id(tg_id)
    if not user_id:
        await callback.answer("❌ У клиента нет личного аккаунта в боте", show_alert=True)
        return
    db.create_subscription(user_id, "Бесплатный")
    _invalidate(tg_id)
    await callback.answer("✅ Сброшено на Бесплатный")
    await _show_card(callback, tg_id)


# ── Выдача модулей/пакетов клиенту ──────────────────────────────
def _grant_catalog(db):
    cat = []
    for m in db.get_all_billing_modules():
        if m['is_active']:
            cat.append(('module', m['key'], f"{m['icon']} {m['name']}", m['price_monthly']))
    for e in db.get_all_billing_extensions():
        if e['is_active']:
            cat.append(('extension', e['key'], f"{e['icon']} {e['name']}", e['price_monthly']))
    for b in db.get_all_billing_bundles():
        if b['is_active']:
            cat.append(('bundle', b['key'], f"{b['icon']} {b['name']}", b['price_monthly']))
    return cat


@payment_system_router.callback_query(F.data.startswith("cust_grant_"))
async def cust_grant_menu(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    tg_id = int(callback.data.split("_")[2])
    cat = _grant_catalog(_get_db())
    if not cat:
        await callback.answer("Нет активных позиций для выдачи", show_alert=True)
        return
    await callback.answer()
    rows = []
    for idx, (it, key, label, price) in enumerate(cat):
        rows.append([InlineKeyboardButton(
            text=f"{label} · {price:.0f}₽", callback_data=f"gbp_{idx}_{tg_id}")])
    rows.append([back_button(f"cust_card_{tg_id}")])
    await callback.message.edit_text(
        "➕ <b>Выдать доступ</b>\n\nВыберите позицию:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=rows), parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("gbp_"))
async def cust_grant_pick(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    parts = callback.data.split("_")
    idx, tg_id = int(parts[1]), int(parts[2])
    cat = _grant_catalog(_get_db())
    if idx >= len(cat):
        await callback.answer("Позиция устарела, откройте заново", show_alert=True)
        return
    _it, _key, label, _price = cat[idx]
    await callback.answer()
    rows = [
        [InlineKeyboardButton(text="30 дней", callback_data=f"gbd_{idx}_{tg_id}_30")],
        [InlineKeyboardButton(text="90 дней", callback_data=f"gbd_{idx}_{tg_id}_90")],
        [InlineKeyboardButton(text="365 дней", callback_data=f"gbd_{idx}_{tg_id}_365")],
        [InlineKeyboardButton(text="♾ Бессрочно", callback_data=f"gbd_{idx}_{tg_id}_0")],
        [back_button(f"cust_grant_{tg_id}")],
    ]
    await callback.message.edit_text(
        f"➕ <b>{he(label)}</b>\n\nНа какой срок выдать доступ?",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=rows), parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("gbd_"))
async def cust_grant_do(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    parts = callback.data.split("_")
    idx, tg_id, days = int(parts[1]), int(parts[2]), int(parts[3])
    db = _get_db()
    cat = _grant_catalog(db)
    if idx >= len(cat):
        await callback.answer("Позиция устарела, откройте заново", show_alert=True)
        return
    it, key, label, _price = cat[idx]
    sub_id = db.grant_billing_item(
        user_telegram_id=tg_id, item_type=it, item_key=key,
        duration_days=days, price_paid=0.0,
        granted_by=f"sa:{callback.from_user.id}", note="bot grant")
    if not sub_id:
        await callback.answer("❌ Ошибка выдачи", show_alert=True)
        return
    _invalidate(tg_id)
    dur = "бессрочно" if days == 0 else f"на {days} дн."
    await _notify_user(callback.bot, tg_id, f"🎉 Вам выдан доступ: {he(label)} ({dur}).")
    await callback.answer("✅ Выдано")
    await _show_card(callback, tg_id)


@payment_system_router.callback_query(F.data.startswith("cust_revoke_"))
async def cust_revoke_confirm(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    parts = callback.data.split("_")
    sub_id, tg_id = int(parts[2]), int(parts[3])
    await callback.answer()
    rows = [
        [InlineKeyboardButton(text="✅ Да, отозвать",
                              callback_data=f"cust_dorevoke_{sub_id}_{tg_id}")],
        [InlineKeyboardButton(text="🔙 Отмена", callback_data=f"cust_card_{tg_id}")],
    ]
    await callback.message.edit_text(
        "❌ <b>Отозвать доступ?</b>\n\nДля модуля каскадно отзываются его расширения.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=rows), parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("cust_dorevoke_"))
async def cust_revoke_do(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    parts = callback.data.split("_")
    sub_id, tg_id = int(parts[2]), int(parts[3])
    res = _get_db().revoke_billing_item(sub_id, cascade=True)
    _invalidate(tg_id)
    if res.get('revoked'):
        casc = res.get('cascaded', 0)
        await _notify_user(callback.bot, tg_id,
                           "ℹ️ Один из ваших платных доступов был отозван администратором.")
        await callback.answer("✅ Отозвано" + (f" (+{casc} расш.)" if casc else ""))
    else:
        await callback.answer("Нечего отзывать")
    await _show_card(callback, tg_id)


# ════════════════════════════════════════════════════════════════════
#  CRUD каталога: модули / расширения / пакеты
# ════════════════════════════════════════════════════════════════════
def _billing_hub_view(db):
    nm = len(db.get_all_billing_modules())
    ne = len(db.get_all_billing_extensions())
    nb = len(db.get_all_billing_bundles())
    text = ("🧩 <b>Модули и пакеты</b>\n\n"
            "Каталог модульного биллинга (зеркало веб-кабинета).\n\n"
            f"📦 Модулей: {nm}\n🧩 Расширений: {ne}\n🎁 Пакетов: {nb}")
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"📦 Модули ({nm})", callback_data="bm_list")],
        [InlineKeyboardButton(text=f"🧩 Расширения ({ne})", callback_data="be_list")],
        [InlineKeyboardButton(text=f"🎁 Пакеты ({nb})", callback_data="bb_list")],
        [back_button("payment_system_admin")],
    ])
    return text, kb


@payment_system_router.callback_query(F.data == "billing_modules_admin")
async def billing_modules_admin(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer()
    text, kb = _billing_hub_view(_get_db())
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


# ── Модули ──────────────────────────────────────────────────────
def _modules_list_view(db):
    rows = []
    for m in db.get_all_billing_modules():
        st = "✅" if m['is_active'] else "❌"
        rows.append([InlineKeyboardButton(
            text=f"{st} {m['icon']} {m['name']} · {m['price_monthly']:.0f}₽",
            callback_data=f"bm_v_{m['id']}")])
    rows.append([InlineKeyboardButton(text="➕ Добавить модуль", callback_data="bm_add")])
    rows.append([back_button("billing_modules_admin")])
    return "📦 <b>Модули</b>", InlineKeyboardMarkup(inline_keyboard=rows)


def _module_detail(db, mid):
    m = next((x for x in db.get_all_billing_modules() if x['id'] == mid), None)
    if not m:
        return None, None
    st = "✅ активен" if m['is_active'] else "❌ выключен"
    text = (f"📦 <b>Модуль: {he(m['name'])}</b>\n\n"
            f"Ключ: <code>{he(m['key'])}</code>\n"
            f"Иконка: {m['icon']}\n"
            f"Описание: {he(m['description'] or '—')}\n"
            f"Цена: {m['price_monthly']:.0f} ₽/мес\n"
            f"Порядок: {m['sort_order']}\n"
            f"Статус: {st}")
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✏️ Название", callback_data=f"bm_f_{mid}_name"),
         InlineKeyboardButton(text="😀 Иконка", callback_data=f"bm_f_{mid}_icon")],
        [InlineKeyboardButton(text="📝 Описание", callback_data=f"bm_f_{mid}_description"),
         InlineKeyboardButton(text="💰 Цена", callback_data=f"bm_f_{mid}_price")],
        [InlineKeyboardButton(text="🔢 Порядок", callback_data=f"bm_f_{mid}_sort")],
        [InlineKeyboardButton(text="🔄 Вкл/Выкл", callback_data=f"bm_t_{mid}"),
         InlineKeyboardButton(text="🗑 Удалить", callback_data=f"bm_dx_{mid}")],
        [back_button("bm_list")],
    ])
    return text, kb


@payment_system_router.callback_query(F.data == "bm_list")
async def bm_list(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer()
    text, kb = _modules_list_view(_get_db())
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bm_v_"))
async def bm_view(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    mid = int(callback.data.split("_")[2])
    text, kb = _module_detail(_get_db(), mid)
    if not text:
        await callback.answer("Модуль не найден", show_alert=True)
        return
    await callback.answer()
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bm_t_"))
async def bm_toggle(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    mid = int(callback.data.split("_")[2])
    _get_db().toggle_billing_module(mid)
    await callback.answer("Переключено")
    text, kb = _module_detail(_get_db(), mid)
    if text:
        await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bm_dx_"))
async def bm_delete_confirm(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    mid = int(callback.data.split("_")[2])
    await callback.answer()
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Да, удалить", callback_data=f"bm_d_{mid}")],
        [InlineKeyboardButton(text="🔙 Отмена", callback_data=f"bm_v_{mid}")],
    ])
    await callback.message.edit_text(
        "🗑 Удалить модуль вместе со всеми его расширениями?", reply_markup=kb)


@payment_system_router.callback_query(F.data.startswith("bm_d_"))
async def bm_delete(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    mid = int(callback.data.split("_")[2])
    _get_db().delete_billing_module(mid)
    await callback.answer("Удалено")
    text, kb = _modules_list_view(_get_db())
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bm_f_"))
async def bm_field_start(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    parts = callback.data.split("_")
    mid, field = int(parts[2]), parts[3]
    await callback.answer()
    await _field_start(callback, state, 'module', mid, field, f"bm_v_{mid}")


@payment_system_router.callback_query(F.data == "bm_add")
async def bm_add_start(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer()
    await _wizard_begin(callback, state, 'module', {})


# ── Расширения ──────────────────────────────────────────────────
def _extensions_list_view(db):
    rows = []
    for e in db.get_all_billing_extensions():
        st = "✅" if e['is_active'] else "❌"
        rows.append([InlineKeyboardButton(
            text=f"{st} {e['icon']} {e['name']} · {e['price_monthly']:.0f}₽ [{e['module_key']}]",
            callback_data=f"be_v_{e['id']}")])
    rows.append([InlineKeyboardButton(text="➕ Добавить расширение", callback_data="be_addp")])
    rows.append([back_button("billing_modules_admin")])
    return "🧩 <b>Расширения</b>", InlineKeyboardMarkup(inline_keyboard=rows)


def _ext_detail(db, eid):
    e = next((x for x in db.get_all_billing_extensions() if x['id'] == eid), None)
    if not e:
        return None, None
    st = "✅ активно" if e['is_active'] else "❌ выключено"
    text = (f"🧩 <b>Расширение: {he(e['name'])}</b>\n\n"
            f"Модуль: <code>{he(e['module_key'])}</code>\n"
            f"Ключ: <code>{he(e['key'])}</code>\n"
            f"Иконка: {e['icon']}\n"
            f"Описание: {he(e['description'] or '—')}\n"
            f"Цена: {e['price_monthly']:.0f} ₽/мес\n"
            f"Порядок: {e['sort_order']}\n"
            f"Статус: {st}")
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✏️ Название", callback_data=f"be_f_{eid}_name"),
         InlineKeyboardButton(text="😀 Иконка", callback_data=f"be_f_{eid}_icon")],
        [InlineKeyboardButton(text="📝 Описание", callback_data=f"be_f_{eid}_description"),
         InlineKeyboardButton(text="💰 Цена", callback_data=f"be_f_{eid}_price")],
        [InlineKeyboardButton(text="🔢 Порядок", callback_data=f"be_f_{eid}_sort")],
        [InlineKeyboardButton(text="🔄 Вкл/Выкл", callback_data=f"be_t_{eid}"),
         InlineKeyboardButton(text="🗑 Удалить", callback_data=f"be_dx_{eid}")],
        [back_button("be_list")],
    ])
    return text, kb


@payment_system_router.callback_query(F.data == "be_list")
async def be_list(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer()
    text, kb = _extensions_list_view(_get_db())
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data == "be_addp")
async def be_add_pick(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    db = _get_db()
    mods = db.get_all_billing_modules()
    if not mods:
        await callback.answer("Сначала создайте модуль", show_alert=True)
        return
    await callback.answer()
    rows = [[InlineKeyboardButton(text=f"{m['icon']} {m['name']}",
                                  callback_data=f"be_add_{m['key']}")] for m in mods]
    rows.append([back_button("be_list")])
    await callback.message.edit_text(
        "🧩 К какому модулю добавить расширение?",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))


@payment_system_router.callback_query(F.data.startswith("be_add_"))
async def be_add_start(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    modkey = callback.data[len("be_add_"):]
    await callback.answer()
    await _wizard_begin(callback, state, 'extension', {'module_key': modkey})


@payment_system_router.callback_query(F.data.startswith("be_v_"))
async def be_view(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    eid = int(callback.data.split("_")[2])
    text, kb = _ext_detail(_get_db(), eid)
    if not text:
        await callback.answer("Расширение не найдено", show_alert=True)
        return
    await callback.answer()
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("be_t_"))
async def be_toggle(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    eid = int(callback.data.split("_")[2])
    _get_db().toggle_billing_extension(eid)
    await callback.answer("Переключено")
    text, kb = _ext_detail(_get_db(), eid)
    if text:
        await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("be_dx_"))
async def be_delete_confirm(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    eid = int(callback.data.split("_")[2])
    await callback.answer()
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Да, удалить", callback_data=f"be_d_{eid}")],
        [InlineKeyboardButton(text="🔙 Отмена", callback_data=f"be_v_{eid}")],
    ])
    await callback.message.edit_text("🗑 Удалить расширение?", reply_markup=kb)


@payment_system_router.callback_query(F.data.startswith("be_d_"))
async def be_delete(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    eid = int(callback.data.split("_")[2])
    _get_db().delete_billing_extension(eid)
    await callback.answer("Удалено")
    text, kb = _extensions_list_view(_get_db())
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("be_f_"))
async def be_field_start(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    parts = callback.data.split("_")
    eid, field = int(parts[2]), parts[3]
    await callback.answer()
    await _field_start(callback, state, 'extension', eid, field, f"be_v_{eid}")


# ── Пакеты ──────────────────────────────────────────────────────
def _bundles_list_view(db):
    rows = []
    for b in db.get_all_billing_bundles():
        st = "✅" if b['is_active'] else "❌"
        rows.append([InlineKeyboardButton(
            text=f"{st} {b['icon']} {b['name']} · {b['price_monthly']:.0f}₽",
            callback_data=f"bb_v_{b['id']}")])
    rows.append([InlineKeyboardButton(text="➕ Добавить пакет", callback_data="bb_add")])
    rows.append([back_button("billing_modules_admin")])
    return "🎁 <b>Пакеты</b>", InlineKeyboardMarkup(inline_keyboard=rows)


def _bundle_detail(db, bid):
    b = next((x for x in db.get_all_billing_bundles() if x['id'] == bid), None)
    if not b:
        return None, None
    mods = {m['key']: m for m in db.get_all_billing_modules()}
    exts = {e['key']: e for e in db.get_all_billing_extensions()}
    inc = b['includes']
    inc_names = []
    for k in inc.get('modules', []):
        inc_names.append(f"📦 {mods[k]['name']}" if k in mods else f"📦 {k}")
    for k in inc.get('extensions', []):
        inc_names.append(f"🧩 {exts[k]['name']}" if k in exts else f"🧩 {k}")
    st = "✅ активен" if b['is_active'] else "❌ выключен"
    text = (f"🎁 <b>Пакет: {he(b['name'])}</b>\n\n"
            f"Ключ: <code>{he(b['key'])}</code>\n"
            f"Иконка: {b['icon']}\n"
            f"Описание: {he(b['description'] or '—')}\n"
            f"Цена: {b['price_monthly']:.0f} ₽/мес\n"
            f"Порядок: {b['sort_order']}\n"
            f"Статус: {st}\n\n"
            "Состав: " + (", ".join(he(x) for x in inc_names) if inc_names else "<i>пусто</i>"))
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✏️ Название", callback_data=f"bb_f_{bid}_name"),
         InlineKeyboardButton(text="😀 Иконка", callback_data=f"bb_f_{bid}_icon")],
        [InlineKeyboardButton(text="📝 Описание", callback_data=f"bb_f_{bid}_description"),
         InlineKeyboardButton(text="💰 Цена", callback_data=f"bb_f_{bid}_price")],
        [InlineKeyboardButton(text="🔢 Порядок", callback_data=f"bb_f_{bid}_sort")],
        [InlineKeyboardButton(text="🧩 Состав пакета", callback_data=f"bb_inc_{bid}")],
        [InlineKeyboardButton(text="🔄 Вкл/Выкл", callback_data=f"bb_t_{bid}"),
         InlineKeyboardButton(text="🗑 Удалить", callback_data=f"bb_dx_{bid}")],
        [back_button("bb_list")],
    ])
    return text, kb


@payment_system_router.callback_query(F.data == "bb_list")
async def bb_list(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer()
    text, kb = _bundles_list_view(_get_db())
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data == "bb_add")
async def bb_add_start(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await callback.answer()
    await _wizard_begin(callback, state, 'bundle', {})


@payment_system_router.callback_query(F.data.startswith("bb_v_"))
async def bb_view(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    bid = int(callback.data.split("_")[2])
    text, kb = _bundle_detail(_get_db(), bid)
    if not text:
        await callback.answer("Пакет не найден", show_alert=True)
        return
    await callback.answer()
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bb_t_"))
async def bb_toggle(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    bid = int(callback.data.split("_")[2])
    _get_db().toggle_billing_bundle(bid)
    await callback.answer("Переключено")
    text, kb = _bundle_detail(_get_db(), bid)
    if text:
        await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bb_dx_"))
async def bb_delete_confirm(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    bid = int(callback.data.split("_")[2])
    await callback.answer()
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Да, удалить", callback_data=f"bb_d_{bid}")],
        [InlineKeyboardButton(text="🔙 Отмена", callback_data=f"bb_v_{bid}")],
    ])
    await callback.message.edit_text("🗑 Удалить пакет?", reply_markup=kb)


@payment_system_router.callback_query(F.data.startswith("bb_d_"))
async def bb_delete(callback: CallbackQuery):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    bid = int(callback.data.split("_")[2])
    _get_db().delete_billing_bundle(bid)
    await callback.answer("Удалено")
    text, kb = _bundles_list_view(_get_db())
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bb_f_"))
async def bb_field_start(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    parts = callback.data.split("_")
    bid, field = int(parts[2]), parts[3]
    await callback.answer()
    await _field_start(callback, state, 'bundle', bid, field, f"bb_v_{bid}")


# ── Состав пакета (toggle-редактор) ─────────────────────────────
def _bundle_inc_text(draft):
    inc_m = draft.get('inc_modules', [])
    inc_e = draft.get('inc_extensions', [])
    return ("🎁 <b>Состав пакета</b>\n\n"
            "Отметьте включённые модули и расширения, затем «Сохранить».\n\n"
            f"Выбрано: 📦 {len(inc_m)} · 🧩 {len(inc_e)}")


def _bundle_inc_kb(draft):
    db = _get_db()
    inc_m = set(draft.get('inc_modules', []))
    inc_e = set(draft.get('inc_extensions', []))
    rows = []
    for m in db.get_all_billing_modules():
        mark = "☑️" if m['key'] in inc_m else "⬜"
        rows.append([InlineKeyboardButton(text=f"{mark} 📦 {m['name']}",
                                          callback_data=f"bbim_{m['key']}")])
    for e in db.get_all_billing_extensions():
        mark = "☑️" if e['key'] in inc_e else "⬜"
        rows.append([InlineKeyboardButton(text=f"{mark} 🧩 {e['name']}",
                                          callback_data=f"bbie_{e['key']}")])
    rows.append([InlineKeyboardButton(text="💾 Сохранить", callback_data="bbisave")])
    rows.append([InlineKeyboardButton(text="❌ Отмена", callback_data="billing_cancel")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


@payment_system_router.callback_query(F.data.startswith("bb_inc_"))
async def bundle_inc_edit_start(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    bid = int(callback.data[len("bb_inc_"):])
    db = _get_db()
    b = next((x for x in db.get_all_billing_bundles() if x['id'] == bid), None)
    if not b:
        await callback.answer("Пакет не найден", show_alert=True)
        return
    draft = {'edit_id': bid,
             'inc_modules': list(b['includes'].get('modules', [])),
             'inc_extensions': list(b['includes'].get('extensions', []))}
    await state.set_state(PaymentSystemStates.billing_create)
    await state.update_data(wiz_entity='bundle', wiz_step=len(_WIZ_FIELDS),
                            wiz_draft=draft, anchor_msg_id=callback.message.message_id)
    await callback.answer()
    await callback.message.edit_text(_bundle_inc_text(draft),
                                     reply_markup=_bundle_inc_kb(draft), parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bbim_"))
async def bundle_inc_toggle_module(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    key = callback.data[len("bbim_"):]
    data = await state.get_data()
    draft = data.get('wiz_draft', {})
    inc = set(draft.get('inc_modules', []))
    inc.discard(key) if key in inc else inc.add(key)
    draft['inc_modules'] = list(inc)
    await state.update_data(wiz_draft=draft)
    await callback.answer()
    await callback.message.edit_text(_bundle_inc_text(draft),
                                     reply_markup=_bundle_inc_kb(draft), parse_mode="HTML")


@payment_system_router.callback_query(F.data.startswith("bbie_"))
async def bundle_inc_toggle_ext(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    key = callback.data[len("bbie_"):]
    data = await state.get_data()
    draft = data.get('wiz_draft', {})
    inc = set(draft.get('inc_extensions', []))
    inc.discard(key) if key in inc else inc.add(key)
    draft['inc_extensions'] = list(inc)
    await state.update_data(wiz_draft=draft)
    await callback.answer()
    await callback.message.edit_text(_bundle_inc_text(draft),
                                     reply_markup=_bundle_inc_kb(draft), parse_mode="HTML")


@payment_system_router.callback_query(F.data == "bbisave")
async def bundle_inc_save(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    data = await state.get_data()
    draft = data.get('wiz_draft', {})
    db = _get_db()
    includes_json = json.dumps({'modules': draft.get('inc_modules', []),
                                'extensions': draft.get('inc_extensions', [])})
    if draft.get('edit_id'):
        b = next((x for x in db.get_all_billing_bundles() if x['id'] == draft['edit_id']), None)
        ok = bool(b) and db.upsert_billing_bundle(
            b['key'], b['name'], b['icon'], b['description'], includes_json,
            b['price_monthly'], b['sort_order'], int(b['is_active']))
        await clear_state_keep_org(state)
        await callback.answer("Сохранено" if ok else "Ошибка")
        if b:
            text, kb = _bundle_detail(db, draft['edit_id'])
            if text:
                await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")
                return
    else:
        ok = db.upsert_billing_bundle(
            draft['key'], draft['name'], draft['icon'], draft.get('description', ''),
            includes_json, draft['price_monthly'])
        await clear_state_keep_org(state)
        await callback.answer("Создано" if ok else "Ошибка (ключ занят?)")
    text, kb = _bundles_list_view(db)
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


# ── Общий мастер создания и редактирования полей ────────────────
async def _wizard_begin(callback, state, entity, draft):
    await state.set_state(PaymentSystemStates.billing_create)
    await state.update_data(wiz_entity=entity, wiz_step=0, wiz_draft=draft,
                            anchor_msg_id=callback.message.message_id)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data="billing_cancel")]])
    await callback.message.edit_text(
        f"➕ <b>Новый {_WIZ_TITLE[entity]}</b>\n\nВведите {_WIZ_PROMPTS['key']}:",
        reply_markup=kb, parse_mode="HTML")


async def _field_start(callback, state, entity, eid, field, back_cb):
    await state.set_state(PaymentSystemStates.billing_edit_field)
    await state.update_data(be_entity=entity, be_id=eid, be_field=field,
                            anchor_msg_id=callback.message.message_id)
    prompts = {'name': 'название', 'icon': 'иконку (emoji)',
               'description': 'описание (или «-» чтобы очистить)',
               'price': 'цену в ₽/мес', 'sort': 'порядок (число)'}
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔙 Отмена", callback_data=back_cb)]])
    await callback.message.edit_text(
        f"✏️ Введите {prompts.get(field, field)}:", reply_markup=kb)


@payment_system_router.callback_query(F.data == "billing_cancel")
async def billing_cancel(callback: CallbackQuery, state: FSMContext):
    if not _sa(callback.from_user.id):
        await callback.answer(_DENY, show_alert=True)
        return
    await clear_state_keep_org(state)
    await callback.answer("Отменено")
    text, kb = _billing_hub_view(_get_db())
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@payment_system_router.message(PaymentSystemStates.billing_create)
async def process_billing_create(message: Message, state: FSMContext):
    if not _sa(message.from_user.id):
        return
    data = await state.get_data()
    entity = data.get('wiz_entity')
    step = data.get('wiz_step', 0)
    draft = data.get('wiz_draft', {})
    val = (message.text or "").strip()
    field = _WIZ_FIELDS[step]

    if field == 'key':
        k = val.lower().replace(' ', '_')
        if not k or not all(c.isalnum() or c == '_' for c in k):
            await fsm_edit(state, message, "❌ Ключ: только латиница/цифры/_. Повторите.")
            return
        draft['key'] = k
    elif field == 'price':
        try:
            draft['price_monthly'] = float(val.replace(',', '.'))
        except ValueError:
            await fsm_edit(state, message, "❌ Цена должна быть числом. Повторите.")
            return
    elif field == 'icon':
        draft['icon'] = val if val and val != '-' else _DEFAULT_ICON[entity]
    elif field == 'description':
        draft['description'] = '' if val == '-' else val
    else:
        draft[field] = val

    step += 1
    cancel_kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Отмена", callback_data="billing_cancel")]])

    if step < len(_WIZ_FIELDS):
        await state.update_data(wiz_step=step, wiz_draft=draft)
        await fsm_edit(state, message,
                       f"Введите {_WIZ_PROMPTS[_WIZ_FIELDS[step]]}:", reply_markup=cancel_kb)
        return

    if entity == 'bundle':
        draft.setdefault('inc_modules', [])
        draft.setdefault('inc_extensions', [])
        await state.update_data(wiz_step=step, wiz_draft=draft)
        await fsm_edit(state, message, _bundle_inc_text(draft),
                       reply_markup=_bundle_inc_kb(draft))
        return

    db = _get_db()
    if entity == 'module':
        ok = db.upsert_billing_module(
            draft['key'], draft['name'], draft['icon'],
            draft.get('description', ''), draft['price_monthly'])
    else:
        ok = db.upsert_billing_extension(
            draft['module_key'], draft['key'], draft['name'], draft['icon'],
            draft.get('description', ''), draft['price_monthly'])
    await clear_state_keep_org(state)
    list_cb = {'module': 'bm_list', 'extension': 'be_list'}[entity]
    done_kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📋 К списку", callback_data=list_cb)]])
    if ok:
        await fsm_edit(state, message,
                       f"✅ {_WIZ_TITLE[entity].capitalize()} создан.", reply_markup=done_kb)
    else:
        await fsm_edit(state, message,
                       "❌ Не удалось создать (возможно, ключ занят).", reply_markup=done_kb)


@payment_system_router.message(PaymentSystemStates.billing_edit_field)
async def process_billing_edit_field(message: Message, state: FSMContext):
    if not _sa(message.from_user.id):
        return
    data = await state.get_data()
    entity = data.get('be_entity')
    eid = data.get('be_id')
    field = data.get('be_field')
    val = (message.text or "").strip()
    db = _get_db()

    if field == 'price':
        try:
            fval = float(val.replace(',', '.'))
        except ValueError:
            await fsm_edit(state, message, "❌ Цена должна быть числом.")
            return
    elif field == 'sort':
        try:
            fval = int(val)
        except ValueError:
            await fsm_edit(state, message, "❌ Порядок должен быть числом.")
            return
    elif field == 'icon':
        fval = val if val and val != '-' else _DEFAULT_ICON.get(entity, '📦')
    elif field == 'description':
        fval = '' if val == '-' else val
    else:
        fval = val

    kkey = _FIELD_KEYMAP.get(field)
    ok = False
    detail = (None, None)
    if entity == 'module':
        m = next((x for x in db.get_all_billing_modules() if x['id'] == eid), None)
        if m and kkey:
            m = dict(m)
            m[kkey] = fval
            ok = db.upsert_billing_module(
                m['key'], m['name'], m['icon'], m['description'], m['price_monthly'],
                m['sort_order'], int(m['is_active']), m['features_json'])
            detail = _module_detail(db, eid)
    elif entity == 'extension':
        e = next((x for x in db.get_all_billing_extensions() if x['id'] == eid), None)
        if e and kkey:
            e = dict(e)
            e[kkey] = fval
            ok = db.upsert_billing_extension(
                e['module_key'], e['key'], e['name'], e['icon'], e['description'],
                e['price_monthly'], e['sort_order'], int(e['is_active']))
            detail = _ext_detail(db, eid)
    elif entity == 'bundle':
        b = next((x for x in db.get_all_billing_bundles() if x['id'] == eid), None)
        if b and kkey:
            b = dict(b)
            b[kkey] = fval
            ok = db.upsert_billing_bundle(
                b['key'], b['name'], b['icon'], b['description'], b['includes_json'],
                b['price_monthly'], b['sort_order'], int(b['is_active']))
            detail = _bundle_detail(db, eid)

    if ok and detail[0]:
        await fsm_edit(state, message, detail[0], reply_markup=detail[1])
    else:
        await fsm_edit(state, message, "❌ Не удалось сохранить изменение.")
    await clear_state_keep_org(state)
